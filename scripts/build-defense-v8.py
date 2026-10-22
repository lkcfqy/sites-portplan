#!/usr/bin/env python3
"""Generate result-driven Korean V8 materials only after complete reviewed evidence.

No solver, browser, experiment generation, network access or old-material writes.
The published wrapper, restored formal index and independent review must bind to
the same bytes. All validation and PDF layout finish in memory before any write.
"""
from __future__ import annotations

import argparse
from collections import Counter
from html import escape
from io import BytesIO
import hashlib
import json
import math
from pathlib import Path
import shlex

ROOT = Path(__file__).resolve().parents[1]
METHODS = ('stableRetainV5', 'lockedRecoveryV6', 'costPolishV7', 'adaptiveV8', 'highsV8')
LABELS = dict(zip(METHODS, ('V5 유효 계획 유지', 'V6 실행 잠금 복구', 'V7 비용 다듬기', 'V8 조건부 의사결정', '독립 native HiGHS')))
PANELS = {'same-distribution': '동일 생성분포 · 새 시드', 'stress': '새 충격 스트레스'}
PAIRS = {'adaptiveVsV7': ('costPolishV7', 'adaptiveV8'), 'adaptiveVsV5': ('stableRetainV5', 'adaptiveV8'),
         'adaptiveVsV6': ('lockedRecoveryV6', 'adaptiveV8'), 'adaptiveVsHiGHS': ('highsV8', 'adaptiveV8'),
         'highsVsV7': ('costPolishV7', 'highsV8')}
OUTPUTS = ('defense-v8.ko.html', 'defense-v8.ko.pdf', 'submission-v8.ko.md', 'defense-v8.en.md', 'contribution-log-v8.md', 'defense-v8.ko.pptx')


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def relative_path(root, value):
    require(isinstance(value, str), 'Evidence path must be a string')
    path = (root / value).resolve()
    require(path.is_relative_to(root.resolve()), 'Evidence path escapes project: ' + value)
    require(path.is_file(), 'Required completed evidence is missing: ' + value)
    return path


def load(root, name, hashes):
    path = relative_path(root, name)
    hashes[name] = sha(path)
    return json.loads(path.read_text())


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def close(actual, expected, label):
    require(finite(actual) and finite(expected) and abs(actual - expected) <= max(1e-5, abs(expected) * 1e-10),
            'Independent point estimate differs: ' + label)


def same_interval(actual, expected, label, tolerance):
    if expected is None:
        require(actual is None, 'Independent interval null differs: ' + label)
        return
    require(isinstance(actual, list) and isinstance(expected, list) and len(actual) == len(expected) == 2,
            'Malformed independent interval: ' + label)
    for a, b in zip(actual, expected):
        require(finite(a) and finite(b) and abs(a - b) <= tolerance,
                'Independent interval differs: ' + label)


def fnum(value, digits=0):
    if not finite(value):
        return 'null'
    return f'{value:,.{digits}f}'


def pct(value):
    return fnum(value, 2) + '%' if finite(value) else 'null'


def ci(value):
    return '[' + ', '.join(fnum(v, 2) for v in value) + ']' if isinstance(value, list) and len(value) == 2 else 'null'


def verified_inputs(root=ROOT, development=None, development_review=None, bootstrap_review=None):
    hashes = {}
    wrapper = load(root, 'data/v8-results.json', hashes)
    require(wrapper.get('schemaVersion') == 'portplan-v8-published-results' and wrapper.get('executionState') == 'complete',
            'Published V8 results are not a completed formal wrapper')
    reference = wrapper.get('formalIndex', {})
    formal_path = relative_path(root, reference.get('path'))
    require(sha(formal_path) == reference.get('sha256'), 'Restored formal-index SHA differs; restore the exact evidence first')
    hashes[reference['path']] = reference['sha256']
    index = json.loads(formal_path.read_text())
    require(index.get('executionState') == 'complete', 'Formal index is not complete')
    require(wrapper.get('summary') == index.get('summary') and wrapper.get('protocol') == index.get('protocol'),
            'Published summary/protocol differs from the formal index')
    protocol = load(root, 'data/protocol-v8.json', hashes)
    require(protocol == index['protocol'], 'Protocol file differs from frozen formal protocol')
    require(hashes['data/protocol-v8.json'] == index.get('protocolHash'), 'Protocol byte commitment differs')
    require(tuple(protocol.get('methods', [])) == METHODS, 'Unexpected formal method set/order')
    require(protocol['budget']['hardWallMs'] == 5000, 'Unexpected whole-episode budget')
    repeats = protocol.get('algorithmRepeats', [])
    require(len(repeats) == 3 and len({r['seed'] for r in repeats}) == 3, 'Missing predeclared algorithm repeats')
    cases = index.get('cases', [])
    require(len(cases) == 144 and len({c['id'] for c in cases}) == 144, 'Missing/duplicate formal business scenarios')
    require(Counter(c['panel'] for c in cases) == Counter({'same-distribution': 72, 'stress': 72}), 'Panel coverage differs')
    episode_count = 0
    for case in cases:
        require(len(case.get('repeats', [])) == len(repeats), 'Missing repeat in ' + case['id'])
        for i, repetition in enumerate(case['repeats']):
            require(repetition.get('repeatIndex') == i + 1 and repetition.get('algorithmSeed') == repeats[i]['seed'], 'Reordered repeat seed')
            require(set(repetition.get('policies', {})) == set(METHODS), 'Missing policy receipt in ' + case['id'])
            episode_count += len(repetition['policies'])
    require(episode_count == 2160 and wrapper.get('scenarioCount') == len(cases) and wrapper.get('episodeCount') == episode_count,
            'Published scenario/episode counts differ from complete receipts')
    review = load(root, 'data/review-v8.json', hashes)
    require(review.get('valid') is True and review.get('inputIndexSha256') == reference['sha256'],
            'Independent review is not valid or belongs to different result bytes')
    require(review.get('summary', {}).get('episodes') == episode_count, 'Independent review did not cover every episode')
    require(review.get('solverRuns') == 0 and review.get('importsJavaScript') is False, 'Review scope is not the independent read-only ledger')
    summary = index['summary']
    require(summary.get('inferentialState') == 'complete-scenario-cluster-bootstrap', 'Inferential summaries are pending')
    for panel in PANELS:
        published = summary['panels'][panel]
        checked = review.get('pointEstimates', {}).get(panel)
        require(checked is not None and published['scenarioCount'] == 72, 'Missing independently reviewed panel')
        for method in METHODS:
            produced, independent = published['byMethod'][method], checked['byMethod'][method]
            for key in ('recordedEpisodeCount', 'completeEpisodeCount', 'fullyCompletedScenarioCount', 'anyCompletedScenarioCount'):
                require(produced[key] == independent[key], 'Reviewed method count differs: ' + panel + '/' + method + '/' + key)
        for name, (baseline, candidate) in PAIRS.items():
            pair, independent = published['comparisons'][name], checked['comparisons'][name]
            require((pair['baselineMethod'], pair['candidateMethod']) == (baseline, candidate), 'Comparison identity differs')
            require(pair['completion']['gainedEpisodeCount'] == independent['gainedEpisodes'] and
                    pair['completion']['lostEpisodeCount'] == independent['lostEpisodes'], 'Completion gain/loss differs')
            for key in ('total', 'incremental'):
                fee, validated = pair['fees'][key], independent['fees'][key]
                require(fee['pairedScenarioCount'] == validated['pairedScenarios'] and fee['pairedEpisodeCount'] == validated['pairedEpisodes'], 'Fee denominator differs')
                require(fee['scenarioWinTieLoss'] == validated['scenarioWinTieLoss'], 'Fee win/tie/loss differs')
                close(fee['baselineScenarioMeanSumKRW'], validated['baselineScenarioMeanSumKRW'], panel + '/' + name + '/baseline')
                close(fee['candidateScenarioMeanSumKRW'], validated['candidateScenarioMeanSumKRW'], panel + '/' + name + '/candidate')
            require([pair['revisions']['baselineObservedSum'], pair['revisions']['candidateObservedSum']] == independent['revisions'], 'Revision totals differ')
    statistical_review = None
    if bootstrap_review:
        statistical_review = load(root, bootstrap_review, hashes)
        require(statistical_review.get('schemaVersion') == 'portplan-independent-v8-bootstrap-review' and
                statistical_review.get('valid') is True and statistical_review.get('readOnly') is True and
                statistical_review.get('importsJavaScript') is False and statistical_review.get('solverRuns') == 0,
                'Optional bootstrap review is not a valid independent read-only review')
        require(statistical_review.get('inputIndexSha256') == reference['sha256'],
                'Optional bootstrap review belongs to different formal index bytes')
        require(statistical_review.get('frozenManifestSha256') == review.get('frozenManifestSha256') and
                bool(statistical_review.get('frozenManifestSha256')),
                'Statistical and physical reviews bind to different frozen manifests')
        require(statistical_review.get('episodeCount') == episode_count and
                statistical_review.get('businessScenarioCount') == len(cases), 'Bootstrap review coverage differs')
        procedure = statistical_review.get('procedure', {})
        declared = protocol['analysis']['bootstrap']
        require(procedure.get('samplingUnit') == 'whole business scenario' and
                procedure.get('seed') == declared['seed'] and procedure.get('draws') == declared['draws'],
                'Bootstrap review sampling procedure differs from the frozen protocol')
        require(procedure.get('tolerances') == {'feeKRW': 1e-6, 'percentagePoints': 1e-9, 'completionRate': 1e-12},
                'Unexpected bootstrap review numerical tolerances')
        require(set(statistical_review.get('panels', {})) == set(PANELS), 'Missing/extra bootstrap panel')
        for panel in PANELS:
            produced = summary['panels'][panel]
            checked = statistical_review['panels'][panel]
            require(checked.get('scenarioCount') == produced['scenarioCount'] == 72, 'Bootstrap panel scenario count differs')
            require(set(checked.get('byMethod', {})) == set(METHODS) and
                    set(checked.get('comparisons', {})) == set(PAIRS), 'Bootstrap method/comparison coverage differs')
            for method in METHODS:
                for key in ('recordedEpisodeCount', 'completeEpisodeCount', 'fullyCompletedScenarioCount', 'anyCompletedScenarioCount'):
                    require(checked['byMethod'][method][key] == produced['byMethod'][method][key], 'Bootstrap completion eligibility differs')
                a, b = checked['byMethod'][method]['meanScenarioCompletionRate'], produced['byMethod'][method]['meanScenarioCompletionRate']
                require(finite(a) and finite(b) and abs(a-b) <= 1e-12, 'Bootstrap scenario completion rate differs')
            for name in PAIRS:
                pair, validated = produced['comparisons'][name], checked['comparisons'][name]
                same_interval(validated['completionRateGain95CI'], pair['completion']['uncertainty']['meanScenarioRateGain95CI'],
                              panel + '/' + name + '/completion', 1e-12)
                for key in ('total', 'incremental'):
                    original, audited = pair['fees'][key], validated['fees'][key]
                    require(audited['pairedScenarioCount'] == original['pairedScenarioCount'] and
                            audited['pairedEpisodeCount'] == original['pairedEpisodeCount'], 'Bootstrap paired denominator differs')
                    uncertainty = original['uncertainty']
                    require(uncertainty['unit'] == 'scenario' and audited['scenarioCount'] == uncertainty['scenarioCount'] == original['pairedScenarioCount'] and
                            audited['draws'] == uncertainty['draws'] == declared['draws'], 'Bootstrap draw/unit counts differ')
                    same_interval(audited['meanSavings95CI'], uncertainty['meanSavings95CI'], panel + '/' + name + '/' + key + '/KRW', 1e-6)
                    same_interval(audited['weightedSavingsPct95CI'], uncertainty['weightedSavingsPct95CI'], panel + '/' + name + '/' + key + '/percent', 1e-9)
    diagnosis = load(root, 'data/failure-diagnosis-v8.json', hashes)
    require(diagnosis.get('valid') is True and diagnosis.get('fixture') is False and diagnosis.get('sourceArchivesReadOnly') is True,
            'G0 diagnosis is not completed real archive evidence')
    require(diagnosis.get('newHeldOutExperiments') == 0 and diagnosis.get('onlinePolicyReruns') == 0, 'G0 scope unexpectedly changed')
    require(len(diagnosis['cases']) == diagnosis['summary']['failureCaseCount'], 'G0 case coverage differs')
    for path, expected in diagnosis.get('inputHashes', {}).items():
        require(sha(relative_path(root, path)) == expected, 'G0 source commitment differs: ' + path)
    require(Counter(c['category'] for c in diagnosis['cases']) == Counter(diagnosis['summary']['failureCategories']), 'G0 category counts differ')
    dev = None
    if development:
        require(bool(development_review), 'Development evidence requires its independently bound review')
        dev = load(root, development, hashes)
        require(dev.get('schemaVersion') == 'portplan-v8-development-mechanism-results' and
                dev.get('executionState') == 'complete' and dev.get('dataRole') == 'development_only', 'Development experiment is not complete or has the wrong data role')
        dev_review=load(root,development_review,hashes)
        require(dev_review.get('valid') is True and dev_review.get('resultSHA256')==hashes[development] and
                dev_review.get('solverCalls')==0 and dev_review.get('importsJavaScript') is False,
                'Development independent review is invalid or bound to different bytes')
        require(len(dev.get('results',[]))==dev_review.get('resultCount')==dev.get('plannedJobCount'), 'Development review coverage differs')
        require(dev.get('summary',{}).get('resultCount')==len(dev['results']), 'Development summary coverage differs')
    else:
        require(not development_review, 'Development review was supplied without its result')
    return {'root': root, 'wrapper': wrapper, 'index': index, 'protocol': protocol, 'summary': summary,
            'review': review, 'bootstrapReview': statistical_review, 'diagnosis': diagnosis, 'development': dev, 'hashes': hashes,
            'materialOptions': {'development': development, 'development-review': development_review, 'bootstrap-review': bootstrap_review}}


def pair_label(pair):
    return LABELS[pair['candidateMethod']] + ' / ' + LABELS[pair['baselineMethod']]


def comparison_rows(panel):
    return [[pair_label(p), f"{p['fees']['total']['pairedScenarioCount']} / {p['fees']['total']['pairedEpisodeCount']}",
             pct(p['fees']['total']['weightedSavingsPct']), pct(p['fees']['incremental']['weightedSavingsPct']),
             '/'.join(str(p['fees']['total']['scenarioWinTieLoss'][k]) for k in ('win', 'tie', 'loss'))]
            for p in panel['comparisons'].values()]


def build_slides(e):
    p, s, d, r = e['protocol'], e['summary'], e['diagnosis'], e['review']
    n, repeats, methods = len(e['index']['cases']), len(p['algorithmRepeats']), len(p['methods'])
    episodes = r['summary']['episodes']
    same=s['panels']['same-distribution'];stress=s['panels']['stress']
    observed_completions=(f"동일분포 V8 {same['byMethod']['adaptiveV8']['completeEpisodeCount']}/{same['byMethod']['adaptiveV8']['recordedEpisodeCount']} vs "
                          f"V7 {same['byMethod']['costPolishV7']['completeEpisodeCount']}/{same['byMethod']['costPolishV7']['recordedEpisodeCount']}; "
                          f"스트레스 V8 {stress['byMethod']['adaptiveV8']['completeEpisodeCount']}/{stress['byMethod']['adaptiveV8']['recordedEpisodeCount']} vs "
                          f"V7 {stress['byMethod']['costPolishV7']['completeEpisodeCount']}/{stress['byMethod']['costPolishV7']['recordedEpisodeCount']}.")
    v8_declined=all(panel['byMethod']['adaptiveV8']['completeEpisodeCount']<panel['byMethod']['costPolishV7']['completeEpisodeCount'] for panel in (same,stress))
    conclusion='V8의 완주 개선 가설은 지지되지 않았습니다. ' if v8_declined else '완주와 비용의 결과를 함께 해석합니다. '
    slides = []
    def add(title, text, *, table=None, chart=None, flow=None, note='', english=''):
        slides.append({'title': title, 'text': text, 'table': table, 'chart': chart, 'flow': flow, 'note': note, 'english': english})
    add('오늘 어떤 컨테이너를 반출하고,\n언제 적출·반납해야 할까요?',
        '사용자: 알려진 수입 컨테이너와 일별 운송·창고를 함께 조정하는 계획 담당자. PortPlan은 요율 조회를 완전한 날짜 계획과 실행 이력으로 연결합니다.',
        flow=['반출 결정', '창고 적출', '공컨테이너 반납'], note=conclusion+observed_completions+' 제품 기능 추가를 알고리즘 성능 개선이나 기업 실현 절감으로 바꾸어 주장하지 않습니다.',
        english='PortPlan supports a planner coordinating known import containers and daily transport/warehouse resources. It connects published tariffs to a complete pickup–unload–return plan. This is a synthetic daily-model research prototype, not realized business savings. '+
        (f"The V8 completion-improvement hypothesis was not supported: same-distribution V8 {same['byMethod']['adaptiveV8']['completeEpisodeCount']}/{same['byMethod']['adaptiveV8']['recordedEpisodeCount']} versus V7 {same['byMethod']['costPolishV7']['completeEpisodeCount']}/{same['byMethod']['costPolishV7']['recordedEpisodeCount']}; stress V8 {stress['byMethod']['adaptiveV8']['completeEpisodeCount']}/{stress['byMethod']['adaptiveV8']['recordedEpisodeCount']} versus V7 {stress['byMethod']['costPolishV7']['completeEpisodeCount']}/{stress['byMethod']['costPolishV7']['recordedEpisodeCount']}." if v8_declined else 'Read the completion and conditional fee results together.'))
    add('세 행동은 같은 자원을 공유합니다.',
        '반출과 반납은 각각 운송 작업을 사용합니다. 적출은 여러 날 연속 점유할 수 있으며, 반출부터 반납 전까지 체류 공간을 사용합니다.',
        table=(['행동/상태', '공유 자원', '지킬 제약'], [['반출 p', '운송 작업', '방출 가능일 이후'], ['적출 u', '연속 창고 점유', '완료가 납기 이내'], ['반납 r', '운송 작업', '적출 완료와 반납 지연 이후'], ['체류 [p,r)', 'staging', '계획기간 내 모든 날짜의 용량']]),
        note='일 단위 원자적 실행 모델입니다. 도로 경로·예약 슬롯·기사 근무시간·신규 주문 도착은 모델링하지 않았습니다.',
        english='Pickup and return each consume a daily transport movement. Unloading occupies consecutive days; staging spans pickup to before return. This is not a route, appointment, driver-hours or newly arriving order model.')
    add('가격을 먼저 정확히 계산하고,\n전체 계획에서 선택합니다.',
        'PNC 공개 표준 보관료의 최고 도달 구간 소급 방식과 Maersk 한국 표준 D&D 해석을 함께 계산합니다. 한 상자의 요금 노출은 전체 계획의 지연 비용과 다릅니다.',
        flow=['공식 요율·기산일', '완전한 날짜 후보', '비용·자원·잠금 재검사'],
        note='개별 계약·실제 청구서·세금·운송비·창고 임차료는 확인하지 않았습니다. 목표는 두 요금의 선언된 부분 문제입니다.',
        english='Published PNC retroactive tiers and an explicit Maersk D&D interpretation determine model fees. A one-container tariff exposure is not a global delay cost. Private contracts, actual invoices, taxes and transport/warehouse operating costs are not validated.')
    ds = d['summary']
    add('G0: 실패를 모두 같은 문제로\n취급하지 않습니다.',
        f"기존 V7 {ds['archivedPrimaryCaseCount']}개 중 {ds['failureCaseCount']}개 실패를 독립 Python trip 모델로 사후 진단했습니다. D0 실패 {ds['failureDayZeroCount']}개, 실행 이후 실패 {ds['failureAfterExecutionCount']}개입니다.",
        table=(['진단 층', '가능', '불가능', '미결'], [[name, counts.get('feasible', 0), counts.get('infeasible', 0), counts.get('unknown', 0)] for name, counts in ds['layerFeasibilityCounts'].items()]),
        note=f"{ds['requestedLayerCount']}개 층 요청 / 중복 제외 {ds['uniqueDiagnosticModelCount']}개 모델. 추가 오프라인 진단 예산이며 온라인 성과가 아닙니다.",
        english='G0 reconstructs all archived V7 failures in four layers: nominal from start, observed snapshot with executed prefix, hindsight actual from start, and hindsight actual with the same prefix. The latter two are offline counterfactuals, never online information.')
    actual_bad = ds['layerFeasibilityCounts']['actualHindsightFromStart'].get('infeasible', 0)
    development_probes = (e.get('development') or {}).get('summary',{}).get('originalSeedBudgetPairs',[])
    add('G0는 개선 가능한 대상을\n구체적으로 좁힙니다.',
        f"실제 미래를 전부 알아도 불가능한 기존 실패는 {actual_bad}개입니다. 원래 자원·납기 모델이 같다면 전체 완주 상한은 {ds['archivedPrimaryCaseCount']-actual_bad}/{ds['archivedPrimaryCaseCount']} 이하입니다.",
        table=(['기존 실패 분류', '사례 수'], [[key, value] for key, value in ds['failureCategories'].items()]),
        note=('G0 사후 진단 자체는 원 알고리즘 가시간 실험이 아닙니다. 이후 원 seed·고정 V7 코드의 추가 예산 탐침은 별도 개발 결과로 보고합니다. 기존 입력은 새 holdout이 아닙니다.' if development_probes else
              '다른 solver가 추가 오프라인 예산으로 계획을 찾았다는 것은 온라인 정책의 자동 복구나 원 알고리즘 가시간 성공을 뜻하지 않습니다. 기존 입력은 개발 자료입니다.'),
        english='The diagnosis bounds what can be recovered under unchanged resources and deadlines. A feasible offline witness alone does not establish that more time given to the original algorithm would recover it. '+
        ('Subsequent original-seed frozen-V7 extra-budget probes were actually performed and are reported separately as old-input development evidence; they do not add formal holdout scenarios.' if development_probes else 'Old V7 inputs remain development data.'))
    dev = e['development']
    if dev:
        main=dev['summary']['panelSummaries']['stratified-mechanism-200ms']
        dm={'retainOrRecoverV7':'단순 유지/복구','alwaysPolishOrRecoverV7':'매번 다듬기/복구','alwaysRebuildV7':'매번 재구성','adaptiveV8':'V8 조건부 결정'}
        against_retain=main['adaptiveVs']['retainOrRecoverV7']
        dev_table=(['개발 방법','적시 완전 / 시도','늦은 완전 후보','적시 후보 날짜 변경'],
                   [[dm[m],f"{v['timelyCompleteCount']}/{v['taskSeedRunCount']}",v['lateCompleteCount'],v['changedEventCountOnTimelyResults']] for m,v in main['methods'].items()])
        dev_text=(f"동일 snapshot 개발 {dev['taskCount']}개, 전체 호출 {dev['summary']['resultCount']}개를 별도 검사했습니다. "
                  f"V8과 단순 유지의 적시 공동 {against_retain['jointTimelyRunCount']}쌍: "
                  f"더 저렴/동일/더 비쌈 {against_retain['lowerFeeCount']}/{against_retain['equalFeeCount']}/{against_retain['higherFeeCount']}.")
        dev_note='warm 동일 이력 개발 검사입니다. 이 결과에서 V8이 단순 유지보다 낫다고 주장하지 않습니다. 날짜 변경 감소는 더 비싼 요금과 함께 나타날 수 있으며 냉 전체 episode와 합치지 않습니다.'
        dev_english=f"The independently reviewed development run has {dev['taskCount']} fixed snapshots and {dev['summary']['resultCount']} calls. V8 versus simple retain/recover has {against_retain['jointTimelyRunCount']} timely common pairs, fee lower/equal/higher {against_retain['lowerFeeCount']}/{against_retain['equalFeeCount']}/{against_retain['higherFeeCount']}. V8 did not outperform simple retention here; these warm development snapshots do not add new holdout scenarios."
        if development_probes:
            improved=[probe for probe in development_probes if probe['foundOnlyWith1000ms']]
            failed_both=[probe for probe in development_probes if probe['at200ms']=='no_complete_snapshot' and probe['at1000ms']=='no_complete_snapshot']
            dev_note+=f" 원 seed 추가 예산 {len(development_probes)*2}회: 1,000ms에서만 적시 완전 {len(improved)}개, 양쪽 미완료 {len(failed_both)}개. 원래 내부 cache의 정확한 재현은 아닙니다."
            dev_english+=f" An additional {len(development_probes)*2} original-seed fresh-factory frozen-V7 calls compared 200ms and 1000ms on {len(development_probes)} old snapshots; {len(improved)} became timely complete only at 1000ms and {len(failed_both)} remained incomplete at both budgets."
            for probe in improved:
                dev_english+=f" {probe['taskId']} at D{probe['day']} preserves {probe['originalExecutedEventCount']} executed events and returns a model total fee of KRW {probe['totalFeeAt1000ms']} at 1000ms."
            dev_english+=' These extra-budget development probes do not restore the private archived V7 cache or replay a whole cold episode, and do not increase the prospective denominator.'
    else:
        dev_table=(['항목','상태/범위'],[['개발 결과','별도 검증 결과 없음 · not tested'],['원자료','기존 입력 · 개발용'],['실행 의미','warm 단일 snapshot / 동일 실행 잠금'],['공식 새 평가','새 시드 냉 전체 episode · 별도 분모']])
        dev_text='실행된 날짜를 같게 고정한 단일 snapshot 검사는 복구 메커니즘을 확인합니다. 초기 계획부터 달라지는 전체 경로 실험과 결론을 합치지 않습니다.'
        dev_note='독립 검증된 개발 결과를 입력하지 않으면 not tested로 남깁니다. 사후 전지적 해답을 온라인 입력으로 사용하지 않습니다.'
        dev_english='No independently validated development result was supplied. Same-prefix mechanism evaluation remains not tested in these materials.'
    add('같은 실행 이력의 개발 검사에서도\n단순 기준과 비교합니다.',dev_text,table=dev_table,note=dev_note,english=dev_english)
    add('V8: 유효한 계획을 유지하고,\n필요할 때만 개선·복구합니다.',
        '현재 계획을 잠금과 자원으로 재검사합니다. 유효하면 비용 기회를 확인하고 날짜 변경·단기 압력 보호를 적용합니다. 무효이면 잠금 복구를 시도합니다.',
        flow=['현재 계획 유효?\n예 → 유지/조건부 개선', '무효 → 잠금 복구', '완전 후보만 수락'],
        note='단기 압력은 현재 명목 예측의 대리 지표입니다. 미래 완주 보장·전체 경로 비용 비증가·전역 최적 증명이 아닙니다.',
        english='V8 audits its cached plan, retains or conditionally polishes valid dates, and rebuilds invalid plans under exact locks. Date-change and near-term pressure guards are heuristics, not future feasibility or full-path cost guarantees.')
    add('분모와 반복을 먼저 고정했습니다.',
        f"{n}개 업무 시나리오 × {methods}개 정책 × {repeats}개 알고리즘/타이밍 반복 = {episodes:,}개 episode. 반복은 새로운 업무 시나리오가 아닙니다.",
        table=(['패널', '시나리오', '범위'], [[PANELS[panel['id']], panel['scenarioCount'],
            {'same-distribution':'기존 생성기와 충격 유형, 새로운 독립 생성 시드',
             'stress':'반복 감소·미공지 회복 지연·연속 자원 충격·점진 회복; 모든 초기 주문은 알려짐'}[panel['id']]] for panel in p['panels']]),
        note=f"전체 냉 episode 부모 기한 {p['budget']['hardWallMs']:,}ms. 초기화·실행·IPC·부모 검사를 포함합니다. 호출별 내부 시계는 같지 않습니다.",
        english=f'The complete formal run contains {n} business scenarios, {methods} fixed methods and {repeats} retained repeats, totaling {episodes} episodes. Panels stay separate. The common parent deadline includes startup, all decisions, cleanup, message receipt and terminal validation; internal call clocks differ.')
    for panel in PANELS:
        ps = s['panels'][panel]
        add(PANELS[panel] + '\n먼저 완주와 미확보 이력을 확인합니다.',
            '완주는 모든 물리 행동과 최종 검사까지 기한 내 수락한 경우입니다. 미확보 이력은 물리적으로 실패했다고 단정하지 않습니다.',
            chart=[{'label': LABELS[m], 'value': ps['byMethod'][m]['completeEpisodeCount'], 'denominator': ps['byMethod'][m]['recordedEpisodeCount'], 'unknown': ps['byMethod'][m]['executionHistoryUnavailableCount'], 'all3': ps['byMethod'][m]['fullyCompletedScenarioCount']} for m in METHODS],
            note='막대는 episode 완주 수입니다. 세 반복 모두 완주한 시나리오 수와 미확보 이력도 함께 표시합니다.',
            english='Read completion before fees. Chart bars show completed episodes; all-three-repeat scenario counts and unavailable histories remain separate. No pooled ranking combines the two panels.')
        add(PANELS[panel] + '\n같은 사례의 완주 쌍만 요금을 비교합니다.',
            '양수는 후보가 더 저렴합니다. 시나리오별로 일치하는 완주 반복의 평균을 먼저 구하고 시나리오를 동일 가중합니다.',
            table=(['후보 / 기준', '시나리오 / 쌍', '총요금 절감', '추가요금 절감', '낮음/동일/높음'], comparison_rows(ps)),
            note='비교마다 분모가 다릅니다. 총/추가 요금의 원화 분모·영 비용·모든 악화 ID·구간은 보고서에 공개합니다. 실패 요금은 null입니다.',
            english='Positive savings means baseline minus candidate. Average matching jointly completed repeats inside each scenario, then weight scenarios equally. Every comparison has its own selected completed-pair denominator; incomplete costs are null.')
    losses=[]
    for panel, ps in s['panels'].items():
        for name, pair in ps['comparisons'].items():
            fee=pair['fees']['total']
            losses.append([panel, name, len(fee['negativeScenarioIds']), fnum(fee['worstScenarioLossKRW']), f"{pair['completion']['gainedEpisodeCount']} / {pair['completion']['lostEpisodeCount']}"])
    high_revisions=[]
    for panel,ps in s['panels'].items():
        rev=ps['comparisons']['highsVsV7']['revisions']
        high_revisions.append(f"{PANELS[panel]} {fnum(rev['baselineObservedSum'])}→{fnum(rev['candidateObservedSum'])} ({rev['pairedEpisodeCount']}쌍)")
    add('불리한 결과도 같은 화면에 남깁니다.',
        '표의 완주 증감은 episode 기준입니다. 최악 요금 악화는 비교 가능한 시나리오 평균 기준입니다. 단위와 분모를 바꿔 섞지 않습니다.',
        table=(['패널', '비교', '총요금 악화 시나리오', '최악 증가 ₩', '완주 얻음/잃음'], losses),
        note='전체 악화·완주 손실 ID는 보고서에 보존합니다. HiGHS/V7 공통 완주 쌍의 날짜 변경: '+'; '.join(high_revisions)+'. 더 많은 완주·조건부 요금 개선과 변경 대가를 함께 봅니다.',
        english='Report adverse fee scenarios, worst scenario-average loss and gained/lost completions for every declared comparison. All negative scenario/repeat identifiers and lost-completion identifiers are retained in the Korean report. HiGHS also sharply increases date revisions within the paired completed paths: '+
        '; '.join(f"{panel}: {ps['comparisons']['highsVsV7']['revisions']['baselineObservedSum']} to {ps['comparisons']['highsVsV7']['revisions']['candidateObservedSum']} across {ps['comparisons']['highsVsV7']['revisions']['pairedEpisodeCount']} episode pairs" for panel,ps in s['panels'].items())+'.')
    timing=[]
    for panel, ps in s['panels'].items():
        for method in METHODS:
            m=ps['byMethod'][method]
            timing.append([panel, LABELS[method], fnum(m['wallTimeMs']['mean'],1), m['hardDeadlineCount'], f"{m['revisionCount']['observedSum']} / {m['revisionCount']['unknownEpisodeCount']}"])
    add('계산 기한과 날짜 변경의\n대가도 공개합니다.',
        '날짜 변경은 미실행 사건 날짜의 변경 횟수입니다. 작업자 시간·재배차 비용·서로 다른 컨테이너 수로 해석하지 않습니다.',
        table=(['패널', '정책', '평균 부모 ms', '강제 기한', '관측 변경 / 미확보 episode'], timing),
        note='HiGHS native 시작도 전체 부모 기한에 포함됩니다. 중간 호출은 협력식 예산이며, 동일 CPU 또는 절대 실시간 보장은 주장하지 않습니다.',
        english='Timing includes native startup in the common parent budget. Date revisions count pending event-date edits, not labor cost or distinct containers. Unknown counters remain unknown rather than zero.')
    rc=r['summary']
    ci_checked = e.get('bootstrapReview') is not None
    ci_boundary = '별도 Python 통계 검산에서 동일 원본의 구간을 재계산함' if ci_checked else '생산 코드의 구간이며 독립 재계산하지 않음'
    ci_english = ('Bootstrap intervals were separately reproduced by the independent Python statistical review bound to the same formal index.'
                  if ci_checked else 'Bootstrap intervals were not independently recomputed.')
    add('독립 구현은 무엇을 확인했을까요?',
        f"Python 장부가 {rc['episodes']:,}개 episode, 수락 계획 {rc['acceptedPlans']:,}개, 기록된 물리 행동 {rc['physicalEvents']:,}개를 재구성했습니다. 완주 {rc['completeEpisodes']:,}개와 이력 미확보 {rc['historyUnavailable']:,}개를 구분했습니다.",
        table=(['확인한 범위', '경계'], [['요율·일별 자원·실행 잠금', 'JS 행렬을 재사용하지 않는 내부 독립 구현'], ['V8 후보 수락 및 점추정', '공식 결과의 독립 재계산'], ['scenario-cluster bootstrap 구간', ci_boundary], ['현장 효과·전역 최적·제3자 인증', '확인하지 않음']]),
        note='구간은 완주 쌍에 조건부인 시나리오 cluster bootstrap입니다. 완주 선택 편향을 제거하지 않으며 관측 밖 운영 집단에 외삽하지 않습니다.',
        english='The internal Python implementation independently reconstructs input generation, causal observations, execution locks, fees, resources, V8 acceptance guards and point estimates. '+ci_english+' Completed-pair selection bias remains; independent interval reproduction does not establish operational validity or third-party certification.')
    add('내 입력으로 성공과 중단을\n직접 확인할 수 있습니다.',
        '작업대 주문 가져오기 → 모의 용량 사건 설정 → 두 정책 순차 실행 → 날짜별 관측·잠금·권고·실패 검사 → 정확한 입력/결과 내보내기.',
        flow=['주문 + 명시적 가정', '당일 정보·실행 잠금', '권고 + 실패 + 내보내기'],
        note='화면의 10개 주문 완주/중단 예시는 합성 학습용입니다. 공식 연구 시나리오나 실제 현장 실적이 아닙니다. 변경한 용량 가정은 이전 입력과 비교 표시합니다.',
        english='The local page imports a custom scenario, configures simulator-owned events, compares two sequential workers, and exports exact inputs/results. The ten-container examples are illustrative synthetic teaching fixtures, not formal experimental evidence or field deployment.')
    add('증거와 기여의 경계를\n끝까지 유지합니다.',
        '핵심은 알려진 합성 주문의 일별 비용·완주·안정성입니다. 공개 PSPLIB 자원 모델과 벌크항만 로그는 별도의 부분 검증이며 실제 컨테이너 절감으로 합치지 않습니다.',
        table=(['주체/주장', '현재 상태'], [['설계·구현·실험·자료 작성', 'AI 수행; 원 연구·공개 코드·라이선스 귀속 유지'], ['사용자 본인 개발·검증·답변 숙련', 'pending · 대신 인정하지 않음'], ['전문가·기업 검증, 현장 비용 절감', '미수행'], ['개인 참가·AI 해석·등록/서약', '본 자료로 확인하지 않음']]),
        note='알고리즘 최초·SOTA·모든 경우 우월·현실 ROI·우승 확률을 주장하지 않습니다.',
        english='AI performed implementation, experiments and material preparation; original research and software remain attributed. Entrant understanding, expert/enterprise validation, field acceptance, registration and applicable AI conditions are pending and are not certified by these documents.')
    add('심사위원이 같은 결정을\n다시 계산할 수 있게 전달합니다.',
        '전체 입력·모든 반복·실패·실행 장부·고정 source·독립 검산·모든 악화 ID를 연결합니다. 새 재현은 새 출력 경로에 기록하며 최초 결과를 덮어쓰지 않습니다.',
        table=(['자료', '찾을 내용'], [['한국어 보고서', '정확한 분모·원화·구간·모든 불리한 ID'], ['영어 이해 노트', '발표 주제별 가정·결과·한계'], ['기여 기록', 'AI와 사용자 미확인 상태'], ['V8 formal archive + review', '해시로 연결된 원 입력·모든 episode·독립 검사']]),
        note='발표는 오늘의 결정 → 세 행동 → 실제 진단 → 고정 실험 → 손실·한계 → 직접 재현 순서로 진행합니다. 이 자료는 제출이나 서명이 아닙니다.',
        english='Deliver complete inputs, every repeat, all failures, frozen source, independent audits and adverse identifiers with exact hash bindings. Reproductions write new paths. The materials are preparation, not a submission, signature or entrant-understanding attestation.')
    require(len(slides)==18, 'Main deck must contain exactly 18 pages')
    return slides


def markdown_table(headers, rows):
    cell=lambda value:str(value).replace('|', '\\|').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(map(cell,headers))+' |','| '+' | '.join(['---']*len(headers))+' |',
                      *['| '+' | '.join(map(cell,row))+' |' for row in rows]])


def material_texts(e, slides):
    provenance='AI 작성 초안. 참가자 본인의 창작·검증·이해·참가자격 확인서가 아닙니다.'
    report=['# PortPlan V8 · 오늘의 반출·적출·반납 결정', '', provenance, '',
            '공식 완결 결과와 별도 Python 검산이 동일한 원본 SHA-256에 연결된 뒤 작성했습니다. 합성 일 단위 모형 연구이며 기업 실현 절감액이 아닙니다.', '']
    english=['# PortPlan V8 — understanding and defense notes', '',
             'AI-created preparation from complete formal evidence. Not evidence of entrant understanding, authorship, registration or field acceptance.', '']
    for i, slide in enumerate(slides,1):
        report.extend([f"## {i:02d}. {slide['title'].replace(chr(10),' ')}", '',slide['text'],''])
        if slide['flow']:report.extend([' → '.join(slide['flow']).replace('\n',' / '),''])
        if slide['table']:report.extend([markdown_table(*slide['table']),''])
        if slide['chart']:report.extend([markdown_table(['정책','완주 episode / 전체','세 반복 모두 완주 시나리오','미확보 이력'],
            [[c['label'],f"{c['value']} / {c['denominator']}",c['all3'],c['unknown']] for c in slide['chart']]),''])
        report.extend([slide['note'],''])
        english.extend([f"## {i:02d}. {slide['title'].replace(chr(10),' ')}",'',slide['english'],''])
    ci_checked = e.get('bootstrapReview') is not None
    ci_korean = ('구간은 동일 formal index의 별도 Python 통계 검산에서 독립 재계산했습니다. 물리·요율 검산과는 별도이며 완주 쌍 선택 편향은 남습니다.'
                 if ci_checked else '독립 검산은 점추정까지이며 구간은 독립 재계산하지 않았습니다. 완주 쌍 선택 편향을 제거하지 않습니다.')
    ci_english = ('Intervals were separately reproduced by the independent Python statistical review bound to the same formal index; completed-pair selection bias remains.'
                  if ci_checked else 'Intervals were not independently recalculated.')
    ci_label = '독립 검산 95% 구간 %' if ci_checked else '생산 95% 구간 %'
    report.extend(['## 정확한 요금·완주 분모·구간·악화 목록','',
        '절감 = 기준 − 후보. 각 시나리오에서 일치한 완주 반복만 평균한 후 시나리오를 동일 가중합니다. 총요금과 추가요금의 원화 분모가 다릅니다. 0원 분모의 비율은 null입니다.',
        '구간은 percentile scenario-cluster bootstrap입니다. '+ci_korean,''])
    for panel, ps in e['summary']['panels'].items():
        report.extend(['### '+PANELS[panel],''])
        for name,pair in ps['comparisons'].items():
            report.extend(['#### '+name+' · '+pair_label(pair),''])
            headers=['요금','쌍 시나리오/반복','기준 평균합 ₩','후보 평균합 ₩','절감 ₩','절감 %','시나리오 낮음/동일/높음','최악 평균 증가 ₩','0원 기준 시나리오/반복',ci_label]
            rows=[]
            for key in ('total','incremental'):
                f=pair['fees'][key]
                rows.append([key,f"{f['pairedScenarioCount']} / {f['pairedEpisodeCount']}",fnum(f['baselineScenarioMeanSumKRW'],2),
                    fnum(f['candidateScenarioMeanSumKRW'],2),fnum(f['scenarioWeightedSavingsKRW'],2),pct(f['weightedSavingsPct']),
                    '/'.join(str(f['scenarioWinTieLoss'][k]) for k in ('win','tie','loss')),fnum(f['worstScenarioLossKRW'],2),
                    f"{f['zeroBaselineScenarioCount']} / {f['zeroBaselineEpisodeCount']}",ci(f.get('uncertainty',{}).get('weightedSavingsPct95CI'))])
            report.extend([markdown_table(headers,rows),''])
            completion=pair['completion']
            report.extend([f"완주 얻음/잃음 episode: {completion['gainedEpisodeCount']} / {completion['lostEpisodeCount']}. "
                f"시나리오 평균 완주율 차이: {fnum(completion['meanScenarioRateGain'],6)}; {'독립 검산' if ci_checked else '생산'} 95% 구간: {ci(completion.get('uncertainty',{}).get('meanScenarioRateGain95CI'))}.",
                f"비교 가능한 날짜 변경: {pair['revisions']['baselineObservedSum']} → {pair['revisions']['candidateObservedSum']}; {pair['revisions']['pairedEpisodeCount']}쌍.",''])
            for key in ('total','incremental'):
                f=pair['fees'][key]
                report.extend([f"{key} 원 반복 단위 비용합(시나리오 동일가중 합과 다름): 기준 {fnum(f['baselineEpisodeSumKRW'],2)} ₩ / 후보 {fnum(f['candidateEpisodeSumKRW'],2)} ₩.",''])
                for field in ('negativeScenarioIds','negativeEpisodeIds'):
                    report.extend([f'**{key} / {field} — 전체 보존**',''])
                    report.extend(['- `'+item+'`' for item in f[field]] or ['없음.'])
                    report.append('')
            for field in ('gainedEpisodeIds','lostEpisodeIds','gainedFromKnownFailureEpisodeIds','gainedFromUnknownHistoryEpisodeIds'):
                report.extend(['**'+field+' — 전체 보존**',''])
                report.extend(['- `'+item+'`' for item in completion[field]] or ['없음.']);report.append('')
    report.extend(['## G0 전체 실패 층별 상태',''])
    diag=e['diagnosis']
    def feasibility(case,layer):return diag['models'][case['layers'][layer]]['feasibility']
    report.extend([markdown_table(['V7 실패 ID','실패일','원 명목','관측+기존 잠금','전지적 실제','전지적 실제+잠금'],
        [[c['id'],c['originalOnlineFailure']['day'],*[feasibility(c,k) for k in ('nominalFromStart','observedExecutedPrefix','actualHindsightFromStart','actualHindsightExecutedPrefix')]] for c in diag['cases']]),'',
        'G0 독립 오프라인 해답 자체는 원 온라인 알고리즘의 가시간 성공을 뜻하지 않습니다. 이후 실제 수행한 원 seed 추가 예산 개발 검사는 다음 절에 별도 보고합니다.' if (e.get('development') or {}).get('summary',{}).get('originalSeedBudgetPairs') else
        '원 알고리즘 추가 예산 재실행 자료는 이 생성 입력에 없습니다. 독립 오프라인 해답은 원 온라인 알고리즘의 가시간 성공이 아닙니다.',''])
    if e['development']:
        report.extend(['## 동일 이력 개발 자료 원 요약','',
                       '아래는 과거 입력 기반의 별도 개발 자료입니다. warm snapshot 검사는 냉 전체 episode 새 평가와 합치지 않습니다.',
                       '```json',json.dumps(e['development'].get('summary',{}),ensure_ascii=False,indent=2),'```',''])
        probes=e['development'].get('summary',{}).get('originalSeedBudgetPairs',[])
        if probes:
            probe_records=[result for result in e['development']['results'] if result.get('probe') is True and result.get('panel')=='original-seed-budget-probe']
            require(len(probe_records)==len(probes)*2,'Original-seed budget probe records differ from the reported pairs')
            report.extend(['## 원 알고리즘·원 seed 추가 예산 개발 검사 — 실제 수행','',
                f"고정 V7 코드의 fresh factory를 원 scenario.seed로 실행한 {len(probe_records)}개 실제 호출입니다. 같은 기존 snapshot에서 200ms와 1000ms를 비교했습니다. 새 holdout·새 업무 시나리오·전체 냉 episode 성과로 세지 않습니다.",'',
                markdown_table(['기존 task ID','원 seed','실패일','기존 실행 잠금 수','200ms 상태','1000ms 상태','1000ms 완전 총요금 ₩'],
                    [[probe['taskId'],probe['seed'],probe['day'],probe['originalExecutedEventCount'],probe['at200ms'],probe['at1000ms'],fnum(probe['totalFeeAt1000ms'])] for probe in probes]),'',
                'D3 결과는 원 실행 잠금을 보존한 단일 관측 snapshot의 추가 예산 검사입니다. private한 원 V7 cache 상태를 복원하지 않았으므로 아카이브 내부 상태 또는 전체 온라인 episode의 정확한 재실행이라고 부르지 않습니다. 세 D0가 더 큰 예산에서도 실패한 사실과 한 D3의 완전 후보를 모두 보존합니다.',
                'G0 원본의 originalAlgorithmExtraBudgetReplay 필드는 진단 생성 시점의 역사적 상태이며, 이후 수행된 위 개발 실험의 현재 상태를 대신하지 않습니다.',''])
    else:report.extend(['## 동일 이력 개발 자료','별도 완료·검증 자료를 생성 입력으로 제공하지 않았습니다. not tested로 남깁니다.',''])
    hashes=markdown_table(['증거 입력','SHA-256'],list(e['hashes'].items()))
    material_command='python scripts/build-defense-v8.py --require-results'
    for option,value in e.get('materialOptions',{}).items():
        if value:material_command+=' --'+option+' '+shlex.quote(value)
    formal_directory=Path(e['wrapper']['formalIndex']['path']).parent.as_posix()
    report.extend(['## 원본 연결', '',hashes,'',
        '복원된 formal index의 각 inputArtifact와 episode artifact가 원 SHA로 연결됩니다. 공개 wrapper는 전체 장부를 대체하지 않습니다. 원자료 복원과 독립 검산은 새 성과 생성이 아닙니다.',
        '', '## 재현', '', '```bash','python scripts/restore-evidence-v8.py',
        'python scripts/review-v8.py --run '+shlex.quote(formal_directory)+' --output artifacts/review-v8-reproduction.json',
        material_command,'```','',
        '새 실험은 고정 source·protocol을 확인한 별도 출력 경로에서 수행합니다. 최초 결과를 덮어쓰지 않습니다. 개인 참가·AI 조건·등록·서약·본인 발표 숙련은 별도 확인 대상입니다.'])
    english.extend(['## Exact evidence bindings','',hashes,'',
        'See submission-v8.ko.md for every pair-specific total/incremental denominator, all adverse scenario and repeat IDs, gained/lost completions and conditional bootstrap intervals. '+ci_english+' No pooled panel leaderboard, full-path non-regression, SOTA, operational ROI or first-place probability is supported.'])
    contribution=['# PortPlan V8 · 실제 기여 기록','',provenance,'',
        markdown_table(['작업','실제 수행 주체','범위'],[
            ['G0 독립 실패 진단','AI','기존 V7 실패의 사후 모델·잠금·요금 검산; 새 holdout 아님'],
            ['V8 의사결정·native HiGHS·시뮬레이션 구현','AI','source·protocol·실행 장부에 실제 범위 보존'],
            ['고정 새 실험 및 독립 Python 검산','AI','완결 원자료와 검산의 동일 SHA 연결'],
            ['scenario-cluster 구간의 별도 통계 검산','AI' if ci_checked else '미확인',
             '동일 원본 SHA에 연결된 별도 Python 구간 재계산' if ci_checked else '독립 구간 검산 자료를 생성 입력으로 제공하지 않음'],
            ['한국어 자료·PDF·영어 노트 작성','AI','완결 결과에 따른 자동 생성; 발표 숙련 증명 아님'],
            ['사용자 본인 개발·검증·이해','pending','대신 완료로 표시하거나 AI 작업을 개인 창작으로 옮기지 않음'],
            ['전문가·기업·현장 검증','미수행','인터뷰·도입 약속·실현 절감·현장 수락 증거 없음'],
            ['등록·서약·본인 발표·AI 규정 확인','pending','이 자료는 제출·서명·자격 확인이 아님']]),'',
        'GLPK, HiGHS, SciPy, 공개 방법과 원 연구·데이터의 저작권 및 notice는 각 소유자에게 귀속합니다. 내부 독립 구현은 제3자 인증이 아닙니다. PDF는 Noto Sans CJK KR의 필요한 글리프를 변환·부분 내장하며 폰트의 SIL Open Font License 귀속을 유지합니다.',
        '', '## 실제 자료 입력','',hashes]
    return '\n'.join(report)+'\n','\n'.join(english)+'\n','\n'.join(contribution)+'\n'


def html_document(slides):
    parts=['<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PortPlan V8</title><style>',
        'body{margin:0;background:#dce8e0;color:#173c40;font-family:"Noto Sans CJK KR","Noto Sans KR",sans-serif}.slide{box-sizing:border-box;width:1280px;min-height:720px;margin:24px auto;padding:44px 62px;background:#f7faf5;position:relative;page-break-after:always}h1{font-size:35px;line-height:1.4;margin:10px 0 22px;white-space:pre-line}.kicker{font-size:13px;letter-spacing:3px;color:#39826d}p{font-size:20px;line-height:1.7}table{width:100%;border-collapse:collapse;font-size:15px;margin:22px 0}th,td{text-align:left;border-bottom:1px solid #d3e0d3;padding:12px 10px;overflow-wrap:anywhere}th{background:#e6f0e4}.note{font-size:14px;background:#e8f2e5;padding:14px 18px;line-height:1.8}.flow{display:flex;gap:24px;align-items:stretch;margin:38px 0}.flow div{flex:1;background:#e0efe0;border-left:5px solid #36836c;padding:24px;font-size:23px;white-space:pre-line}.chart{margin:22px 0}.row{display:grid;grid-template-columns:240px 1fr 250px;align-items:center;gap:16px;margin:19px 0;font-size:17px}.track{height:21px;background:#dfebde}.fill{height:100%;background:#35826d}.footer{font-size:12px;letter-spacing:.5px;margin-top:25px;color:#638176}@media(max-width:1300px){.slide{width:100%;min-height:0;margin:12px auto;padding:30px}h1{font-size:29px}.row{grid-template-columns:180px 1fr 190px}}@media print{@page{size:1280px 720px;margin:0}body{background:white}.slide{width:1280px;height:720px;margin:0;overflow:hidden}}',
        '</style><body>']
    for i,slide in enumerate(slides,1):
        parts.extend([f'<section class="slide" id="slide-{i}"><div class="kicker">{i:02d} / PORTPLAN V8</div>',
                      '<h1>'+escape(slide['title'])+'</h1><p>'+escape(slide['text'])+'</p>'])
        if slide['flow']:parts.append('<div class="flow">'+''.join('<div>'+escape(text)+'</div>' for text in slide['flow'])+'</div>')
        if slide['table']:
            headers,rows=slide['table'];parts.append('<table><thead><tr>'+''.join('<th>'+escape(str(h))+'</th>' for h in headers)+'</tr></thead><tbody>')
            parts.extend('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows);parts.append('</tbody></table>')
        if slide['chart']:
            parts.append('<div class="chart">')
            for c in slide['chart']:
                width=100*c['value']/c['denominator'] if c['denominator'] else 0
                parts.append(f'<div class="row"><span>{escape(c["label"])}</span><div class="track"><div class="fill" style="width:{width:.4f}%"></div></div><span>{c["value"]}/{c["denominator"]} · 3/3 {c["all3"]} · 미확보 {c["unknown"]}</span></div>')
            parts.append('</div>')
        parts.extend(['<div class="note">'+escape(slide['note'])+'</div>',f'<div class="footer">AI 작성 준비 자료 · 본인 검토·현장 검증 미확인 / {i:02d} / {len(slides):02d}</div></section>'])
    parts.append('</body></html>');return '\n'.join(parts)+'\n'


def register_noto(text):
    """Embed Noto KR; convert only the used CFF glyphs to quadratic outlines in memory."""
    from fontTools import subset
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.cu2quPen import Cu2QuPen
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as PDFont
    path=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    require(path.is_file(), 'Noto Sans CJK font is required for a readable Korean PDF')
    font=TTFont(path,fontNumber=1)
    require('KR' in font['name'].getDebugName(1), 'Expected Noto Korean font face')
    chars=set(ord(c) for c in text if not c.isspace());cmap=font.getBestCmap()
    require(not (chars-set(cmap)), 'Noto is missing required PDF glyphs: '+repr(chars-set(cmap)))
    subsetter=subset.Subsetter();subsetter.populate(text=text);subsetter.subset(font)
    order=font.getGlyphOrder();glyphset=font.getGlyphSet();glyf={}
    for name in order:
        pen=TTGlyphPen(glyphset);glyphset[name].draw(Cu2QuPen(pen,1.0,reverse_direction=True));glyf[name]=pen.glyph()
    builder=FontBuilder(font['head'].unitsPerEm,isTTF=True);builder.setupGlyphOrder(order);builder.setupCharacterMap(font.getBestCmap());builder.setupGlyf(glyf)
    builder.setupHorizontalMetrics(font['hmtx'].metrics);builder.setupHorizontalHeader(ascent=font['hhea'].ascent,descent=font['hhea'].descent)
    builder.setupNameTable({'familyName':'Noto Sans CJK KR PortPlan subset','styleName':'Regular','uniqueFontIdentifier':'PortPlanNotoKR','fullName':'PortPlan Noto KR subset','psName':'PortPlanNotoKR',
                          'copyright':'Noto Sans CJK. Copyright Adobe and Google. SIL Open Font License 1.1. Subset outlines converted for embedding.'})
    builder.setupOS2(sTypoAscender=font['hhea'].ascent,sTypoDescender=font['hhea'].descent,usWinAscent=font['hhea'].ascent,usWinDescent=-font['hhea'].descent)
    builder.setupPost();builder.setupMaxp();buffer=BytesIO();builder.save(buffer);buffer.seek(0);pdfmetrics.registerFont(PDFont('PortPlanNoto',buffer))


def pdf_document(slides):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle
    all_text=json.dumps(slides,ensure_ascii=False)+' PORTPLAN V8 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789'
    register_noto(all_text)
    buffer=BytesIO();c=canvas.Canvas(buffer,pagesize=(1280,720),invariant=1,pageCompression=1)
    c.setTitle('PortPlan V8 · Korean competition preparation');c.setAuthor('AI-assisted PortPlan research; entrant verification pending')
    bg=HexColor('#f7faf5');ink=HexColor('#173c40');green=HexColor('#35826d');light=HexColor('#e6f0e4')
    def paragraph(text,x,y,w,size=20,leading=None):
        style=ParagraphStyle('p',fontName='PortPlanNoto',fontSize=size,leading=leading or size*1.55,textColor=ink,wordWrap='CJK')
        obj=Paragraph(escape(str(text)).replace('\n','<br/>'),style);_,h=obj.wrap(w,700)
        obj.drawOn(c,x,y-h);return y-h
    for index,slide in enumerate(slides,1):
        c.setFillColor(bg);c.rect(0,0,1280,720,stroke=0,fill=1)
        c.setFillColor(green);c.setFont('PortPlanNoto',13);c.drawString(62,672,f'{index:02d} / PORTPLAN V8')
        y=paragraph(slide['title'],62,642,1156,33,43)-15
        y=paragraph(slide['text'],62,y,1156,18,27)-19
        if slide['flow']:
            w=(1156-44)/3;h=118
            for j,text in enumerate(slide['flow']):
                x=62+j*(w+22);c.setFillColor(light);c.roundRect(x,y-h,w,h,7,stroke=0,fill=1);paragraph(text,x+18,y-22,w-36,23,34)
                if j<2:c.setStrokeColor(green);c.setLineWidth(2);c.line(x+w+4,y-h/2,x+w+18,y-h/2);c.line(x+w+13,y-h/2+5,x+w+18,y-h/2);c.line(x+w+13,y-h/2-5,x+w+18,y-h/2)
            y-=h+24
        if slide['chart']:
            for row in slide['chart']:
                paragraph(row['label'],62,y-5,252,17,23)
                c.setFillColor(light);c.rect(322,y-24,580,20,stroke=0,fill=1)
                c.setFillColor(green);c.rect(322,y-24,580*row['value']/row['denominator'] if row['denominator'] else 0,20,stroke=0,fill=1)
                paragraph(f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y-4,294,15,22);y-=53
            y-=8
        if slide['table']:
            headers,rows=slide['table'];dense=len(rows)>7;font_size=11 if dense else 15
            style=ParagraphStyle('cell',fontName='PortPlanNoto',fontSize=font_size,leading=font_size*1.4,textColor=ink,wordWrap='CJK')
            data=[[Paragraph(escape(str(v)),style) for v in row] for row in [headers,*rows]]
            widths=[1156/len(headers)]*len(headers)
            if len(headers)==2:widths=[550,606]
            table=Table(data,colWidths=widths,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),light),('VALIGN',(0,0),(-1,-1),'TOP'),
                ('BOTTOMPADDING',(0,0),(-1,-1),5 if dense else 9),('TOPPADDING',(0,0),(-1,-1),5 if dense else 9),('LEFTPADDING',(0,0),(-1,-1),10),
                ('LINEBELOW',(0,0),(-1,-1),0.5,HexColor('#d3e0d3'))]))
            _,h=table.wrap(1156,600);require(y-h>100,f'PDF table overflows page {index}: {y-h:.1f}');table.drawOn(c,62,y-h);y-=h+18
        note_style=ParagraphStyle('note',fontName='PortPlanNoto',fontSize=12,leading=19,textColor=ink,wordWrap='CJK')
        note=Paragraph(escape(slide['note']),note_style);_,h=note.wrap(1116,300)
        require(y-h-28>64,f'PDF note overflows page {index}: {y-h-28:.1f}')
        c.setFillColor(light);c.roundRect(62,y-h-28,1156,h+25,5,stroke=0,fill=1);note.drawOn(c,82,y-h-14)
        c.setFillColor(HexColor('#638176'));c.setFont('PortPlanNoto',10)
        c.drawString(62,32,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인');c.drawRightString(1218,32,f'PORTPLAN V8 / {index:02d} / {len(slides):02d}');c.showPage()
    c.save();return buffer.getvalue()


def pptx_document(slides):
    """Editable native text, tables and vector shapes; no raster page pictures."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from reportlab.pdfbase import pdfmetrics

    require(len(slides)==18, 'Editable deck must contain exactly 18 slides')
    # pdf_document normally registers the same outline metrics first; direct
    # renderer use remains possible without rendering a PDF or writing a font.
    if 'PortPlanNoto' not in pdfmetrics.getRegisteredFontNames():
        register_noto(json.dumps(slides,ensure_ascii=False)+' PORTPLAN V8 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789')
    presentation=Presentation();presentation.slide_width=Inches(1280/96);presentation.slide_height=Inches(720/96)
    presentation.core_properties.title='PortPlan V8 · Korean competition preparation'
    presentation.core_properties.author='AI-assisted PortPlan research; entrant verification pending'
    presentation.core_properties.subject='Editable result-driven research deck; synthetic daily model'
    colors={'background':'F7FAF5','ink':'173C40','green':'35826D','light':'E6F0E4','muted':'638176'}
    unit=lambda px: Inches(px/96)

    def wrapped(value,width,size):
        lines=[]
        for raw in str(value).split('\n'):
            line='';used=0
            for char in raw:
                advance=pdfmetrics.stringWidth(char,'PortPlanNoto',size)
                if line and used+advance>width:
                    lines.append(line.rstrip());line='';used=0
                line+=char;used+=advance
            lines.append(line.rstrip())
        return '\n'.join(lines),len(lines)

    def style_text(frame,text,size,leading,color='ink'):
        frame.clear();frame.text=text;frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        frame.vertical_anchor=MSO_ANCHOR.TOP
        for paragraph in frame.paragraphs:
            paragraph.font.name='Noto Sans CJK KR';paragraph.font.size=Pt(size*.75)
            paragraph.font.color.rgb=RGBColor.from_string(colors[color])
            paragraph.line_spacing=Pt(leading*.75);paragraph.space_before=Pt(0);paragraph.space_after=Pt(0)
            properties=paragraph._p.get_or_add_pPr().get_or_add_defRPr()
            east_asian=properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
            if east_asian is None:
                east_asian=OxmlElement('a:ea');properties.append(east_asian)
            east_asian.set('typeface','Noto Sans CJK KR')

    def text(slide,value,x,y,width,size=18,leading=27,color='ink',height=None,name='Editable text'):
        content,line_count=wrapped(value,width-2,size)
        measured=line_count*leading+3
        box=slide.shapes.add_textbox(unit(x),unit(y),unit(width),unit(height or measured));box.name=name
        style_text(box.text_frame,content,size,leading,color)
        return y+(height or measured)

    def rect(slide,x,y,width,height,color='light',kind=MSO_SHAPE.RECTANGLE,name='Editable vector'):
        shape=slide.shapes.add_shape(kind,unit(x),unit(y),unit(width),unit(height));shape.name=name
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(colors[color]);shape.line.fill.background()
        return shape

    for index,source in enumerate(slides,1):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string(colors['background'])
        text(slide,f'{index:02d} / PORTPLAN V8',62,31,1156,13,20,'green',name='Slide index')
        y=text(slide,source['title'],62,72,1156,33,43,name='Slide title')+15
        y=text(slide,source['text'],62,y,1156,18,27,name='Main explanation')+19
        if source['flow']:
            width=(1156-44)/3
            for j,label in enumerate(source['flow']):
                x=62+j*(width+22);rect(slide,x,y,width,118,name='Flow node '+str(j+1))
                rect(slide,x,y,5,118,'green',name='Flow node accent')
                text(slide,label,x+18,y+22,width-36,23,34,height=80,name='Editable flow label')
                if j<2:rect(slide,x+width+4,y+52,14,14,'green',MSO_SHAPE.CHEVRON,'Editable flow arrow')
            y+=142
        if source['chart']:
            for row in source['chart']:
                text(slide,row['label'],62,y+5,252,17,23,name='Chart method label')
                rect(slide,322,y+4,580,20,name='Editable completion bar track')
                width=580*row['value']/row['denominator'] if row['denominator'] else 0
                if width:rect(slide,322,y+4,width,20,'green',name='Editable completion bar')
                text(slide,f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y+4,294,15,22,name='Chart exact value')
                y+=53
            y+=8
        if source['table']:
            headers,rows=source['table'];dense=len(rows)>7;size=11 if dense else 15;leading=size*1.4;padding=5 if dense else 9
            widths=[1156/len(headers)]*len(headers)
            if len(headers)==2:widths=[550,606]
            formatted=[];heights=[]
            for row in [headers,*rows]:
                cells=[wrapped(value,widths[j]-20,size) for j,value in enumerate(row)]
                formatted.append([value for value,_ in cells]);heights.append(max(count for _,count in cells)*leading+2*padding+3)
            height=sum(heights);require(y+height<620,f'Editable table overflows slide {index}')
            native=slide.shapes.add_table(len(formatted),len(headers),unit(62),unit(y),unit(1156),unit(height));native.name='Editable evidence table'
            table=native.table;table.first_row=True
            for j,width in enumerate(widths):table.columns[j].width=unit(width)
            for i,row in enumerate(formatted):
                table.rows[i].height=unit(heights[i])
                for j,value in enumerate(row):
                    cell=table.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(colors['light' if i==0 else 'background'])
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
                    style_text(cell.text_frame,value,size,leading)
                    # Keep table padding after generic text-frame styling.
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
            y+=height+18
        note,note_lines=wrapped(source['note'],1116,12);note_height=note_lines*19+3
        require(y+note_height+28<656,f'Editable note overflows slide {index}')
        rect(slide,62,y,1156,note_height+25,name='Editable boundary note background')
        text(slide,note,82,y+11,1116,12,19,height=note_height+3,name='Evidence boundary note')
        text(slide,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인',62,678,900,10,15,'muted',name='Contribution footer')
        text(slide,f'PORTPLAN V8 / {index:02d} / {len(slides):02d}',1040,678,178,10,15,'muted',name='Page footer')
        for shape in slide.shapes:
            require(shape.shape_type!=MSO_SHAPE_TYPE.PICTURE,'Raster slide pictures are prohibited in the editable deck')
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=presentation.slide_width+1 and
                    shape.top+shape.height<=presentation.slide_height+1,f'Editable shape escapes slide {index}: {shape.name}')
    require(len(presentation.slides)==18,'Editable deck page count changed')
    output=BytesIO();presentation.save(output);return output.getvalue()


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--require-results',action='store_true',help='Completeness is always required; retained for release commands.')
    parser.add_argument('--check-only',action='store_true',help='Validate evidence without writing materials.')
    parser.add_argument('--development',help='Optional completed same-prefix development result, never a new holdout.')
    parser.add_argument('--development-review',help='Independent review bound to the optional development result SHA.')
    parser.add_argument('--bootstrap-review',help='Optional independent interval review bound to the same formal index SHA; invalid supplied evidence is rejected.')
    parser.add_argument('--replace-generated',action='store_true',help='Explicitly regenerate these V8 materials only.')
    args=parser.parse_args(argv)
    evidence=verified_inputs(development=args.development,development_review=args.development_review,bootstrap_review=args.bootstrap_review)
    slides=build_slides(evidence)
    if args.check_only:
        print(json.dumps({'valid':True,'pages':len(slides),'writes':0,'evidenceHashes':evidence['hashes']}));return
    targets=[ROOT/'docs'/name for name in OUTPUTS]
    if not args.replace_generated:require(not any(path.exists() for path in targets),'V8 materials already exist; use --replace-generated explicitly')
    report,english,contribution=material_texts(evidence,slides)
    pdf=pdf_document(slides)
    pptx=pptx_document(slides)
    require(pdf.startswith(b'%PDF-'), 'PDF generation failed')
    require(pptx.startswith(b'PK'), 'Editable presentation generation failed')
    outputs={targets[0]:html_document(slides).encode(),targets[1]:pdf,targets[2]:report.encode(),targets[3]:english.encode(),targets[4]:contribution.encode(),targets[5]:pptx}
    for path,data in outputs.items():path.write_bytes(data)
    print(json.dumps({'valid':True,'pages':len(slides),'written':[str(p.relative_to(ROOT)) for p in outputs],
        'outputHashes':{str(p.relative_to(ROOT)):sha(p) for p in outputs}},ensure_ascii=False))


if __name__=='__main__':
    main()
