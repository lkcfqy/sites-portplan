#!/usr/bin/env python3
"""Generate result-driven Korean V9 materials only after complete reviewed evidence.

No solver, browser, experiment generation, network access or old-material writes.
The published wrapper, restored formal index and independent review must bind to
the same bytes. All validation and PDF layout finish in memory before any write.
"""
from __future__ import annotations

import argparse
from collections import Counter
from html import escape
from io import BytesIO
import hashlib
import json
import math
from pathlib import Path
import shlex

ROOT = Path(__file__).resolve().parents[1]
METHODS = ('stableRetainV5', 'costPolishV7', 'adaptiveV8', 'adaptiveV9', 'highsV8')
LABELS = dict(zip(METHODS, ('V5 유효 계획 유지', 'V7 비용 다듬기', 'V8 조건부 의사결정', 'V9 적시 후보 전달', '독립 native HiGHS')))
PANELS = {'same-distribution': '동일 생성분포 · 새 주문 시드', 'stress': '스트레스 · 새 주문 시드'}
PAIRS = {'v9VsV8': ('adaptiveV8', 'adaptiveV9'), 'v9VsV7': ('costPolishV7', 'adaptiveV9'),
         'v9VsV5': ('stableRetainV5', 'adaptiveV9'), 'v9VsHiGHS': ('highsV8', 'adaptiveV9'),
         'highsVsV7': ('costPolishV7', 'highsV8')}
OUTPUTS = ('defense-v9.ko.html', 'defense-v9.ko.pdf', 'submission-v9.ko.md', 'defense-v9.en.md', 'contribution-log-v9.md', 'defense-v9.ko.pptx')


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def relative_path(root, value):
    require(isinstance(value, str), 'Evidence path must be a string')
    path = (root / value).resolve()
    require(path.is_relative_to(root.resolve()), 'Evidence path escapes project: ' + value)
    require(path.is_file(), 'Required completed evidence is missing: ' + value)
    return path


def load(root, name, hashes):
    path = relative_path(root, name)
    hashes[name] = sha(path)
    return json.loads(path.read_text())


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def close(actual, expected, label):
    require(finite(actual) and finite(expected) and abs(actual - expected) <= max(1e-5, abs(expected) * 1e-10),
            'Independent point estimate differs: ' + label)


def same_interval(actual, expected, label, tolerance):
    if expected is None:
        require(actual is None, 'Independent interval null differs: ' + label)
        return
    require(isinstance(actual, list) and isinstance(expected, list) and len(actual) == len(expected) == 2,
            'Malformed independent interval: ' + label)
    for a, b in zip(actual, expected):
        require(finite(a) and finite(b) and abs(a - b) <= tolerance,
                'Independent interval differs: ' + label)


def fnum(value, digits=0):
    if not finite(value):
        return 'null'
    return f'{value:,.{digits}f}'


def pct(value):
    return fnum(value, 2) + '%' if finite(value) else 'null'


def ci(value):
    return '[' + ', '.join(fnum(v, 2) for v in value) + ']' if isinstance(value, list) and len(value) == 2 else 'null'


def verified_inputs(root=ROOT, bootstrap_review='data/bootstrap-review-v9.json'):
    """No fixture, partial run, unbound summary or unreviewed interval may publish."""
    hashes = {}
    wrapper = load(root, 'data/v9-results.json', hashes)
    require(wrapper.get('schemaVersion') == 'portplan-v9-published-results' and
            wrapper.get('executionState') == 'complete', 'Published V9 results are not a completed formal wrapper')
    reference = wrapper.get('formalIndex', {})
    formal_path = relative_path(root, reference.get('path'))
    require(sha(formal_path) == reference.get('sha256'), 'Restored formal-index SHA differs; restore exact evidence first')
    hashes[reference['path']] = reference['sha256']
    index = json.loads(formal_path.read_text())
    require(index.get('executionState') == 'complete', 'Formal V9 index is not complete')
    require(wrapper.get('summary') == index.get('summary') and wrapper.get('protocol') == index.get('protocol'),
            'Published summary/protocol differs from the formal index')
    protocol = load(root, 'data/protocol-v9.json', hashes)
    require(protocol == index['protocol'] and hashes['data/protocol-v9.json'] == index.get('protocolHash'),
            'Protocol bytes differ from frozen formal protocol')
    require(tuple(protocol.get('methods', [])) == METHODS, 'Unexpected formal method set/order')
    require(protocol['budget']['hardWallMs'] == 5000, 'Unexpected whole-episode budget')
    repeats = protocol.get('algorithmRepeats', [])
    require(len(repeats) == 3 and len({r['seed'] for r in repeats}) == 3, 'Missing predeclared algorithm repeats')
    cases = index.get('cases', [])
    require(len(cases) == 144 and len({c['id'] for c in cases}) == 144, 'Missing/duplicate formal business scenarios')
    require(Counter(c['panel'] for c in cases) == Counter({'same-distribution': 72, 'stress': 72}), 'Panel coverage differs')
    episode_count = 0
    for case in cases:
        require(len(case.get('repeats', [])) == len(repeats), 'Missing repeat in ' + case['id'])
        for i, repetition in enumerate(case['repeats']):
            require(repetition.get('repeatIndex') == i + 1 and repetition.get('algorithmSeed') == repeats[i]['seed'], 'Reordered repeat seed')
            require(set(repetition.get('policies', {})) == set(METHODS), 'Missing policy receipt in ' + case['id'])
            episode_count += len(repetition['policies'])
    require(episode_count == 2160 and wrapper.get('scenarioCount') == len(cases) and wrapper.get('episodeCount') == episode_count,
            'Published scenario/episode counts differ from complete receipts')
    review = load(root, 'data/review-v9.json', hashes)
    require(review.get('schemaVersion') == 'portplan-independent-v9-run-review' and review.get('valid') is True and
            review.get('fixture') is False and review.get('inputIndexSha256') == reference['sha256'],
            'Independent physical review is invalid, a fixture or bound to different result bytes')
    require(review.get('summary', {}).get('episodes') == episode_count, 'Independent review did not cover every episode')
    require(review.get('solverRuns') == 0 and review.get('importsJavaScript') is False,
            'Review scope is not the independent read-only ledger')
    summary = index['summary']
    require(summary.get('inferentialState') == 'complete-scenario-cluster-bootstrap', 'Inferential summaries are pending')
    require(set(summary.get('panels', {})) == set(PANELS), 'Missing/extra formal panel')
    for panel in PANELS:
        published = summary['panels'][panel]
        checked = review.get('pointEstimates', {}).get(panel)
        require(checked is not None and published['scenarioCount'] == 72, 'Missing independently reviewed panel')
        require(set(published.get('byMethod', {})) == set(METHODS) and
                set(published.get('comparisons', {})) == set(PAIRS), 'Unexpected method/comparison coverage')
        for method in METHODS:
            produced, independent = published['byMethod'][method], checked['byMethod'][method]
            for key in ('recordedEpisodeCount', 'completeEpisodeCount', 'fullyCompletedScenarioCount', 'anyCompletedScenarioCount'):
                require(produced[key] == independent[key], 'Reviewed method count differs: ' + panel + '/' + method + '/' + key)
        for name, (baseline, candidate) in PAIRS.items():
            pair, independent = published['comparisons'][name], checked['comparisons'][name]
            require((pair['baselineMethod'], pair['candidateMethod']) == (baseline, candidate), 'Comparison identity differs')
            require(pair['completion']['gainedEpisodeCount'] == independent['gainedEpisodes'] and
                    pair['completion']['lostEpisodeCount'] == independent['lostEpisodes'], 'Completion gain/loss differs')
            for key in ('total', 'incremental'):
                fee, validated = pair['fees'][key], independent['fees'][key]
                require(fee['pairedScenarioCount'] == validated['pairedScenarios'] and fee['pairedEpisodeCount'] == validated['pairedEpisodes'], 'Fee denominator differs')
                require(fee['scenarioWinTieLoss'] == validated['scenarioWinTieLoss'], 'Fee win/tie/loss differs')
                close(fee['baselineScenarioMeanSumKRW'], validated['baselineScenarioMeanSumKRW'], panel + '/' + name + '/baseline')
                close(fee['candidateScenarioMeanSumKRW'], validated['candidateScenarioMeanSumKRW'], panel + '/' + name + '/candidate')
            require([pair['revisions']['baselineObservedSum'], pair['revisions']['candidateObservedSum']] == independent['revisions'], 'Revision totals differ')
    require(bool(bootstrap_review), 'A separate independent bootstrap review is required for V9 materials')
    statistics = load(root, bootstrap_review, hashes)
    require(statistics.get('schemaVersion') == 'portplan-independent-v9-bootstrap-review' and
            statistics.get('valid') is True and statistics.get('fixture', False) is False and
            statistics.get('readOnly') is True and statistics.get('importsJavaScript') is False and
            statistics.get('solverRuns') == 0, 'Bootstrap review is not a valid independent read-only formal review')
    require(statistics.get('inputIndexSha256') == reference['sha256'], 'Bootstrap review belongs to different index bytes')
    require(statistics.get('frozenManifestSha256') == review.get('frozenManifestSha256') and
            bool(statistics.get('frozenManifestSha256')), 'Statistical and physical reviews bind to different frozen manifests')
    require(statistics.get('episodeCount') == episode_count and statistics.get('businessScenarioCount') == len(cases),
            'Bootstrap review coverage differs')
    procedure = statistics.get('procedure', {})
    declared = protocol['analysis']['bootstrap']
    require(procedure.get('samplingUnit') == 'whole business scenario' and procedure.get('seed') == declared['seed'] and
            procedure.get('draws') == declared['draws'], 'Bootstrap sampling procedure differs from frozen protocol')
    require(procedure.get('tolerances') == {'feeKRW': 1e-6, 'percentagePoints': 1e-9, 'completionRate': 1e-12},
            'Unexpected bootstrap numerical tolerances')
    require(set(statistics.get('panels', {})) == set(PANELS), 'Missing/extra bootstrap panel')
    for panel in PANELS:
        produced, checked = summary['panels'][panel], statistics['panels'][panel]
        require(checked.get('scenarioCount') == produced['scenarioCount'] == 72, 'Bootstrap panel scenario count differs')
        require(set(checked.get('byMethod', {})) == set(METHODS) and set(checked.get('comparisons', {})) == set(PAIRS),
                'Bootstrap method/comparison coverage differs')
        for method in METHODS:
            for key in ('recordedEpisodeCount', 'completeEpisodeCount', 'fullyCompletedScenarioCount', 'anyCompletedScenarioCount'):
                require(checked['byMethod'][method][key] == produced['byMethod'][method][key], 'Bootstrap completion eligibility differs')
            close(checked['byMethod'][method]['meanScenarioCompletionRate'], produced['byMethod'][method]['meanScenarioCompletionRate'],
                  panel + '/' + method + '/scenario-completion-rate')
        for name in PAIRS:
            pair, validated = produced['comparisons'][name], checked['comparisons'][name]
            same_interval(validated['completionRateGain95CI'], pair['completion']['uncertainty']['meanScenarioRateGain95CI'],
                          panel + '/' + name + '/completion', 1e-12)
            for key in ('total', 'incremental'):
                original, audited = pair['fees'][key], validated['fees'][key]
                uncertainty = original['uncertainty']
                require(audited['pairedScenarioCount'] == original['pairedScenarioCount'] and
                        audited['pairedEpisodeCount'] == original['pairedEpisodeCount'], 'Bootstrap paired denominator differs')
                require(uncertainty['unit'] == 'scenario' and audited['scenarioCount'] == uncertainty['scenarioCount'] == original['pairedScenarioCount'] and
                        audited['draws'] == uncertainty['draws'] == declared['draws'], 'Bootstrap draw/unit counts differ')
                same_interval(audited['meanSavings95CI'], uncertainty['meanSavings95CI'], panel + '/' + name + '/' + key + '/KRW', 1e-6)
                same_interval(audited['weightedSavingsPct95CI'], uncertainty['weightedSavingsPct95CI'], panel + '/' + name + '/' + key + '/percent', 1e-9)
    # This is a separate previous experiment, never a V9 paired denominator.
    previous = load(root, 'data/v8-results.json', hashes)
    require(previous.get('schemaVersion') == 'portplan-v8-published-results' and previous.get('executionState') == 'complete',
            'Historical V8 wrapper is not complete')
    old_reference = previous.get('formalIndex', {})
    old_path = relative_path(root, old_reference.get('path'))
    require(sha(old_path) == old_reference.get('sha256'), 'Historical V8 index hash differs')
    hashes[old_reference['path']] = old_reference['sha256']
    old_index = json.loads(old_path.read_text())
    require(old_index.get('executionState') == 'complete' and previous.get('summary') == old_index.get('summary'),
            'Historical V8 summary differs from its complete first run')
    old_review = load(root, 'data/review-v8.json', hashes)
    require(old_review.get('valid') is True and old_review.get('fixture') is False and
            old_review.get('inputIndexSha256') == old_reference['sha256'], 'Historical V8 review is not bound to the historical run')
    diagnosis = load(root, 'data/failure-diagnosis-v8.json', hashes)
    require(diagnosis.get('valid') is True and diagnosis.get('fixture') is False and diagnosis.get('sourceArchivesReadOnly') is True,
            'G0 diagnosis is not completed real archive evidence')
    require(diagnosis.get('newHeldOutExperiments') == 0 and diagnosis.get('onlinePolicyReruns') == 0, 'G0 scope changed')
    for path, expected in diagnosis.get('inputHashes', {}).items():
        require(sha(relative_path(root, path)) == expected, 'G0 source commitment differs: ' + path)
    require(len(diagnosis['cases']) == diagnosis['summary']['failureCaseCount'] and
            Counter(c['category'] for c in diagnosis['cases']) == Counter(diagnosis['summary']['failureCategories']), 'G0 coverage differs')
    resource = None
    resource_paths = ('data/resource-validation-v9.json', 'data/resource-review-v9.json')
    if any((root / name).exists() for name in resource_paths):
        teaching, resource_review = (load(root, name, hashes) for name in resource_paths)
        require(teaching.get('schemaVersion') == 'portplan-v9-resource-teaching-validation' and teaching.get('fixture') is True,
                'Resource illustration must be explicitly synthetic teaching data')
        require(resource_review.get('schemaVersion') == 'portplan-independent-resource-recovery-v9-review' and
                resource_review.get('valid') is True and resource_review.get('readOnly') is True and
                resource_review.get('importsJavaScript') is False and resource_review.get('solverRuns') == 0 and
                resource_review.get('inputSha256') == hashes[resource_paths[0]], 'Resource review is invalid or bound to different bytes')
        records = teaching.get('records', [])
        require(len(records) == resource_review.get('recordCount') == 4 and
                {record['fixture'] for record in records} == {'rescue', 'impossible'} and
                {record['input']['objectiveMode'] for record in records} == {'procurementOnly', 'feesPlusProcurement'} and
                len({(record['fixture'], record['input']['objectiveMode']) for record in records}) == 4 and
                all(len(record['input']['snapshot']['containers']) == 1 for record in records),
                'Resource teaching coverage differs from one container x two variants x two objectives')
        checked = {record['name']: record for record in resource_review.get('records', [])}
        require(set(checked) == {record['name'] for record in records} and
                all(checked[record['name']].get('valid') is True and
                    checked[record['name']].get('optimalityAudited') is True and
                    checked[record['name']].get('status') == record['result']['status'] for record in records),
                'Resource finite enumeration did not cover every teaching record')
        for name, expected in teaching.get('implementationHashes', {}).items():
            require(sha(relative_path(root, name)) == expected, 'Resource source commitment differs: ' + name)
        resource = {'teaching': teaching, 'review': resource_review}
    return {'root': root, 'wrapper': wrapper, 'index': index, 'protocol': protocol, 'summary': summary,
            'review': review, 'bootstrapReview': statistics, 'previousV8': previous, 'diagnosis': diagnosis, 'hashes': hashes,
            'resource': resource, 'materialOptions': {'bootstrap-review': bootstrap_review}}

def pair_label(pair):
    return LABELS[pair['candidateMethod']] + ' / ' + LABELS[pair['baselineMethod']]


def comparison_rows(panel):
    return [[pair_label(p), f"{p['fees']['total']['pairedScenarioCount']} / {p['fees']['total']['pairedEpisodeCount']}",
             pct(p['fees']['total']['weightedSavingsPct']), pct(p['fees']['incremental']['weightedSavingsPct']),
             '/'.join(str(p['fees']['total']['scenarioWinTieLoss'][k]) for k in ('win', 'tie', 'loss'))]
            for p in panel['comparisons'].values()]


def completion_text(summary, method):
    return '; '.join(f"{PANELS[panel]} {summary['panels'][panel]['byMethod'][method]['completeEpisodeCount']}/"
                     f"{summary['panels'][panel]['byMethod'][method]['recordedEpisodeCount']}" for panel in PANELS)


def conclusion(e):
    differences = [e['summary']['panels'][p]['byMethod']['adaptiveV9']['completeEpisodeCount'] -
                   e['summary']['panels'][p]['byMethod']['adaptiveV8']['completeEpisodeCount'] for p in PANELS]
    if all(value > 0 for value in differences):
        return '이 고정 평가에서 V9는 두 패널 모두 V8보다 완주 수가 많았습니다.'
    if all(value == 0 for value in differences):
        return '이 고정 평가에서 V9의 V8 대비 완주 수 개선은 관측되지 않았습니다.'
    if all(value < 0 for value in differences):
        return '이 고정 평가에서 V9는 두 패널 모두 V8보다 완주 수가 적었습니다.'
    return 'V9의 V8 대비 완주 결과는 패널별로 다릅니다. 패널을 합쳐 순위를 만들지 않습니다.'


def build_slides(e):
    p, s, r, previous = e['protocol'], e['summary'], e['review'], e['previousV8']['summary']
    n, repeats, methods = len(e['index']['cases']), len(p['algorithmRepeats']), len(p['methods'])
    episodes = r['summary']['episodes']
    old_observed = 'V8 ' + completion_text(previous, 'adaptiveV8') + ' / V7 ' + completion_text(previous, 'costPolishV7')
    observed = ('V9 ' + completion_text(s, 'adaptiveV9') + ' / V8 ' + completion_text(s, 'adaptiveV8') +
                ' / V7 ' + completion_text(s, 'costPolishV7'))
    slides = []
    def add(title, text, *, table=None, chart=None, flow=None, note='', english=''):
        slides.append({'title': title, 'text': text, 'table': table, 'chart': chart, 'flow': flow, 'note': note, 'english': english})
    add('오늘의 반출·적출·반납을\n실행 가능한 날짜로 연결합니다.',
        'PortPlan은 알려진 수입 컨테이너와 일별 운송·창고·체류 공간을 함께 계획하는 연구용 도구입니다. 이번 V9는 완전 후보를 제시간에 전달하는 문제를 검증합니다.',
        flow=['현재 정보·실행 잠금', '완전한 날짜 후보', '기한 내 검사·수락'],
        note=conclusion(e) + ' ' + observed + '. 합성 모델 결과이며 실제 기업 절감·우승 확률이 아닙니다.',
        english='PortPlan turns known import-container orders, published fees and daily shared resources into pickup–unload–return dates. V9 tests timely delivery of complete audited candidates. '+observed+'. This is a synthetic daily-model research result, not realized business savings or a first-place probability.')
    add('세 행동의 날짜와 요금을\n같은 자원 장부에서 계산합니다.',
        '반출·반납은 각각 운송 작업을 사용합니다. 적출은 연속 창고 일수를 차지하고, 반출부터 반납 전까지 체류 공간을 점유합니다.',
        table=(['행동/모형', '계산할 항목', '경계'], [
            ['반출 p / 적출 u / 반납 r', '방출·납기·선후행·실행 잠금', '일 단위 원자적 실행'],
            ['운송·연속 창고·체류 [p,r)', '날짜별 용량; container 또는 TEU', '도로 경로·슬롯·기사 근무시간 미포함'],
            ['PNC 보관료 / Maersk D&D', '공개 요율·기산일·누진/소급 가정', '개별 계약·청구서·세금 미확인']]),
        note='비용 목표는 선언된 두 요금입니다. 기존 주문은 시작 시 모두 알려지며 신규 주문 도착을 평가하지 않습니다.',
        english='Daily transport, consecutive unloading and half-open staging constraints share one ledger with the explicitly interpreted PNC and Maersk tariffs. Route, appointment, driver-hour, private-contract, invoice, tax and incoming-order claims are outside this model.')
    add('V8의 실제 후퇴를\n개선의 출발점으로 남깁니다.',
        '이전 최초 평가에서 V8은 V7보다 완주가 적었습니다. 내부 구성기가 후보를 찾았다는 사실만으로 외부 계획기가 기한 내 검사·수락했다고 볼 수 없었습니다.',
        table=(['이전 V8 평가 · 별도 입력', 'V8 완주', 'V7 완주'], [
            [PANELS[panel], f"{previous['panels'][panel]['byMethod']['adaptiveV8']['completeEpisodeCount']}/216",
             f"{previous['panels'][panel]['byMethod']['costPolishV7']['completeEpisodeCount']}/216"] for panel in PANELS]),
        note='이전 결과와 실패를 수정하지 않습니다. 예산 소진·늦은 인수·추가 검사 비용을 V9의 개발 동기로 사용하며, 새 성능 분모에 합치지 않습니다.',
        english='The previous V8 first run is retained unchanged: '+old_observed+'. Budget exhaustion between inner construction and outer acceptance motivates V9. An internal candidate timestamp is not evidence of timely receipt, copy, independent reconstruction and outer acceptance. Old scenarios are now development/regression data, not new V9 holdout evidence.')
    add('V9 가설: 먼저 받은 완전 후보를\n외부 검사 뒤 보존합니다.',
        '구성기의 동기식 onCandidate가 외부 gate에 후보를 전달합니다. 외부 gate는 복사·재가격·물리/잠금 검사·선택이 끝난 시각으로 수락을 결정합니다.',
        flow=['구성기 후보 → 실제 수신', '독립 날짜 복사·검사', '기한 내 수락만 보존'],
        note='공식 평가에서는 observer와 부모 후보 stream을 쓰지 않습니다. 부모 IPC 확인이나 내부 firstIncumbentMs를 일별 수락 증거로 바꾸지 않습니다. 가설의 지지는 실제 결과로 판단합니다.',
        english='V9 synchronously receives constructor candidates at an outer gate, makes an isolated copy, reconstructs prices and constraints, applies selection and accepts only by the original outer deadline. Formal observers and parent candidate streams are disabled. Internal first-incumbent timing and enqueue/IPC acknowledgements are not acceptance evidence.')
    add('유지·개선·복구의 조건은\n전달 방식과 구분합니다.',
        '현재 계획이 유효하면 유지하거나 더 저렴한 후보만 검토합니다. 실행한 날짜는 잠그고, 미실행 날짜 변경 수와 단기 명목 압력을 보호합니다.',
        table=(['조건', '선언된 수락 규칙'], [
            ['유효 cache', '기본 유지; 비용이 엄격히 낮은 후보만 개선 검토'],
            ['날짜 변경·비용 허용', '최대 6개 미실행 사건 변경; 0.5% 허용 경계'],
            ['단기 압력', '현재 명목 예측의 2일 대리 지표 비악화'],
            ['무효 cache / 초기 계획', '실행 잠금과 날짜 하한을 지킨 완전 복구 후보']]),
        note='현재 snapshot의 수락 조건입니다. 미래 완주 보장·전체 경로 비용 비증가·전역 최적을 뜻하지 않습니다.',
        english='V9 retains the adaptive date-change, fee-tolerance and near-term-pressure rules while changing timely candidate delivery. The six-event cap, 0.5% boundary and two-day nominal proxy are local heuristics. They do not guarantee future completion, full-path cost non-regression or a global optimum.')
    add('자원이 부족하면 명시한 메뉴에서\n구매 비용을 계산합니다.',
        '각 추가 자원 package에 가격·단위·날짜·가용일·허용 여부를 입력합니다. 승인된 현재/미래 package만 선택하여 고정 실행 이력을 지킨 복구안을 제안합니다.',
        table=(['선택 목표', '증명할 수 있는 범위'], [
            ['procurementOnly', '명시된 유한 메뉴 내 추가 구매비 최소'],
            ['feesPlusProcurement', '현재 모델의 추가 요금 + 구매비 최소'],
            ['feasible / time limit', '검사한 후보; 최소 비용이 증명된 것은 아님'],
            ['optimal / infeasible proof', '해당 모델·허용 메뉴에 한정; 실제 공급 보장 아님']]),
        note='모형 권고이며 구매·배차를 실행하지 않습니다. 메뉴 밖 대안, 미래 미공지 충격, 실제 견적·조달 가능성·현장 수익은 검증하지 않았습니다.',
        english='The finite-menu recovery model accepts explicit prices, units, days, earliest availability and authorization. It proposes approved packages under exact past locks. Minimum procurement or combined fee-plus-procurement claims require the corresponding proof and remain local to the declared menu/model. No purchase, dispatch, supplier availability or field ROI is validated.')
    add('분모와 예산은\n새 입력을 만들기 전에 고정했습니다.',
        f"{n}개 업무 시나리오 × {methods}개 정책 × {repeats}개 반복 = {episodes:,}개 episode. 두 패널 모두 V8의 생성기·충격 유형을 유지하고 새 주문 시드를 사용했습니다.",
        table=(['패널', '업무 시나리오', '평가 범위'], [
            ['동일 생성분포', 72, '기존 생성기·충격 유형; 새 독립 입력 시드'],
            ['스트레스', 72, 'V8과 같은 반복 감소·회복 연장/점진 회복·연속 자원 충격'],
            ['반복/예산', '3회 / 냉 5,000ms', '반복은 새 업무 표본 아님; 초기화·모든 결정·IPC·최종 검사 포함']]),
        note='정책은 직렬 실행하고 순서를 순환합니다. 일별 내부 시계가 완전히 같거나 실제 운영 실시간 보장을 한다고 주장하지 않습니다.',
        english=f'{n} new business scenarios, {methods} fixed methods and {repeats} retained algorithm/timing repeats yield {episodes} episodes. Both panels reuse the V8 generator and shock-pattern catalog with new order seeds; the stress mechanisms are not new. Both panels stay separate. The cold parent 5000ms includes startup, every decision, cleanup, terminal audits and receipt. Native startup is charged. The three repeats are not three independent business scenarios.')
    for panel in PANELS:
        ps = s['panels'][panel]
        add(PANELS[panel] + '\n완주와 이력 미확보를 먼저 확인합니다.',
            '완주는 물리 행동과 최종 검사까지 부모 기한 내 수락한 경우입니다. 이력을 확보하지 못한 episode를 물리 불가능이나 0원으로 바꾸지 않습니다.',
            chart=[{'label': LABELS[m], 'value': ps['byMethod'][m]['completeEpisodeCount'],
                    'denominator': ps['byMethod'][m]['recordedEpisodeCount'],
                    'unknown': ps['byMethod'][m]['executionHistoryUnavailableCount'],
                    'all3': ps['byMethod'][m]['fullyCompletedScenarioCount']} for m in METHODS],
            note='막대는 episode 완주 수입니다. 3회 모두 완주한 시나리오와 미확보 이력을 별도로 표시합니다.',
            english='Completion is reported before fees. Bars count complete episodes; all-three-repeat business scenarios and unavailable execution histories are separate. Missing terminal history is unknown, not evidence of physical infeasibility or zero cost.')
        add(PANELS[panel] + '\n공동 완주한 동일 반복만 요금을 비교합니다.',
            '양수는 후보가 더 저렴합니다. 시나리오마다 일치하는 완주 반복을 먼저 평균하고 시나리오를 동일 가중합니다.',
            table=(['후보 / 기준', '시나리오 / 쌍', '총요금 절감', '추가요금 절감', '낮음/동일/높음'], comparison_rows(ps)),
            note='비교마다 선택된 분모가 다릅니다. 실패 비용은 null입니다. 원화 분모·모든 악화 ID·조건부 구간은 보고서에 공개합니다.',
            english='Savings are baseline minus candidate. First average matched jointly completed repeats inside each business scenario, then weight scenarios equally. Pair-specific selected denominators differ; no ranking across unequal fee subsets is supported. Incomplete episode fees remain null.')
    losses=[]
    for panel, ps in s['panels'].items():
        for name, pair in ps['comparisons'].items():
            fee=pair['fees']['total']
            losses.append([panel, name, len(fee['negativeScenarioIds']), fnum(fee['worstScenarioLossKRW']),
                           f"{pair['completion']['gainedEpisodeCount']} / {pair['completion']['lostEpisodeCount']}"])
    add('완주 손실과 비용 악화를\n모든 비교에서 보존합니다.',
        '완주 증감은 episode 기준이며, 최악 요금 악화는 공동 완주 시나리오 평균 기준입니다. 유리한 분모만 골라 결론을 바꾸지 않습니다.',
        table=(['패널', '비교', '총요금 악화 시나리오', '최악 증가 ₩', '완주 얻음/잃음'], losses),
        note='총/추가 요금의 모든 악화 시나리오·반복 ID와 완주 손실 ID를 별도 보고서에 그대로 남깁니다.',
        english='Every declared comparison preserves gained/lost completions, fee losses and the worst scenario-average increase. All adverse total/incremental fee scenario/repeat IDs and lost-completion IDs remain in the report; completion and conditional fees never share an invented denominator.')
    timing=[]
    for panel, ps in s['panels'].items():
        for method in METHODS:
            m=ps['byMethod'][method]
            timing.append([panel, LABELS[method], fnum(m['wallTimeMs']['mean'],1), m['hardDeadlineCount'],
                           f"{m['revisionCount']['observedSum']} / {m['revisionCount']['unknownEpisodeCount']}"])
    add('더 많은 완주에도 계산과\n날짜 변경의 대가가 있습니다.',
        '날짜 변경은 미실행 사건 날짜의 변경 횟수입니다. 서로 다른 상자 수·작업자 시간·재배차 비용으로 바꾸어 해석하지 않습니다.',
        table=(['패널', '정책', '평균 부모 ms', '강제 기한', '관측 변경 / 미확보 episode'], timing),
        note='native HiGHS는 같은 모델의 서버 기준이며 웹의 GLPK와 구분합니다. native 시작 시간도 전체 부모 예산에 포함합니다.',
        english='Date revisions are edits of pending event dates, not distinct containers, worker time or dispatch cost. Unknown counters remain unknown. Native HiGHS is a same-model server baseline, separate from browser GLPK; its startup is included in the cold parent budget.')
    add('후보를 찾은 시각과\n검사·수락한 시각을 구분합니다.',
        '각 V9 일별 후보의 실제 수신·복사 완료·감사 완료·선택 완료·수락 시각을 전체 receipt sequence와 날짜로 보존합니다. 늦은 후보와 거부 이유도 남깁니다.',
        table=(['보존하는 증거', '확인할 경계'], [
            ['candidateDelivery.events 전체', '수신·복사·감사·선택의 순서와 예산 경계'],
            ['선택된 sequence와 날짜', '독립 재가격·자원·물리 잠금 재구성'],
            ['늦음·무효·guard 거부', '성공 후보만 선별하지 않음'],
            ['전체 episode receipt 없음', '과거 이력·비용 미확보; 후보 부재 증명 아님']]),
        note='별도 Python 검산은 기록된 시각의 일관성과 날짜를 재구성합니다. 외부 계측 인증이나 절대 CPU/실시간 보장이 아닙니다.',
        english='The untruncated V9 candidateDelivery receipt sequence preserves arrival, isolated-copy, audit, selection and acceptance endpoints with dates and rejections. The Python audit checks recorded timestamp consistency and reconstructs dates and constraints. It is not external timing certification; missing whole-episode receipts cannot prove that no inner candidate existed.')
    rc=r['summary']
    add('물리와 통계는\n서로 다른 Python 검산으로 확인합니다.',
        f"독립 장부가 {rc['episodes']:,}개 episode, 수락 계획 {rc['acceptedPlans']:,}개, 기록된 물리 행동 {rc['physicalEvents']:,}개를 재구성했습니다. 구간은 별도 통계 검산에 연결합니다.",
        table=(['검산', '범위와 한계'], [
            ['물리·요율·잠금·V9 gate', 'JS 행렬 미재사용; 내부 독립 구현'],
            ['scenario-cluster 95% bootstrap', '시나리오 단위 재표집; 별도 Python 재계산'],
            ['공동 완주 조건', '비용 구간의 완주 선택 편향은 남음'],
            ['현장·전역 최적·외부 인증', '이 검산으로 확인하지 않음']]),
        note='두 검산은 같은 formal index와 frozen manifest SHA에 묶습니다. 432개 반복을 432개 독립 시나리오로 세지 않습니다.',
        english='Separate read-only Python physics and statistical reviews bind to the same formal index and frozen manifest. They independently reconstruct causal inputs, locks, resources, tariffs, adaptive acceptance and scenario-cluster intervals. Completed-pair selection bias remains. Internal independent implementation is not third-party certification or operational validity.')
    add('화면에서 내 입력과\n명시적 보완 자원을 검토합니다.',
        '주문·모의 용량 사건을 입력하고 두 정책을 순차 비교합니다. 날짜별 정보·실행 잠금·권고·중단 이유를 확인하고 정확한 입력과 결과를 내보냅니다.',
        flow=['입력·가정 확인', '계획·실패·메뉴 권고', '입력·장부 내보내기'],
        note='작은 성공/중단 예시는 합성 학습용입니다. 공식 144개 연구 입력이나 실제 현장 실적이 아닙니다. 보완 자원 제안은 구매가 아닙니다.',
        english='Custom inputs and simulator-owned events support two sequential worker comparisons, daily known information, locked actions, failure inspection and exact export. Small illustrative success/failure fixtures are teaching examples, not formal business scenarios or field evidence. Resource recommendations do not purchase or dispatch anything.')
    add('개인 기여와 현실 검증을\n대신 완료로 표시하지 않습니다.',
        '이번 증거는 알려진 합성 주문의 일별 완주·요금·안정성에 한정됩니다. 기존 PSPLIB와 벌크항만 로그의 부분 검증을 실제 컨테이너 절감으로 합치지 않습니다.',
        table=(['주체/주장', '현재 상태'], [
            ['구현·실험·검산·자료 작성', 'AI 수행; 원 연구·코드·라이선스 귀속 유지'],
            ['사용자 본인 개발·이해·발표 숙련', 'pending · 대신 인정하지 않음'],
            ['전문가·기업·실제 청구·실현 절감', '미수행'],
            ['참가 자격·등록·서약·AI 규정', '본 자료로 확인하거나 제출하지 않음']]),
        note='최초 알고리즘·SOTA·모든 경우 우월·현실 ROI·우승 확률을 주장하지 않습니다.',
        english='AI performed the implementation, experiments, internal audits and preparation. Entrant authorship, understanding, defense readiness, competition eligibility, registration, declarations and applicable AI rules remain unverified. Public partial-model datasets do not establish container field savings; no first-place probability, SOTA or universal superiority is supported.')
    add('같은 결정을 다시 계산할\n전체 증거를 전달합니다.',
        '새 입력·모든 반복·실패·기한 손실·완전한 후보 receipt·고정 source·독립 검산을 SHA로 연결합니다. 재현은 새 경로에 기록합니다.',
        table=(['자료', '확인할 내용'], [
            ['한국어 보고서', '정확한 원화·분모·구간·모든 악화/손실 ID'],
            ['동일 원고의 PDF / 편집 가능 PPTX', '18쪽 이하; 표·문자·도형이 편집 가능'],
            ['영어 이해 노트 / 기여 기록', '가정·결과·한계·미확인 본인 작업'],
            ['V9 formal archive + 두 검산', '원 입력·모든 episode·고정 source의 해시 연결']]),
        note=conclusion(e) + ' 다음 판단도 완주·조건부 요금·변경 대가를 함께 사용합니다. 자료 준비는 대회 제출이나 서명이 아닙니다.',
        english='Deliver the exact input, every repeat, every adverse outcome, full candidate receipts, frozen source, independent physics review and independent interval reproduction. PDF and editable PPTX come from the same slide objects. New reproductions write new paths. Prepared materials are not a competition submission, signature or entrant-readiness attestation.')
    require(len(slides) == 18, 'Main deck must contain exactly 18 pages')
    return slides

def markdown_table(headers, rows):
    cell=lambda value:str(value).replace('|', '\\|').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(map(cell,headers))+' |','| '+' | '.join(['---']*len(headers))+' |',
                      *['| '+' | '.join(map(cell,row))+' |' for row in rows]])


def material_texts(e, slides):
    provenance='AI 작성 준비 자료. 참가자 본인의 창작·검증·이해·참가자격 확인서가 아닙니다.'
    report=['# PortPlan V9 · 적시 후보 전달과 오늘의 반출·적출·반납', '', provenance, '',
            '완결 공식 결과, 독립 물리 검산, 별도 bootstrap 검산이 동일 원본 SHA-256에 연결된 뒤 작성했습니다. 알려진 합성 주문의 일 단위 모형 연구이며 실제 기업 실현 절감액이 아닙니다.',
            '', conclusion(e), '']
    english=['# PortPlan V9 — understanding and defense notes', '',
             'AI-created preparation from complete formal evidence and two separately bound read-only reviews. Not evidence of entrant understanding, authorship, registration, competition eligibility or field acceptance.', '']
    for i, slide in enumerate(slides,1):
        report.extend([f"## {i:02d}. {slide['title'].replace(chr(10),' ')}", '',slide['text'],''])
        if slide['flow']:report.extend([' → '.join(slide['flow']).replace('\n',' / '),''])
        if slide['table']:report.extend([markdown_table(*slide['table']),''])
        if slide['chart']:report.extend([markdown_table(['정책','완주 episode / 전체','세 반복 모두 완주 시나리오','미확보 이력'],
            [[c['label'],f"{c['value']} / {c['denominator']}",c['all3'],c['unknown']] for c in slide['chart']]),''])
        report.extend([slide['note'],''])
        english.extend([f"## {i:02d}. {slide['title'].replace(chr(10),' ')}",'',slide['english'],''])
    report.extend(['## 정확한 요금·완주 분모·조건부 구간·모든 악화 목록','',
        '절감 = 기준 − 후보. 각 시나리오에서 일치한 공동 완주 반복만 평균한 뒤 시나리오를 동일 가중합니다. 총요금과 추가요금의 원화 분모가 다릅니다. 0원 분모의 비율은 null입니다.',
        '모든 95% 구간은 percentile scenario-cluster bootstrap이며 별도 Python 통계 검산에서 동일 원본으로 재계산했습니다. 비용 구간은 선택된 공동 완주 쌍에 조건부입니다. 검산이 완주 선택 편향을 제거하거나 관측 밖 운영 집단으로의 외삽을 정당화하지 않습니다.',
        '날짜 변경의 원가를 관측하지 않았습니다. 날짜 변경 수를 비용에 합산하거나 작업자 시간·재배차비로 환산하지 않습니다.',''])
    for panel, ps in e['summary']['panels'].items():
        report.extend(['### '+PANELS[panel],''])
        report.extend([markdown_table(['정책','전체 episode','완주','3회 모두 완주 시나리오','1회 이상 완주 시나리오','이력 미확보'],
            [[LABELS[m],ps['byMethod'][m]['recordedEpisodeCount'],ps['byMethod'][m]['completeEpisodeCount'],
              ps['byMethod'][m]['fullyCompletedScenarioCount'],ps['byMethod'][m]['anyCompletedScenarioCount'],
              ps['byMethod'][m]['executionHistoryUnavailableCount']] for m in METHODS]),''])
        for name,pair in ps['comparisons'].items():
            report.extend(['#### '+name+' · '+pair_label(pair),''])
            headers=['요금','쌍 시나리오/반복','기준 평균합 ₩','후보 평균합 ₩','절감 ₩','절감 %','낮음/동일/높음','최악 평균 증가 ₩','0원 기준 시나리오/반복','독립 검산 95% 구간 %']
            rows=[]
            for key in ('total','incremental'):
                f=pair['fees'][key]
                rows.append([key,f"{f['pairedScenarioCount']} / {f['pairedEpisodeCount']}",fnum(f['baselineScenarioMeanSumKRW'],2),
                    fnum(f['candidateScenarioMeanSumKRW'],2),fnum(f['scenarioWeightedSavingsKRW'],2),pct(f['weightedSavingsPct']),
                    '/'.join(str(f['scenarioWinTieLoss'][k]) for k in ('win','tie','loss')),fnum(f['worstScenarioLossKRW'],2),
                    f"{f['zeroBaselineScenarioCount']} / {f['zeroBaselineEpisodeCount']}",ci(f.get('uncertainty',{}).get('weightedSavingsPct95CI'))])
            report.extend([markdown_table(headers,rows),''])
            c=pair['completion']; rev=pair['revisions']
            report.extend([f"완주 얻음/잃음 episode: {c['gainedEpisodeCount']} / {c['lostEpisodeCount']}. "
                f"시나리오 평균 완주율 차이: {fnum(c['meanScenarioRateGain'],6)}; 독립 검산 95% 구간: {ci(c.get('uncertainty',{}).get('meanScenarioRateGain95CI'))}.",
                f"공동 완주 반복의 관측 날짜 변경 합: {rev['baselineObservedSum']} → {rev['candidateObservedSum']}; {rev['pairedEpisodeCount']}쌍.",''])
            for key in ('total','incremental'):
                f=pair['fees'][key]; uncertainty=f.get('uncertainty',{})
                report.extend([f"{key} 원 반복 단위 비용합: 기준 {fnum(f['baselineEpisodeSumKRW'],2)} ₩ / 후보 {fnum(f['candidateEpisodeSumKRW'],2)} ₩. 이 합은 시나리오 동일가중 합과 다릅니다.",
                    f"{key} 시나리오 평균 절감 원화 95% 구간: {ci(uncertainty.get('meanSavings95CI'))}; "
                    f"bootstrap 시나리오 {uncertainty.get('scenarioCount')}개 / draw {uncertainty.get('draws')}회.",''])
                for field in ('negativeScenarioIds','negativeEpisodeIds'):
                    report.extend([f'**{key} / {field} — 전체 보존**',''])
                    report.extend(['- `'+item+'`' for item in f[field]] or ['없음.']);report.append('')
            for field in ('gainedEpisodeIds','lostEpisodeIds','gainedFromKnownFailureEpisodeIds','gainedFromUnknownHistoryEpisodeIds'):
                report.extend(['**'+field+' — 전체 보존**',''])
                report.extend(['- `'+item+'`' for item in c[field]] or ['없음.']);report.append('')
    report.extend(['## 이전 V8 평가 — 새 평가와 합치지 않음','',
        'V8 ' + completion_text(e['previousV8']['summary'],'adaptiveV8') + '.',
        'V7 ' + completion_text(e['previousV8']['summary'],'costPolishV7') + '.',
        '이전 최초 입력·반복·실패·비용 악화는 변경하지 않습니다. 이 자료는 V9 개발/회귀 검사에 사용할 수 있지만 새 독립 시나리오로 다시 세지 않습니다. 내부 구성 후보 시각, 외부 수신·복사·검사·수락 시각과 전체 냉 episode 수락은 서로 다른 경계입니다.',
        '', '## G0 전체 실패 층별 상태 — 기존 자료의 사후 진단',''])
    diag=e['diagnosis']
    def feasibility(case,layer):return diag['models'][case['layers'][layer]]['feasibility']
    report.extend([markdown_table(['기존 V7 실패 ID','실패일','원 명목','관측+기존 잠금','전지적 실제','전지적 실제+잠금'],
        [[c['id'],c['originalOnlineFailure']['day'],*[feasibility(c,k) for k in ('nominalFromStart','observedExecutedPrefix','actualHindsightFromStart','actualHindsightExecutedPrefix')]] for c in diag['cases']]),'',
        'G0는 독립 추가 오프라인 trip 모델 진단입니다. 전지적 해답이나 추가 오프라인 예산을 V9 온라인 입력·성공·신규 holdout 성과로 바꾸지 않습니다.',''])
    report.extend(['## 후보 전달 증거와 검산 범위','',
        '공식 일별 수락은 내부 동기식 constructor → outer gate입니다. 외부 observer·부모 candidate stream·enqueue acknowledgement는 공식 수락 규칙에 없습니다. 각 endpoint는 기록된 상대 시각의 일관성을 검사하며 외부 계측 인증으로 해석하지 않습니다.',
        '전체 candidateDelivery.events, 제안 날짜, canonical 날짜, 재구성 비용, 수락·거부 이유와 선택 sequence는 개별 episode 원본에 보존됩니다. 미확보 전체 이력은 실제 후보 부재를 증명하지 않습니다.',
        '', '### 물리·후보 검산 원 요약', '', '```json', json.dumps(e['review'].get('summary',{}),ensure_ascii=False,indent=2), '```',
        '', '### 통계 검산 절차', '', '```json', json.dumps(e['bootstrapReview'].get('procedure',{}),ensure_ascii=False,indent=2), '```',
        '', '## 명시적 자원 메뉴 권고 — 구매·현장 성과 아님','',
        '메뉴 항목은 resource, days, units, priceKRW, authorized, earliestAvailableDay를 명시합니다. 승인되지 않은 항목, 과거 용량을 바꾸는 항목, 가용일보다 이른 항목은 선택 대상에서 제외합니다. 이미 실행한 사건 날짜는 그대로 유지합니다.',
        '`procurementOnly`의 minimumProcurementProven과 `feesPlusProcurement`의 minimumObjectiveProven은 서로 다른 주장입니다. 검증된 feasible 후보라도 대응 증명이 없으면 최소라고 말하지 않습니다. infeasibleWithinMenuProven은 해당 관측·날짜 도메인·명시 메뉴의 불가능성에 한정됩니다.',
        '권고는 counterfactual이며 구매·배차를 실행하지 않습니다. 추가 요금과 구매비를 함께 표시하더라도 실제 견적·공급 확약·미래 불확실성·업무 도입 효과·현장 ROI를 확인한 것은 아닙니다. 이 기능의 작은 검증 예제를 공식 144개 전체 정책 비교의 분모에 넣지 않습니다.',''])
    if e.get('resource'):
        records = e['resource']['teaching']['records']
        procurement = next(record for record in records if record['fixture'] == 'rescue' and record['input']['objectiveMode'] == 'procurementOnly')
        combined = next(record for record in records if record['fixture'] == 'rescue' and record['input']['objectiveMode'] == 'feesPlusProcurement')
        p_result, c_result = procurement['result'], combined['result']
        p_return, c_return = p_result['rows'][0]['returnDay'], c_result['rows'][0]['returnDay']
        report.extend(['### 독립 전수 열거로 확인한 합성 교육 예제 4건','',
            '한 컨테이너 × 두 변형(rescue / impossible) × 두 목적함수입니다. 네 개 업무 시나리오나 실제 주문이 아닙니다. 실제 견적·구매 권한이 아니라 입력한 합성 가격과 합성 허용 값입니다.', '',
            markdown_table(['합성 record','요청 목적','상태','구매비 ₩','추가 요금 ₩','추가 요금 + 구매비 ₩','반납일'],
                [[record['name'],record['input']['objectiveMode'],record['result']['status'],
                  fnum(record['result'].get('procurementCostKRW')),fnum(record['result'].get('incrementalCost')),
                  fnum(record['result'].get('combinedIncrementalCostKRW')),
                  ', '.join('D'+str(row['returnDay']) for row in record['result'].get('rows',[])) or 'null'] for record in records]),'',
            f'구매비만 최소화하는 목적은 요금의 2차 최적화를 요구하지 않습니다. 이 실제 결과의 procurementOnly는 D{p_return} 반납을 반환하였고, feesPlusProcurement는 D{c_return}를 반환했습니다. 두 결과와 목적 차이를 그대로 보존합니다. 공동 목적의 최적값을 최소 구매비와 혼동하지 않습니다.', '',
            '독립 Python은 이 작은 날짜 도메인의 모든 적법 날짜와 전체 허용 package 부분집합을 열거했습니다. 이 예제의 최적성 확인을 더 큰 임의 입력이나 알 수 없는 미래 충격의 보장으로 확장하지 않습니다.',''])
        english.extend(['## Separately enumerated resource teaching records','',
            f"The four records are one synthetic container x two variants x two requested objectives, not four business scenarios or real orders. Minimum procurement is KRW {fnum(p_result['procurementCostKRW'])}; the combined incremental-fee-plus-procurement minimum is KRW {fnum(c_result['combinedIncrementalCostKRW'])}. These are different objectives. The actual procurement-only result returns on D{p_return}, with incremental fees KRW {fnum(p_result['incrementalCost'])}; the combined-objective result returns on D{c_return}, with incremental fees KRW {fnum(c_result['incrementalCost'])}. Both outcomes remain disclosed. Neither is a purchase, quote validation, future-robustness guarantee or realized saving.",''])
    hashes=markdown_table(['증거 입력','SHA-256'],list(e['hashes'].items()))
    material_command='python scripts/build-defense-v9.py --require-results'
    for option,value in e.get('materialOptions',{}).items():
        if value:material_command+=' --'+option+' '+shlex.quote(value)
    formal_directory=Path(e['wrapper']['formalIndex']['path']).parent.as_posix()
    report.extend(['## 원본 연결','',hashes,'',
        '복원된 formal index의 inputArtifact와 episode artifact는 원 SHA로 연결됩니다. wrapper는 전체 장부를 대체하지 않습니다. 자료 복원과 독립 검산은 새 실험 성과 생성이 아닙니다.',
        '', '## 재현','', '```bash',
        'python scripts/restore-evidence-v9.py',
        'python scripts/review-v9.py --run '+shlex.quote(formal_directory)+' --output artifacts/review-v9-reproduction.json',
        'python scripts/review-bootstrap-v9.py --run '+shlex.quote(formal_directory)+' --output artifacts/bootstrap-v9-reproduction.json',
        material_command,'```','',
        '복원 CLI의 옵션과 통계 검산 파일명은 저장된 각 script의 --help를 확인합니다. 최초 결과를 덮어쓰지 않습니다. 개인 참가·AI 조건·등록·서약·본인 발표 숙련은 별도 확인 대상입니다.'])
    english.extend(['## Exact evidence bindings','',hashes,'',
        'See submission-v9.ko.md for every pair-specific total/incremental denominator, all adverse scenario and repeat IDs, gained/lost completions, revision tradeoffs and conditional scenario-cluster intervals. Both independent reviews bind to the same formal index and frozen manifest. Completed-pair selection bias remains. No pooled panel leaderboard, full-path non-regression, universal recovery, SOTA, field ROI or first-place probability is supported.',
        '', 'The previous V8 experiment remains unchanged and is a separate development/regression source. Finite-menu minimum procurement requires the explicit corresponding proof; recommendations are not actual purchases, dispatches or supplier commitments.'])
    contribution=['# PortPlan V9 · 실제 기여 기록','',provenance,'',
        markdown_table(['작업','실제 수행 주체','범위'],[
            ['이전 V8 후퇴·G0 진단 보존','AI','과거 결과를 변경하지 않음; 새 holdout과 분리'],
            ['V9 후보 전달·잠금·guard 구현','AI','동기식 gate·전체 receipt·수락/거부 증거 보존'],
            ['명시적 가격·권한 메뉴 자원 보완 구현','AI','현재 모델의 반사실 권고; 실제 구매·배차 없음'],
            ['고정 새 평가·물리 검산','AI','완결 원본과 검산을 같은 SHA로 연결'],
            ['별도 scenario-cluster 통계 검산','AI','동일 원본 SHA와 frozen manifest의 구간 재계산'],
            ['한국어 HTML/PDF/PPTX·영어 노트','AI','동일 18개 slide 객체에서 문자·표·도형 편집 가능 자료 생성'],
            ['사용자 본인 개발·검증·이해·발표','pending','AI 작업을 개인 창작이나 이해 증명으로 옮기지 않음'],
            ['전문가·기업·청구서·현장 검증','미수행','인터뷰·실현 절감·현장 수락 증거 없음'],
            ['등록·서약·본인 발표·AI 규정 확인','pending','자료 준비이며 대회 제출·서명·자격 인증 아님']]),'',
        'GLPK, HiGHS, SciPy, 원 연구·데이터·공개 방법의 저작권과 notice는 각 소유자에게 귀속합니다. 내부 독립 구현은 외부 제3자 인증이 아닙니다. PDF는 Noto Sans CJK KR의 필요한 glyph를 변환·부분 내장하며 SIL Open Font License 귀속을 유지합니다.',
        '', '## 실제 자료 입력','',hashes]
    return '\n'.join(report)+'\n','\n'.join(english)+'\n','\n'.join(contribution)+'\n'

def html_document(slides):
    parts=['<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PortPlan V9</title><style>',
        'body{margin:0;background:#dce8e0;color:#173c40;font-family:"Noto Sans CJK KR","Noto Sans KR",sans-serif}.slide{box-sizing:border-box;width:1280px;min-height:720px;margin:24px auto;padding:44px 62px;background:#f7faf5;position:relative;page-break-after:always}h1{font-size:35px;line-height:1.4;margin:10px 0 22px;white-space:pre-line}.kicker{font-size:13px;letter-spacing:3px;color:#39826d}p{font-size:20px;line-height:1.7}table{width:100%;border-collapse:collapse;font-size:15px;margin:22px 0}th,td{text-align:left;border-bottom:1px solid #d3e0d3;padding:12px 10px;overflow-wrap:anywhere}th{background:#e6f0e4}.note{font-size:14px;background:#e8f2e5;padding:14px 18px;line-height:1.8}.flow{display:flex;gap:24px;align-items:stretch;margin:38px 0}.flow div{flex:1;background:#e0efe0;border-left:5px solid #36836c;padding:24px;font-size:23px;white-space:pre-line}.chart{margin:22px 0}.row{display:grid;grid-template-columns:240px 1fr 250px;align-items:center;gap:16px;margin:19px 0;font-size:17px}.track{height:21px;background:#dfebde}.fill{height:100%;background:#35826d}.footer{font-size:12px;letter-spacing:.5px;margin-top:25px;color:#638176}@media(max-width:1300px){.slide{width:100%;min-height:0;margin:12px auto;padding:30px}h1{font-size:29px}.row{grid-template-columns:180px 1fr 190px}}@media print{@page{size:1280px 720px;margin:0}body{background:white}.slide{width:1280px;height:720px;margin:0;overflow:hidden}}',
        '</style><body>']
    for i,slide in enumerate(slides,1):
        parts.extend([f'<section class="slide" id="slide-{i}"><div class="kicker">{i:02d} / PORTPLAN V9</div>',
                      '<h1>'+escape(slide['title'])+'</h1><p>'+escape(slide['text'])+'</p>'])
        if slide['flow']:parts.append('<div class="flow">'+''.join('<div>'+escape(text)+'</div>' for text in slide['flow'])+'</div>')
        if slide['table']:
            headers,rows=slide['table'];parts.append('<table><thead><tr>'+''.join('<th>'+escape(str(h))+'</th>' for h in headers)+'</tr></thead><tbody>')
            parts.extend('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows);parts.append('</tbody></table>')
        if slide['chart']:
            parts.append('<div class="chart">')
            for c in slide['chart']:
                width=100*c['value']/c['denominator'] if c['denominator'] else 0
                parts.append(f'<div class="row"><span>{escape(c["label"])}</span><div class="track"><div class="fill" style="width:{width:.4f}%"></div></div><span>{c["value"]}/{c["denominator"]} · 3/3 {c["all3"]} · 미확보 {c["unknown"]}</span></div>')
            parts.append('</div>')
        parts.extend(['<div class="note">'+escape(slide['note'])+'</div>',f'<div class="footer">AI 작성 준비 자료 · 본인 검토·현장 검증 미확인 / {i:02d} / {len(slides):02d}</div></section>'])
    parts.append('</body></html>');return '\n'.join(parts)+'\n'


def register_noto(text):
    """Embed Noto KR; convert only the used CFF glyphs to quadratic outlines in memory."""
    from fontTools import subset
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.cu2quPen import Cu2QuPen
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as PDFont
    path=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    require(path.is_file(), 'Noto Sans CJK font is required for a readable Korean PDF')
    font=TTFont(path,fontNumber=1)
    require('KR' in font['name'].getDebugName(1), 'Expected Noto Korean font face')
    chars=set(ord(c) for c in text if not c.isspace());cmap=font.getBestCmap()
    require(not (chars-set(cmap)), 'Noto is missing required PDF glyphs: '+repr(chars-set(cmap)))
    subsetter=subset.Subsetter();subsetter.populate(text=text);subsetter.subset(font)
    order=font.getGlyphOrder();glyphset=font.getGlyphSet();glyf={}
    for name in order:
        pen=TTGlyphPen(glyphset);glyphset[name].draw(Cu2QuPen(pen,1.0,reverse_direction=True));glyf[name]=pen.glyph()
    builder=FontBuilder(font['head'].unitsPerEm,isTTF=True);builder.setupGlyphOrder(order);builder.setupCharacterMap(font.getBestCmap());builder.setupGlyf(glyf)
    builder.setupHorizontalMetrics(font['hmtx'].metrics);builder.setupHorizontalHeader(ascent=font['hhea'].ascent,descent=font['hhea'].descent)
    builder.setupNameTable({'familyName':'Noto Sans CJK KR PortPlan subset','styleName':'Regular','uniqueFontIdentifier':'PortPlanNotoKR','fullName':'PortPlan Noto KR subset','psName':'PortPlanNotoKR',
                          'copyright':'Noto Sans CJK. Copyright Adobe and Google. SIL Open Font License 1.1. Subset outlines converted for embedding.'})
    builder.setupOS2(sTypoAscender=font['hhea'].ascent,sTypoDescender=font['hhea'].descent,usWinAscent=font['hhea'].ascent,usWinDescent=-font['hhea'].descent)
    builder.setupPost();builder.setupMaxp();buffer=BytesIO();builder.save(buffer);buffer.seek(0);pdfmetrics.registerFont(PDFont('PortPlanNoto',buffer))


def pdf_document(slides):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle
    all_text=json.dumps(slides,ensure_ascii=False)+' PORTPLAN V9 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789'
    register_noto(all_text)
    buffer=BytesIO();c=canvas.Canvas(buffer,pagesize=(1280,720),invariant=1,pageCompression=1)
    c.setTitle('PortPlan V9 · Korean competition preparation');c.setAuthor('AI-assisted PortPlan research; entrant verification pending')
    bg=HexColor('#f7faf5');ink=HexColor('#173c40');green=HexColor('#35826d');light=HexColor('#e6f0e4')
    def paragraph(text,x,y,w,size=20,leading=None):
        style=ParagraphStyle('p',fontName='PortPlanNoto',fontSize=size,leading=leading or size*1.55,textColor=ink,wordWrap='CJK')
        obj=Paragraph(escape(str(text)).replace('\n','<br/>'),style);_,h=obj.wrap(w,700)
        obj.drawOn(c,x,y-h);return y-h
    for index,slide in enumerate(slides,1):
        c.setFillColor(bg);c.rect(0,0,1280,720,stroke=0,fill=1)
        c.setFillColor(green);c.setFont('PortPlanNoto',13);c.drawString(62,672,f'{index:02d} / PORTPLAN V9')
        y=paragraph(slide['title'],62,642,1156,33,43)-15
        y=paragraph(slide['text'],62,y,1156,18,27)-19
        if slide['flow']:
            w=(1156-44)/3;h=118
            for j,text in enumerate(slide['flow']):
                x=62+j*(w+22);c.setFillColor(light);c.roundRect(x,y-h,w,h,7,stroke=0,fill=1);paragraph(text,x+18,y-22,w-36,23,34)
                if j<2:c.setStrokeColor(green);c.setLineWidth(2);c.line(x+w+4,y-h/2,x+w+18,y-h/2);c.line(x+w+13,y-h/2+5,x+w+18,y-h/2);c.line(x+w+13,y-h/2-5,x+w+18,y-h/2)
            y-=h+24
        if slide['chart']:
            for row in slide['chart']:
                paragraph(row['label'],62,y-5,252,17,23)
                c.setFillColor(light);c.rect(322,y-24,580,20,stroke=0,fill=1)
                c.setFillColor(green);c.rect(322,y-24,580*row['value']/row['denominator'] if row['denominator'] else 0,20,stroke=0,fill=1)
                paragraph(f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y-4,294,15,22);y-=53
            y-=8
        if slide['table']:
            headers,rows=slide['table'];dense=len(rows)>7;font_size=11 if dense else 15
            style=ParagraphStyle('cell',fontName='PortPlanNoto',fontSize=font_size,leading=font_size*1.4,textColor=ink,wordWrap='CJK')
            data=[[Paragraph(escape(str(v)),style) for v in row] for row in [headers,*rows]]
            widths=[1156/len(headers)]*len(headers)
            if len(headers)==2:widths=[550,606]
            table=Table(data,colWidths=widths,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),light),('VALIGN',(0,0),(-1,-1),'TOP'),
                ('BOTTOMPADDING',(0,0),(-1,-1),5 if dense else 9),('TOPPADDING',(0,0),(-1,-1),5 if dense else 9),('LEFTPADDING',(0,0),(-1,-1),10),
                ('LINEBELOW',(0,0),(-1,-1),0.5,HexColor('#d3e0d3'))]))
            _,h=table.wrap(1156,600);require(y-h>100,f'PDF table overflows page {index}: {y-h:.1f}');table.drawOn(c,62,y-h);y-=h+18
        note_style=ParagraphStyle('note',fontName='PortPlanNoto',fontSize=12,leading=19,textColor=ink,wordWrap='CJK')
        note=Paragraph(escape(slide['note']),note_style);_,h=note.wrap(1116,300)
        require(y-h-28>64,f'PDF note overflows page {index}: {y-h-28:.1f}')
        c.setFillColor(light);c.roundRect(62,y-h-28,1156,h+25,5,stroke=0,fill=1);note.drawOn(c,82,y-h-14)
        c.setFillColor(HexColor('#638176'));c.setFont('PortPlanNoto',10)
        c.drawString(62,32,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인');c.drawRightString(1218,32,f'PORTPLAN V9 / {index:02d} / {len(slides):02d}');c.showPage()
    c.save();return buffer.getvalue()


def pptx_document(slides):
    """Editable native text, tables and vector shapes; no raster page pictures."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from reportlab.pdfbase import pdfmetrics

    require(len(slides)==18, 'Editable deck must contain exactly 18 slides')
    # pdf_document normally registers the same outline metrics first; direct
    # renderer use remains possible without rendering a PDF or writing a font.
    if 'PortPlanNoto' not in pdfmetrics.getRegisteredFontNames():
        register_noto(json.dumps(slides,ensure_ascii=False)+' PORTPLAN V9 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789')
    presentation=Presentation();presentation.slide_width=Inches(1280/96);presentation.slide_height=Inches(720/96)
    presentation.core_properties.title='PortPlan V9 · Korean competition preparation'
    presentation.core_properties.author='AI-assisted PortPlan research; entrant verification pending'
    presentation.core_properties.subject='Editable result-driven research deck; synthetic daily model'
    colors={'background':'F7FAF5','ink':'173C40','green':'35826D','light':'E6F0E4','muted':'638176'}
    unit=lambda px: Inches(px/96)

    def wrapped(value,width,size):
        lines=[]
        for raw in str(value).split('\n'):
            line='';used=0
            for char in raw:
                advance=pdfmetrics.stringWidth(char,'PortPlanNoto',size)
                if line and used+advance>width:
                    lines.append(line.rstrip());line='';used=0
                line+=char;used+=advance
            lines.append(line.rstrip())
        return '\n'.join(lines),len(lines)

    def style_text(frame,text,size,leading,color='ink'):
        frame.clear();frame.text=text;frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        frame.vertical_anchor=MSO_ANCHOR.TOP
        for paragraph in frame.paragraphs:
            paragraph.font.name='Noto Sans CJK KR';paragraph.font.size=Pt(size*.75)
            paragraph.font.color.rgb=RGBColor.from_string(colors[color])
            paragraph.line_spacing=Pt(leading*.75);paragraph.space_before=Pt(0);paragraph.space_after=Pt(0)
            properties=paragraph._p.get_or_add_pPr().get_or_add_defRPr()
            east_asian=properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
            if east_asian is None:
                east_asian=OxmlElement('a:ea');properties.append(east_asian)
            east_asian.set('typeface','Noto Sans CJK KR')

    def text(slide,value,x,y,width,size=18,leading=27,color='ink',height=None,name='Editable text'):
        content,line_count=wrapped(value,width-2,size)
        measured=line_count*leading+3
        box=slide.shapes.add_textbox(unit(x),unit(y),unit(width),unit(height or measured));box.name=name
        style_text(box.text_frame,content,size,leading,color)
        return y+(height or measured)

    def rect(slide,x,y,width,height,color='light',kind=MSO_SHAPE.RECTANGLE,name='Editable vector'):
        shape=slide.shapes.add_shape(kind,unit(x),unit(y),unit(width),unit(height));shape.name=name
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(colors[color]);shape.line.fill.background()
        return shape

    for index,source in enumerate(slides,1):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string(colors['background'])
        text(slide,f'{index:02d} / PORTPLAN V9',62,31,1156,13,20,'green',name='Slide index')
        y=text(slide,source['title'],62,72,1156,33,43,name='Slide title')+15
        y=text(slide,source['text'],62,y,1156,18,27,name='Main explanation')+19
        if source['flow']:
            width=(1156-44)/3
            for j,label in enumerate(source['flow']):
                x=62+j*(width+22);rect(slide,x,y,width,118,name='Flow node '+str(j+1))
                rect(slide,x,y,5,118,'green',name='Flow node accent')
                text(slide,label,x+18,y+22,width-36,23,34,height=80,name='Editable flow label')
                if j<2:rect(slide,x+width+4,y+52,14,14,'green',MSO_SHAPE.CHEVRON,'Editable flow arrow')
            y+=142
        if source['chart']:
            for row in source['chart']:
                text(slide,row['label'],62,y+5,252,17,23,name='Chart method label')
                rect(slide,322,y+4,580,20,name='Editable completion bar track')
                width=580*row['value']/row['denominator'] if row['denominator'] else 0
                if width:rect(slide,322,y+4,width,20,'green',name='Editable completion bar')
                text(slide,f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y+4,294,15,22,name='Chart exact value')
                y+=53
            y+=8
        if source['table']:
            headers,rows=source['table'];dense=len(rows)>7;size=11 if dense else 15;leading=size*1.4;padding=5 if dense else 9
            widths=[1156/len(headers)]*len(headers)
            if len(headers)==2:widths=[550,606]
            formatted=[];heights=[]
            for row in [headers,*rows]:
                cells=[wrapped(value,widths[j]-20,size) for j,value in enumerate(row)]
                formatted.append([value for value,_ in cells]);heights.append(max(count for _,count in cells)*leading+2*padding+3)
            height=sum(heights);require(y+height<620,f'Editable table overflows slide {index}')
            native=slide.shapes.add_table(len(formatted),len(headers),unit(62),unit(y),unit(1156),unit(height));native.name='Editable evidence table'
            table=native.table;table.first_row=True
            for j,width in enumerate(widths):table.columns[j].width=unit(width)
            for i,row in enumerate(formatted):
                table.rows[i].height=unit(heights[i])
                for j,value in enumerate(row):
                    cell=table.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(colors['light' if i==0 else 'background'])
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
                    style_text(cell.text_frame,value,size,leading)
                    # Keep table padding after generic text-frame styling.
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
            y+=height+18
        note,note_lines=wrapped(source['note'],1116,12);note_height=note_lines*19+3
        require(y+note_height+28<656,f'Editable note overflows slide {index}')
        rect(slide,62,y,1156,note_height+25,name='Editable boundary note background')
        text(slide,note,82,y+11,1116,12,19,height=note_height+3,name='Evidence boundary note')
        text(slide,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인',62,678,900,10,15,'muted',name='Contribution footer')
        text(slide,f'PORTPLAN V9 / {index:02d} / {len(slides):02d}',1040,678,178,10,15,'muted',name='Page footer')
        for shape in slide.shapes:
            require(shape.shape_type!=MSO_SHAPE_TYPE.PICTURE,'Raster slide pictures are prohibited in the editable deck')
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=presentation.slide_width+1 and
                    shape.top+shape.height<=presentation.slide_height+1,f'Editable shape escapes slide {index}: {shape.name}')
    require(len(presentation.slides)==18,'Editable deck page count changed')
    output=BytesIO();presentation.save(output);return output.getvalue()


def review_generated(evidence, slides):
    """Inspect saved artifacts; do not solve, replay, or modify formal evidence."""
    from datetime import datetime
    from zoneinfo import ZoneInfo
    import unicodedata
    import fitz
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    normalize = lambda value: ''.join(unicodedata.normalize('NFKC', str(value)).split())
    paths = [ROOT / 'docs' / name for name in OUTPUTS]
    require(all(path.is_file() for path in paths), 'Missing generated material')
    document = fitz.open(paths[1])
    presentation = Presentation(paths[5])
    require(len(document) == len(presentation.slides) == len(slides) == 18, 'PDF/PPTX/source page counts differ')
    review = {'schemaVersion': 'portplan-v9-material-review', 'valid': True,
              'reviewedAtKST': datetime.now(ZoneInfo('Asia/Seoul')).isoformat(timespec='seconds'),
              'readOnlyFormalEvidence': True, 'solverRuns': 0, 'newScenarios': 0,
              'formalIndexSha256': evidence['wrapper']['formalIndex']['sha256'],
              'evidenceHashes': evidence['hashes'],
              'generatorSha256': sha(Path(__file__)),
              'slideObjectSha256': hashlib.sha256(json.dumps(slides, ensure_ascii=False, sort_keys=True,
                                                           separators=(',', ':')).encode()).hexdigest(),
              'pdfPageCount': len(document), 'pptxSlideCount': len(presentation.slides),
              'sameSourceSlideObjects': True, 'sourceTextCoverage': True,
              'pdfImages': 0, 'pptxPictures': 0, 'pptxEditableTextShapes': 0, 'pptxEditableTables': 0,
              'pages': [],
              'scope': 'Saved-artifact text, exact native table cells, vector/editable objects, page counts, PDF text bounds and PPTX shape bounds; renderer rejects table/note overflow before writing. PDF/PPTX share the same source slide objects and bound evidence. No PowerPoint application rendering, browser interaction, external clock certification, field validation or entrant understanding is certified.'}
    report_text = paths[2].read_text()
    expected_ids = Counter()
    adverse_id_count = 0
    completion_loss_id_count = 0
    for panel in evidence['summary']['panels'].values():
        for pair in panel['comparisons'].values():
            for key in ('total', 'incremental'):
                for field in ('negativeScenarioIds', 'negativeEpisodeIds'):
                    expected_ids.update(pair['fees'][key][field])
                    adverse_id_count += len(pair['fees'][key][field])
                require(ci(pair['fees'][key]['uncertainty']['weightedSavingsPct95CI']) in report_text,
                        'Conditional fee interval is absent from Korean report')
            for field in ('gainedEpisodeIds', 'lostEpisodeIds', 'gainedFromKnownFailureEpisodeIds', 'gainedFromUnknownHistoryEpisodeIds'):
                expected_ids.update(pair['completion'][field])
            completion_loss_id_count += len(pair['completion']['lostEpisodeIds'])
            require(ci(pair['completion']['uncertainty']['meanScenarioRateGain95CI']) in report_text,
                    'Completion interval is absent from Korean report')
    for identifier, count in expected_ids.items():
        require(report_text.count('- `'+identifier+'`') == count,
                'Adverse/completion identifier multiplicity differs in Korean report: ' + identifier)
    require('스트레스 · 새 주문 시드' in report_text and '새 충격 스트레스' not in report_text,
            'Stress mechanism wording incorrectly implies a new shock catalog')
    review['koreanReportAdverseIdentifierOccurrences'] = adverse_id_count
    review['koreanReportLostCompletionIdentifierOccurrences'] = completion_loss_id_count
    review['allDeclaredConditionalIntervalsPresent'] = True
    for i, (page, slide, source) in enumerate(zip(document, presentation.slides, slides), 1):
        pdf_text = normalize(page.get_text())
        pptx_strings = []
        images = len(page.get_images())
        require(images == 0, 'Unexpected raster image in vector PDF page ' + str(i))
        review['pdfImages'] += images
        out_of_bounds = []
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                for span in line.get('spans', []):
                    x0, y0, x1, y1 = span['bbox']
                    if min(x0, y0) < -0.5 or x1 > page.rect.width + 0.5 or y1 > page.rect.height + 0.5:
                        out_of_bounds.append(span['text'])
        require(not out_of_bounds, f'PDF text escapes page {i}: {out_of_bounds[:2]}')
        tables = []
        texts = {}
        for shape in slide.shapes:
            require(shape.shape_type != MSO_SHAPE_TYPE.PICTURE, 'PPTX contains a raster slide picture')
            require(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= presentation.slide_width + 1 and
                    shape.top + shape.height <= presentation.slide_height + 1, f'PPTX shape escapes slide {i}')
            if shape.has_text_frame:
                pptx_strings.append(shape.text)
                texts[shape.name] = shape.text
                if shape.text.strip(): review['pptxEditableTextShapes'] += 1
            if shape.has_table:
                tables.append(shape.table)
                review['pptxEditableTables'] += 1
                pptx_strings.extend(cell.text for row in shape.table.rows for cell in row.cells)
        pptx_text = normalize(' '.join(pptx_strings))
        required = [source['title'], source['text'], source['note']]
        required.extend(source['flow'] or [])
        for row in source['chart'] or []:
            required.extend([row['label'], f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}"])
        if source['table']:
            headers, rows = source['table']
            required.extend(str(value) for row in [headers, *rows] for value in row)
            require(len(tables) == 1, 'Missing native editable table on slide ' + str(i))
            require(len(tables[0].rows) == len(rows) + 1 and len(tables[0].columns) == len(headers), 'PPTX native table dimensions differ')
            for y, row in enumerate([headers, *rows]):
                for x, value in enumerate(row):
                    require(normalize(tables[0].cell(y, x).text) == normalize(value), f'PPTX exact table cell differs at {i}/{y}/{x}')
        else:
            require(not tables, 'Unexpected native table on slide ' + str(i))
        require(normalize(texts.get('Slide title')) == normalize(source['title']) and
                normalize(texts.get('Main explanation')) == normalize(source['text']) and
                normalize(texts.get('Evidence boundary note')) == normalize(source['note']), 'PPTX main text differs on slide ' + str(i))
        for value in required:
            expected = normalize(value)
            require(expected in pdf_text, f'PDF source text absent on page {i}: {str(value)[:90]}')
            require(expected in pptx_text, f'PPTX source text absent on page {i}: {str(value)[:90]}')
        review['pages'].append({'page': i, 'title': source['title'].replace('\n', ' '),
                                'requiredTextItems': len(required), 'sourceTextCovered': True,
                                'pdfTextWithinPage': True, 'pptxShapesWithinSlide': True,
                                'editableNativeTableCount': len(tables)})
    require(review['pptxEditableTextShapes'] > 100 and review['pptxEditableTables'] >= 8,
            'Deck lacks the expected editable native text and tables')
    contact_path = ROOT / 'artifacts/v9-materials-contact-sheet.png'
    contact_path.parent.mkdir(parents=True, exist_ok=True)
    sheet = fitz.open()
    contact = sheet.new_page(width=1920, height=2292)
    for i in range(len(document)):
        column, row = i % 3, i // 3
        x, y = column * 640, row * 382
        contact.show_pdf_page(fitz.Rect(x+2, y+2, x+638, y+360), document, i)
        contact.insert_text((x+8, y+375), f'PORTPLAN V9 / {i+1:02d} / 18', fontsize=11)
    contact.get_pixmap(alpha=False).save(contact_path)
    sheet.close(); document.close()
    review['outputHashes'] = {str(path.relative_to(ROOT)): sha(path) for path in paths}
    review['contactSheet'] = {'path': str(contact_path.relative_to(ROOT)), 'sha256': sha(contact_path),
                              'scope': 'Full 18-page PDF contact sheet for human visual inspection; no PowerPoint application screenshot.'}
    destination = ROOT / 'data/material-review-v9.json'
    destination.write_text(json.dumps(review, ensure_ascii=False, indent=2)+'\n')
    return review


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--require-results',action='store_true',help='Completeness is always required; retained for release commands.')
    parser.add_argument('--check-only',action='store_true',help='Validate evidence and the narrative without rendering or writing materials.')
    parser.add_argument('--bootstrap-review',default='data/bootstrap-review-v9.json',help='Required independent interval review bound to the same formal index SHA.')
    parser.add_argument('--replace-generated',action='store_true',help='Explicitly regenerate these V9 materials only.')
    args=parser.parse_args(argv)
    evidence=verified_inputs(bootstrap_review=args.bootstrap_review)
    slides=build_slides(evidence)
    if args.check_only:
        print(json.dumps({'valid':True,'pages':len(slides),'writes':0,'evidenceHashes':evidence['hashes']}));return
    targets=[ROOT/'docs'/name for name in OUTPUTS]
    if not args.replace_generated:require(not any(path.exists() for path in targets),'V9 materials already exist; use --replace-generated explicitly')
    report,english,contribution=material_texts(evidence,slides)
    pdf=pdf_document(slides)
    pptx=pptx_document(slides)
    require(pdf.startswith(b'%PDF-'), 'PDF generation failed')
    require(pptx.startswith(b'PK'), 'Editable presentation generation failed')
    outputs={targets[0]:html_document(slides).encode(),targets[1]:pdf,targets[2]:report.encode(),targets[3]:english.encode(),targets[4]:contribution.encode(),targets[5]:pptx}
    for path,data in outputs.items():path.write_bytes(data)
    material_review=review_generated(evidence,slides)
    print(json.dumps({'valid':True,'pages':len(slides),'written':[str(p.relative_to(ROOT)) for p in outputs],
        'outputHashes':{str(p.relative_to(ROOT)):sha(p) for p in outputs},
        'materialReview':'data/material-review-v9.json','editableTextShapes':material_review['pptxEditableTextShapes'],
        'editableTables':material_review['pptxEditableTables']},ensure_ascii=False))


if __name__=='__main__':
    main()
