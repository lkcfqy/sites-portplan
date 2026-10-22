#!/usr/bin/env python3
"""Build V14-study materials for the existing V13 algorithm from completed dual-reviewed evidence.

No solver, new inputs, network, old-material writes or prefilled formal results.
Korean slide renderers retain the V12 vector/editable style in this new file.
"""
from __future__ import annotations
import argparse
from collections import Counter
from html import escape
from io import BytesIO
import hashlib
import json
import math
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
METHODS = ('costPolishV7','adaptiveV9','highsV8','stableHiGHSV10','commitmentHiGHSV11','seededHiGHSV12','guardedHiGHSV13')
LABELS = dict(zip(METHODS,('V7','V9','native HiGHS','V10','V11','V12','V13 guard')))
PANELS = {'same-distribution':'동일 분포 · 새 주문','stress':'스트레스 · 새 주문'}
PANELS_ZH = {'same-distribution':'同分布 · 新订单','stress':'压力扰动 · 新订单'}
PRIMARY_PAIRS = {'v13VsV12':('seededHiGHSV12','guardedHiGHSV13'),'v13VsV11':('commitmentHiGHSV11','guardedHiGHSV13'),
 'v13VsV10':('stableHiGHSV10','guardedHiGHSV13'),'v13VsV9':('adaptiveV9','guardedHiGHSV13'),
 'v13VsHiGHS':('highsV8','guardedHiGHSV13'),'v13VsV7':('costPolishV7','guardedHiGHSV13')}
SECONDARY_PAIRS = {'v12VsV11':('commitmentHiGHSV11','seededHiGHSV12'),'v12VsV10':('stableHiGHSV10','seededHiGHSV12'),
 'v12VsV9':('adaptiveV9','seededHiGHSV12'),'v12VsHiGHS':('highsV8','seededHiGHSV12'),
 'v12VsV7':('costPolishV7','seededHiGHSV12'),'v11VsV10':('stableHiGHSV10','commitmentHiGHSV11'),
 'v10VsHiGHS':('highsV8','stableHiGHSV10'),'highsVsV7':('costPolishV7','highsV8')}
PAIRS = {**PRIMARY_PAIRS,**SECONDARY_PAIRS}
OUTPUTS = ('defense-v14.ko.html','defense-v14.ko.pdf','defense-v14.ko.pptx','submission-v14.ko.md','report-v14.zh.md','defense-v14.en.md','contribution-log-v14.md','entrant-defense-v14.blank.md')
MATERIAL_REVIEW = 'data/material-review-v14.json'


def require(condition, message):
    if not condition:
        raise ValueError(message)

def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()

def relative_path(root, value):
    require(isinstance(value, str), 'Evidence path must be a string')
    path = (root / value).resolve()
    require(path.is_relative_to(root.resolve()), 'Evidence path escapes project: ' + value)
    require(path.is_file(), 'Required completed evidence is missing: ' + value)
    return path

def load(root, name, hashes):
    path = relative_path(root, name)
    hashes[name] = sha(path)
    return json.loads(path.read_text())

def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)

def close(actual, expected, label):
    require(finite(actual) and finite(expected) and abs(actual - expected) <= max(1e-5, abs(expected) * 1e-10),
            'Independent point estimate differs: ' + label)

def same_interval(actual, expected, label, tolerance):
    if expected is None:
        require(actual is None, 'Independent interval null differs: ' + label)
        return
    require(isinstance(actual, list) and isinstance(expected, list) and len(actual) == len(expected) == 2,
            'Malformed independent interval: ' + label)
    for a, b in zip(actual, expected):
        require(finite(a) and finite(b) and abs(a - b) <= tolerance,
                'Independent interval differs: ' + label)

def same_tree(actual, expected, label):
    if isinstance(expected, dict):
        require(isinstance(actual, dict) and set(actual) == set(expected), 'Independent tree keys differ: '+label)
        for key in expected: same_tree(actual[key], expected[key], label+'/'+key)
    elif isinstance(expected, list):
        require(isinstance(actual, list) and len(actual) == len(expected), 'Independent tree length differs: '+label)
        for i, (left, right) in enumerate(zip(actual, expected)): same_tree(left, right, label+'/'+str(i))
    elif finite(expected):
        require(finite(actual) and abs(actual-expected) <= 1e-6, 'Independent tree number differs: '+label)
    else:
        require(type(actual) is type(expected) and actual == expected, 'Independent tree value differs: '+label)

def fnum(value, digits=0):
    if not finite(value):
        return 'null'
    return f'{value:,.{digits}f}'

def pct(value):
    return fnum(value, 2) + '%' if finite(value) else 'null'

def ci(value):
    # Keep the raw sign of small endpoints; rounding must not turn a strictly
    # positive/negative bound into an apparent zero for the declared gates.
    return '[' + ', '.join(format(v, '.10g') if finite(v) else 'null' for v in value) + ']' if isinstance(value, list) and len(value) == 2 else 'null'

def validate_adverse_indices(cases, summary):
    """Derive every negative/lost ID from the complete original scalar receipts."""
    occurrences = 0
    for panel in PANELS:
        selected = [case for case in cases if case['panel'] == panel]
        for name, (baseline, candidate) in PAIRS.items():
            pair = summary['panels'][panel]['comparisons'][name]
            gains, losses = [], []
            for case in selected:
                for repeat in case['repeats']:
                    before, after = repeat['policies'][baseline], repeat['policies'][candidate]
                    identifier = case['id'] + '::repeat-' + str(repeat['repeatIndex'])
                    if before['complete'] and not after['complete']: losses.append(identifier)
                    if after['complete'] and not before['complete']: gains.append(identifier)
            require(pair['completion'].get('lostEpisodeIds') == losses and pair['completion'].get('gainedEpisodeIds') == gains,
                    'Missing or changed complete gain/loss identifiers: '+panel+'/'+name)
            require(pair['completion']['lostEpisodeCount'] == len(losses) and pair['completion']['gainedEpisodeCount'] == len(gains),
                    'Completion loss/gain counts disagree with receipts')
            occurrences += len(losses)
            for key, field in (('total', 'totalCost'), ('incremental', 'incrementalCost')):
                negative_scenarios, negative_episodes = [], []
                for case in selected:
                    differences = []
                    for repeat in case['repeats']:
                        before, after = repeat['policies'][baseline], repeat['policies'][candidate]
                        if before['complete'] and after['complete']:
                            require(finite(before[field]) and finite(after[field]), 'Completed receipt fee is missing')
                            difference = after[field]-before[field]; differences.append(difference)
                            if difference > 1e-6: negative_episodes.append(case['id']+'::repeat-'+str(repeat['repeatIndex']))
                    if differences and sum(differences)/len(differences) > 1e-6: negative_scenarios.append(case['id'])
                fee = pair['fees'][key]
                require(fee.get('negativeScenarioIds') == negative_scenarios and fee.get('negativeEpisodeIds') == negative_episodes,
                        'Missing or changed all negative identifiers: '+panel+'/'+name+'/'+key)
                require(fee['scenarioWinTieLoss']['loss'] == len(negative_scenarios), 'Negative scenario denominator differs')
                occurrences += len(negative_scenarios)+len(negative_episodes)
    return occurrences

def validate_parent_lifecycle(result):
    require(result.get('endToEndCompletedWithinBudget') is True and result.get('episodeCompletedWithinBudget') is True and
            result.get('terminalReceiptAcceptedAsHistory') is True and result.get('normalWorkerExit') is True and
            result.get('exitCode') == 0 and result.get('exitSignal') is None and result.get('deadlineReached') is False,
            'Completed episode lacks normal accepted parent lifecycle')
    cleanup=result.get('processGroupCleanup',{})
    require(cleanup.get('successful') is True and cleanup.get('verifiedProcessGroupGone') is True and
            cleanup.get('verificationResult') == 'ESRCH' and cleanup.get('verificationError') is None,
            'Completed episode lacks verified process-group cleanup')
    for field in ('wallTimeMs','endToEndWallTimeMs','processExitWallTimeMs','processCloseWallTimeMs','cleanupFinishedMs','receiptValidatedMs'):
        require(finite(result.get(field)) and 0<=result[field]<=5000,'Completed lifecycle exceeds 5s or lacks '+field)

def presentation_table(source):
    """Short CI display for slides only; reports retain the source table precision."""
    if not source['table']:
        return None
    headers, rows = source['table']
    columns=[i for i,header in enumerate(headers) if '95% CI' in header]
    if not columns:
        return headers, rows
    def endpoint(value):
        if not finite(value): return 'null'
        if value == 0: return '0'
        # Do not round a tiny strictly positive/negative endpoint into zero.
        return format(value, '.2g') if abs(value) < 0.00005 else f'{value:.4f}'
    displayed=[]
    for row in rows:
        cells=list(row)
        for column in columns:
            value=json.loads(cells[column])
            cells[column]='['+', '.join(endpoint(v) for v in value)+']' if isinstance(value,list) else 'null'
        displayed.append(cells)
    return headers, displayed

def presentation_note(source):
    return source['note']+(' CI는 약 소수 4자리로 표시하고 작은 값의 부호는 보존합니다. 전체 정밀도는 보고서·JSON을 참조합니다.'
                           if source['table'] and any('95% CI' in header for header in source['table'][0]) else '')

def presentation_table_layout(headers, rows):
    dense = len(rows)>7 or len(headers)>=5
    size = 16 if len(rows)>7 or any('95% CI' in header for header in headers) else 18
    widths = [1156/len(headers)]*len(headers)
    if len(headers)==2: widths=[550,606]
    elif len(headers)==6:
        widths=[255,145,145,120,250,241] if '조건부 95% CI' in headers else [260,179.2,179.2,179.2,179.2,179.2]
    elif len(headers)==5: widths=[160,120,285,285,306]
    return size, size*1.35, 4 if dense else 9, widths

def html_document(slides):
    parts=['<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PortPlan V14</title><style>',
        'body{margin:0;background:#dce8e0;color:#173c40;font-family:"Noto Sans CJK KR","Noto Sans KR",sans-serif}.slide{box-sizing:border-box;width:1280px;min-height:720px;margin:24px auto;padding:44px 62px;background:#f7faf5;position:relative;page-break-after:always}h1{font-size:35px;line-height:1.4;margin:10px 0 22px;white-space:pre-line}.kicker{font-size:13px;letter-spacing:3px;color:#39826d}p{font-size:20px;line-height:1.7}table{width:100%;border-collapse:collapse;font-size:15px;margin:22px 0}th,td{text-align:left;border-bottom:1px solid #d3e0d3;padding:12px 10px;overflow-wrap:anywhere}th{background:#e6f0e4}.note{font-size:14px;background:#e8f2e5;padding:14px 18px;line-height:1.8}.flow{display:flex;gap:24px;align-items:stretch;margin:38px 0}.flow div{flex:1;background:#e0efe0;border-left:5px solid #36836c;padding:24px;font-size:23px;white-space:pre-line}.chart{margin:22px 0}.row{display:grid;grid-template-columns:240px 1fr 250px;align-items:center;gap:16px;margin:19px 0;font-size:17px}.track{height:21px;background:#dfebde}.fill{height:100%;background:#35826d}.footer{font-size:12px;letter-spacing:.5px;margin-top:25px;color:#638176}@media(max-width:1300px){.slide{width:100%;min-height:0;margin:12px auto;padding:30px}h1{font-size:29px}.row{grid-template-columns:180px 1fr 190px}}@media print{@page{size:1280px 720px;margin:0}body{background:white}.slide{width:1280px;height:720px;margin:0;overflow:hidden}}',
        '</style><body>']
    for i,slide in enumerate(slides,1):
        parts.extend([f'<section class="slide" id="slide-{i}"><div class="kicker">{i:02d} / PORTPLAN V14</div>',
                      '<h1>'+escape(slide['title'])+'</h1><p>'+escape(slide['text'])+'</p>'])
        if slide['flow']:parts.append('<div class="flow">'+''.join('<div>'+escape(text)+'</div>' for text in slide['flow'])+'</div>')
        if slide['table']:
            headers,rows=presentation_table(slide);size,_,_,_=presentation_table_layout(headers,rows);parts.append(f'<table style="font-size:{size}px"><thead><tr>'+''.join('<th>'+escape(str(h))+'</th>' for h in headers)+'</tr></thead><tbody>')
            parts.extend('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows);parts.append('</tbody></table>')
        if slide['chart']:
            parts.append('<div class="chart">')
            for c in slide['chart']:
                width=100*c['value']/c['denominator'] if c['denominator'] else 0
                parts.append(f'<div class="row"><span>{escape(c["label"])}</span><div class="track"><div class="fill" style="width:{width:.4f}%"></div></div><span>{c["value"]}/{c["denominator"]} · 3/3 {c["all3"]} · 미확보 {c["unknown"]}</span></div>')
            parts.append('</div>')
        parts.extend(['<div class="note">'+escape(presentation_note(slide))+'</div>',f'<div class="footer">AI 작성 준비 자료 · 본인 검토·현장 검증 미확인 / {i:02d} / {len(slides):02d}</div></section>'])
    parts.append('</body></html>');return '\n'.join(parts)+'\n'

def register_noto(text):
    """Embed Noto KR; convert only the used CFF glyphs to quadratic outlines in memory."""
    from fontTools import subset
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.cu2quPen import Cu2QuPen
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as PDFont
    path=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    require(path.is_file(), 'Noto Sans CJK font is required for a readable Korean PDF')
    font=TTFont(path,fontNumber=1)
    require('KR' in font['name'].getDebugName(1), 'Expected Noto Korean font face')
    chars=set(ord(c) for c in text if not c.isspace());cmap=font.getBestCmap()
    require(not (chars-set(cmap)), 'Noto is missing required PDF glyphs: '+repr(chars-set(cmap)))
    subsetter=subset.Subsetter();subsetter.populate(text=text);subsetter.subset(font)
    order=font.getGlyphOrder();glyphset=font.getGlyphSet();glyf={}
    for name in order:
        pen=TTGlyphPen(glyphset);glyphset[name].draw(Cu2QuPen(pen,1.0,reverse_direction=True));glyf[name]=pen.glyph()
    builder=FontBuilder(font['head'].unitsPerEm,isTTF=True);builder.setupGlyphOrder(order);builder.setupCharacterMap(font.getBestCmap());builder.setupGlyf(glyf)
    builder.setupHorizontalMetrics(font['hmtx'].metrics);builder.setupHorizontalHeader(ascent=font['hhea'].ascent,descent=font['hhea'].descent)
    builder.setupNameTable({'familyName':'Noto Sans CJK KR PortPlan subset','styleName':'Regular','uniqueFontIdentifier':'PortPlanNotoKR','fullName':'PortPlan Noto KR subset','psName':'PortPlanNotoKR',
                          'copyright':'Noto Sans CJK. Copyright Adobe and Google. SIL Open Font License 1.1. Subset outlines converted for embedding.'})
    builder.setupOS2(sTypoAscender=font['hhea'].ascent,sTypoDescender=font['hhea'].descent,usWinAscent=font['hhea'].ascent,usWinDescent=-font['hhea'].descent)
    builder.setupPost();builder.setupMaxp();buffer=BytesIO();builder.save(buffer);buffer.seek(0);pdfmetrics.registerFont(PDFont('PortPlanNoto',buffer))

def pdf_document(slides):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle
    all_text=json.dumps(slides,ensure_ascii=False)+' '.join(presentation_note(slide) for slide in slides)+' PORTPLAN V14 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789'
    register_noto(all_text)
    buffer=BytesIO();c=canvas.Canvas(buffer,pagesize=(1280,720),invariant=1,pageCompression=1)
    c.setTitle('PortPlan V14 · Korean competition preparation');c.setAuthor('AI-assisted PortPlan research; entrant verification pending')
    bg=HexColor('#f7faf5');ink=HexColor('#173c40');green=HexColor('#35826d');light=HexColor('#e6f0e4')
    def paragraph(text,x,y,w,size=20,leading=None):
        style=ParagraphStyle('p',fontName='PortPlanNoto',fontSize=size,leading=leading or size*1.55,textColor=ink,wordWrap='CJK')
        obj=Paragraph(escape(str(text)).replace('\n','<br/>'),style);_,h=obj.wrap(w,700)
        obj.drawOn(c,x,y-h);return y-h
    for index,slide in enumerate(slides,1):
        c.setFillColor(bg);c.rect(0,0,1280,720,stroke=0,fill=1)
        c.setFillColor(green);c.setFont('PortPlanNoto',13);c.drawString(62,672,f'{index:02d} / PORTPLAN V14')
        y=paragraph(slide['title'],62,642,1156,33,43)-15
        y=paragraph(slide['text'],62,y,1156,18,27)-19
        if slide['flow']:
            w=(1156-44)/3;h=118
            for j,text in enumerate(slide['flow']):
                x=62+j*(w+22);c.setFillColor(light);c.roundRect(x,y-h,w,h,7,stroke=0,fill=1);paragraph(text,x+18,y-22,w-36,23,34)
                if j<2:c.setStrokeColor(green);c.setLineWidth(2);c.line(x+w+4,y-h/2,x+w+18,y-h/2);c.line(x+w+13,y-h/2+5,x+w+18,y-h/2);c.line(x+w+13,y-h/2-5,x+w+18,y-h/2)
            y-=h+24
        if slide['chart']:
            for row in slide['chart']:
                paragraph(row['label'],62,y-5,252,17,23)
                c.setFillColor(light);c.rect(322,y-24,580,20,stroke=0,fill=1)
                c.setFillColor(green);c.rect(322,y-24,580*row['value']/row['denominator'] if row['denominator'] else 0,20,stroke=0,fill=1)
                paragraph(f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y-4,294,15,22);y-=53
            y-=8
        if slide['table']:
            headers,rows=presentation_table(slide);font_size,leading,padding,widths=presentation_table_layout(headers,rows)
            style=ParagraphStyle('cell',fontName='PortPlanNoto',fontSize=font_size,leading=leading,textColor=ink,wordWrap='CJK')
            data=[[Paragraph(escape(str(v)),style) for v in row] for row in [headers,*rows]]
            table=Table(data,colWidths=widths,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),light),('VALIGN',(0,0),(-1,-1),'TOP'),
                ('BOTTOMPADDING',(0,0),(-1,-1),padding),('TOPPADDING',(0,0),(-1,-1),padding),('LEFTPADDING',(0,0),(-1,-1),10),
                ('LINEBELOW',(0,0),(-1,-1),0.5,HexColor('#d3e0d3'))]))
            _,h=table.wrap(1156,600);require(y-h>100,f'PDF table overflows page {index}: {y-h:.1f}');table.drawOn(c,62,y-h);y-=h+18
        note_style=ParagraphStyle('note',fontName='PortPlanNoto',fontSize=12,leading=19,textColor=ink,wordWrap='CJK')
        note=Paragraph(escape(presentation_note(slide)),note_style);_,h=note.wrap(1116,300)
        require(y-h-28>64,f'PDF note overflows page {index}: {y-h-28:.1f}')
        c.setFillColor(light);c.roundRect(62,y-h-28,1156,h+25,5,stroke=0,fill=1);note.drawOn(c,82,y-h-14)
        c.setFillColor(HexColor('#638176'));c.setFont('PortPlanNoto',10)
        c.drawString(62,32,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인');c.drawRightString(1218,32,f'PORTPLAN V14 / {index:02d} / {len(slides):02d}');c.showPage()
    c.save();return buffer.getvalue()

def pptx_document(slides):
    """Editable native text, tables and vector shapes; no raster page pictures."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from reportlab.pdfbase import pdfmetrics

    require(1 <= len(slides) <= 20, 'Editable deck must contain at most 20 slides')
    # pdf_document normally registers the same outline metrics first; direct
    # renderer use remains possible without rendering a PDF or writing a font.
    if 'PortPlanNoto' not in pdfmetrics.getRegisteredFontNames():
        register_noto(json.dumps(slides,ensure_ascii=False)+' '.join(presentation_note(slide) for slide in slides)+' PORTPLAN V14 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789')
    presentation=Presentation();presentation.slide_width=Inches(1280/96);presentation.slide_height=Inches(720/96)
    presentation.core_properties.title='PortPlan V14 · Korean competition preparation'
    presentation.core_properties.author='AI-assisted PortPlan research; entrant verification pending'
    presentation.core_properties.subject='Editable result-driven research deck; synthetic daily model'
    colors={'background':'F7FAF5','ink':'173C40','green':'35826D','light':'E6F0E4','muted':'638176'}
    unit=lambda px: Inches(px/96)

    def wrapped(value,width,size):
        lines=[]
        for raw in str(value).split('\n'):
            line='';used=0
            for char in raw:
                advance=pdfmetrics.stringWidth(char,'PortPlanNoto',size)
                if line and used+advance>width:
                    lines.append(line.rstrip());line='';used=0
                line+=char;used+=advance
            lines.append(line.rstrip())
        return '\n'.join(lines),len(lines)

    def style_text(frame,text,size,leading,color='ink'):
        frame.clear();frame.text=text;frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        frame.vertical_anchor=MSO_ANCHOR.TOP
        for paragraph in frame.paragraphs:
            paragraph.font.name='Noto Sans CJK KR';paragraph.font.size=Pt(size*.75)
            paragraph.font.color.rgb=RGBColor.from_string(colors[color])
            paragraph.line_spacing=Pt(leading*.75);paragraph.space_before=Pt(0);paragraph.space_after=Pt(0)
            properties=paragraph._p.get_or_add_pPr().get_or_add_defRPr()
            east_asian=properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
            if east_asian is None:
                east_asian=OxmlElement('a:ea');properties.append(east_asian)
            east_asian.set('typeface','Noto Sans CJK KR')

    def text(slide,value,x,y,width,size=18,leading=27,color='ink',height=None,name='Editable text'):
        content,line_count=wrapped(value,width-2,size)
        measured=line_count*leading+3
        box=slide.shapes.add_textbox(unit(x),unit(y),unit(width),unit(height or measured));box.name=name
        style_text(box.text_frame,content,size,leading,color)
        return y+(height or measured)

    def rect(slide,x,y,width,height,color='light',kind=MSO_SHAPE.RECTANGLE,name='Editable vector'):
        shape=slide.shapes.add_shape(kind,unit(x),unit(y),unit(width),unit(height));shape.name=name
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(colors[color]);shape.line.fill.background()
        return shape

    for index,source in enumerate(slides,1):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string(colors['background'])
        text(slide,f'{index:02d} / PORTPLAN V14',62,31,1156,13,20,'green',name='Slide index')
        y=text(slide,source['title'],62,72,1156,33,43,name='Slide title')+15
        y=text(slide,source['text'],62,y,1156,18,27,name='Main explanation')+19
        if source['flow']:
            width=(1156-44)/3
            for j,label in enumerate(source['flow']):
                x=62+j*(width+22);rect(slide,x,y,width,118,name='Flow node '+str(j+1))
                rect(slide,x,y,5,118,'green',name='Flow node accent')
                text(slide,label,x+18,y+22,width-36,23,34,height=80,name='Editable flow label')
                if j<2:rect(slide,x+width+4,y+52,14,14,'green',MSO_SHAPE.CHEVRON,'Editable flow arrow')
            y+=142
        if source['chart']:
            for row in source['chart']:
                text(slide,row['label'],62,y+5,252,17,23,name='Chart method label')
                rect(slide,322,y+4,580,20,name='Editable completion bar track')
                width=580*row['value']/row['denominator'] if row['denominator'] else 0
                if width:rect(slide,322,y+4,width,20,'green',name='Editable completion bar')
                text(slide,f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y+4,294,15,22,name='Chart exact value')
                y+=53
            y+=8
        if source['table']:
            headers,rows=presentation_table(source);size,leading,padding,widths=presentation_table_layout(headers,rows)
            formatted=[];heights=[]
            for row in [headers,*rows]:
                cells=[wrapped(value,widths[j]-20,size) for j,value in enumerate(row)]
                formatted.append([value for value,_ in cells]);heights.append(max(count for _,count in cells)*leading+2*padding+3)
            height=sum(heights);require(y+height<620,f'Editable table overflows slide {index}')
            native=slide.shapes.add_table(len(formatted),len(headers),unit(62),unit(y),unit(1156),unit(height));native.name='Editable evidence table'
            table=native.table;table.first_row=True
            for j,width in enumerate(widths):table.columns[j].width=unit(width)
            for i,row in enumerate(formatted):
                table.rows[i].height=unit(heights[i])
                for j,value in enumerate(row):
                    cell=table.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(colors['light' if i==0 else 'background'])
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
                    style_text(cell.text_frame,value,size,leading)
                    cell.text_frame.auto_size=MSO_AUTO_SIZE.NONE
                    # Keep table padding after generic text-frame styling.
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
            y+=height+18
        note,note_lines=wrapped(presentation_note(source),1116,12);note_height=note_lines*19+3
        require(y+note_height+28<656,f'Editable note overflows slide {index}')
        rect(slide,62,y,1156,note_height+25,name='Editable boundary note background')
        text(slide,note,82,y+11,1116,12,19,height=note_height+3,name='Evidence boundary note')
        text(slide,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인',62,678,900,10,15,'muted',name='Contribution footer')
        text(slide,f'PORTPLAN V14 / {index:02d} / {len(slides):02d}',1040,678,178,10,15,'muted',name='Page footer')
        for shape in slide.shapes:
            require(shape.shape_type!=MSO_SHAPE_TYPE.PICTURE,'Raster slide pictures are prohibited in the editable deck')
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=presentation.slide_width+1 and
                    shape.top+shape.height<=presentation.slide_height+1,f'Editable shape escapes slide {index}: {shape.name}')
    require(len(presentation.slides)==len(slides),'Editable deck page count changed')
    output=BytesIO();presentation.save(output);return output.getvalue()

def markdown_table(headers, rows):
    cell=lambda value:str(value).replace('|','\\|').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(map(cell,headers))+' |','| '+' | '.join(['---']*len(headers))+' |',
                      *['| '+' | '.join(map(cell,row))+' |' for row in rows]])

def audit_rendered(pdf,pptx,slides):
    import fitz
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    import unicodedata
    normalize=lambda text:''.join(unicodedata.normalize('NFKC',str(text)).split())
    document=fitz.open(stream=pdf,filetype='pdf');deck=Presentation(BytesIO(pptx));require(len(document)==len(deck.slides)==len(slides)<=20,'Rendered material page count differs or exceeds 20')
    text_shapes=tables=0; table_font_sizes=[]
    for number,(page,slide,source) in enumerate(zip(document,deck.slides,slides),1):
        require(not page.get_images(),'PDF must use vector/text content, not slide screenshots')
        for block in page.get_text('dict')['blocks']:
            if 'lines' in block:
                x0,y0,x1,y1=block['bbox'];require(x0>=-1 and y0>=-1 and x1<=page.rect.width+1 and y1<=page.rect.height+1,'PDF text escapes page '+str(number))
        ppt_strings=[]
        for shape in slide.shapes:
            require(shape.shape_type!=MSO_SHAPE_TYPE.PICTURE,'PPT must remain native editable')
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=deck.slide_width+1 and shape.top+shape.height<=deck.slide_height+1,'PPT shape escapes slide')
            if shape.has_text_frame:ppt_strings.append(shape.text);text_shapes+=bool(shape.text.strip())
            if shape.has_table:
                tables+=1;ppt_strings.extend(cell.text for row in shape.table.rows for cell in row.cells)
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            if paragraph.text:
                                require(paragraph.font.size is not None and paragraph.font.size.pt>=12,'Editable table text below 16px/12pt minimum')
                                table_font_sizes.append(paragraph.font.size.pt)
        expected=[source['title'],source['text'],presentation_note(source),*(source['flow'] or [])]
        if source['table']:
            headers,rows=presentation_table(source);expected.extend(value for row in [headers,*rows] for value in row)
        pdf_text=normalize(page.get_text());ppt_text=normalize(' '.join(ppt_strings))
        for value in expected:require(normalize(value) in pdf_text and normalize(value) in ppt_text,'Rendered source text missing on page '+str(number)+': '+str(value)[:70])
    require(text_shapes>=len(slides)*3 and tables>=1,'Editable native objects missing')
    return {'valid':True,'schemaVersion':'portplan-v14-material-review','pdfPageCount':len(document),'pptxSlideCount':len(deck.slides),'editableTextShapes':text_shapes,'editableTables':tables,'minimumEditableTableFontPt':min(table_font_sizes),'sameSlideSource':True,'sourceTextCoverage':True,'pdfAndPptBoundsValid':True,'solverRuns':0,'browserInteractionCertified':False,'entrantUnderstandingCertified':False}

def verified_inputs(root=ROOT, bootstrap_review='data/bootstrap-review-v14.json'):
    root = Path(root).resolve()
    hashes = {}
    release = load(root, 'data/v14-release.json', hashes)
    require(all(release.get(key) is True for key in ('formalExperimentComplete','independentReviewComplete','independentBootstrapReviewComplete')),
            'V14 formal evidence and both reviews must complete before materials')
    wrapper = load(root, 'data/v14-results.json', hashes)
    require(wrapper.get('schemaVersion') == 'portplan-v14-published-results' and wrapper.get('executionState') == 'complete' and
            not wrapper.get('fixture'), 'V14 results are not completed formal evidence')
    validate_publication_header(release,wrapper)
    reference = wrapper.get('formalIndex', {})
    formal_path = relative_path(root, reference.get('path'))
    require(sha(formal_path) == reference.get('sha256'), 'Formal index SHA differs')
    hashes[reference['path']] = reference['sha256']; index = json.loads(formal_path.read_text())
    require(index.get('schemaVersion') == 'portplan-v14-streamed-experiment' and index.get('executionState') == 'complete' and
            not index.get('fixture'), 'Restored formal run is incomplete or a fixture')
    require(wrapper.get('summary') == index.get('summary') and wrapper.get('protocol') == index.get('protocol'), 'Published summary/protocol differ from formal bytes')
    protocol = load(root, 'data/protocol-v14.json', hashes)
    require(protocol == index['protocol'] and hashes['data/protocol-v14.json'] == index.get('protocolHash'), 'Protocol bytes differ')
    require(tuple(protocol.get('methods', [])) == METHODS and protocol['budget']['hardWallMs'] == 5000, 'Unexpected methods or hard episode clock')
    require(protocol['budget'].get('parentAcceptance') == {'normalWorkerExitRequired':True,'confirmedProcessGroupGoneRequired':True,
            'processGroupVerificationResult':'ESRCH','cleanupWithinHardWallRequired':True,'endToEndWithinHardWallRequired':True,
            'timelyFinalAloneIsCompletion':False}, 'Missing complete parent lifecycle budget declaration')
    require(all(protocol['callBudgets'][method]['timeBudgetMs'] == 200 for method in METHODS), 'Unexpected cooperative call budget')
    require(protocol.get('executionAuthorization', {}).get('status') in ('ready-for-source-freeze','frozen'), 'Draft development declaration cannot generate formal materials')
    validate_natural_contract(protocol,index)
    frozen_dir = (formal_path.parent/index.get('frozenSource','..')).resolve()
    manifest_path = relative_path(root, str((frozen_dir/'manifest.json').relative_to(root.resolve())))
    manifest_hash = sha(manifest_path); hashes[str(manifest_path.relative_to(root))] = manifest_hash
    require(wrapper.get('frozenManifestSha256') == manifest_hash, 'Published wrapper binds to a different frozen manifest')
    manifest = json.loads(manifest_path.read_text())
    validate_frozen_binding(root,index,manifest,frozen_dir,hashes)
    review = load(root, 'data/review-v14.json', hashes); statistics = load(root, bootstrap_review, hashes)
    for value, schema in ((review,'portplan-independent-v14-run-review'),(statistics,'portplan-independent-v14-bootstrap-review')):
        validate_review_binding(value,schema,reference['sha256'],manifest_hash)
        require(value.get('schemaVersion') == schema and value.get('valid') is True and value.get('fixture',False) is False and
                value.get('readOnly') is True and value.get('solverRuns') == 0 and value.get('importsJavaScript') is False,
                'Incomplete, fixture or non-independent review')
        require(value.get('inputIndexSha256') == reference['sha256'] and value.get('frozenManifestSha256') == manifest_hash,
                'Independent review belongs to different result/frozen bytes')
    for name,value in (('scripts/review-v14.py',review),('scripts/review-bootstrap-v14.py',statistics)):
        require(value.get('implementationSha256') == manifest['implementationHashes'].get(name), 'Independent reviewer was not the frozen implementation: '+name)
    require(review.get('sourceProvenance',{}).get('sourceGitCommit') == manifest['sourceGitCommit'] and review['sourceProvenance'].get('allCopiedBlobsMatchRecordedCommitInventory') is True,'Independent committed-source receipt is missing')
    cases = index.get('cases', []); repeats = protocol.get('algorithmRepeats', [])
    require(len(cases) == len({case['id'] for case in cases}) == 144 and Counter(case['panel'] for case in cases) == Counter({'same-distribution':72,'stress':72}), 'Missing/duplicate business scenarios')
    require(len(repeats) == len({repeat['seed'] for repeat in repeats}) == 3, 'Missing fixed timing repeats')
    reviewed = {item['id']:item for item in review.get('episodes', [])}; count = 0; consumed=set()
    require(len(reviewed) == len(review.get('episodes',[])) == 3024, 'Physical review lacks unique episode coverage')
    for case in cases:
        artifact = case['inputArtifact']; path = relative_path(root, str((formal_path.parent/artifact['path']).relative_to(root)))
        require(sha(path) == artifact['sha256'], 'Scenario input SHA differs')
        require(len(case['repeats']) == 3, 'Missing repeat')
        for number, repeat in enumerate(case['repeats'],1):
            require(repeat['repeatIndex'] == number and repeat['algorithmSeed'] == repeats[number-1]['seed'] and set(repeat['policies']) == set(METHODS), 'Missing/reordered repeat/method')
            for method, receipt in repeat['policies'].items():
                ref = receipt['artifact']; path = relative_path(root, str((formal_path.parent/ref['path']).relative_to(root)))
                require(sha(path) == ref['sha256'], 'Episode artifact SHA differs')
                envelope = json.loads(path.read_text()); actual = envelope['result']; audited = reviewed.get(envelope['episodeId'])
                require(audited and audited['artifactSha256'] == ref['sha256'] and envelope['caseId'] == case['id'] and envelope['method'] == method, 'Episode/review identity binding differs')
                require(type(audited.get('complete')) is bool and audited['complete'] is actual['complete'], 'Independent episode completion differs')
                require(envelope.get('repeatIndex') == number and envelope.get('algorithmSeed') == repeats[number-1]['seed'] and
                        envelope['episodeId'] not in consumed,'Episode repeat/seed is mismatched or duplicated')
                consumed.add(envelope['episodeId'])
                require(type(receipt.get('complete')) is bool and receipt['complete'] is actual['complete'], 'Receipt completion differs')
                if actual.get('executionHistoryUnavailable') is not True:
                    require(actual.get('naturalFullD0') is True and actual.get('developmentEndpointImportEnabled') is False,'Raw episode is not natural D0 with import disabled')
                require(all(meta.get('developmentEndpointImport') is None for meta in actual.get('plannerTrace',[])),'Imported endpoint receipt in formal history')
                if receipt['complete']:
                    validate_parent_lifecycle(receipt);validate_parent_lifecycle(actual)
                for field in ('totalCost','incrementalCost'):
                    if receipt['complete']: close(receipt[field],actual[field],'receipt/'+field)
                    else: require(receipt[field] is None and actual[field] is None, 'Failed full-path fees must remain null')
                count += 1
    require(count == len(consumed) == 3024 and wrapper.get('scenarioCount') == 144 and wrapper.get('episodeCount') == count and
            review.get('summary',{}).get('episodes') == count and statistics.get('episodeCount') == count and statistics.get('businessScenarioCount') == 144,
            'Formal/independent episode coverage differs')
    summary = index['summary']; require(summary.get('inferentialState') == 'complete-scenario-cluster-bootstrap' and set(summary['panels']) == set(PANELS), 'Missing final scenario intervals')
    declared = protocol['analysis']['bootstrap']; procedure = statistics.get('procedure',{})
    require(procedure.get('samplingUnit') == 'whole business scenario' and procedure.get('seed') == declared['seed'] and procedure.get('draws') == declared['draws'], 'Bootstrap unit/seed/draws differ')
    for panel in PANELS:
        produced, physical, statistical = summary['panels'][panel], review['pointEstimates'][panel], statistics['panels'][panel]
        require(produced['scenarioCount'] == 72 and set(produced['byMethod']) == set(METHODS) and set(produced['comparisons']) == set(PAIRS), 'Unexpected panel/method/comparison coverage')
        for method in METHODS:
            for field in ('recordedEpisodeCount','completeEpisodeCount','fullyCompletedScenarioCount','anyCompletedScenarioCount'):
                require(produced['byMethod'][method][field] == physical['byMethod'][method][field] == statistical['byMethod'][method][field], 'Reviewed completion differs')
            require(produced['byMethod'][method]['recordedEpisodeCount'] == 216, 'Missing full-sample method receipts')
            same_tree(produced['byMethod'][method]['initialization'],physical['byMethod'][method]['initialization'],'D0 physical')
            same_tree(produced['byMethod'][method]['initialization'],statistical['byMethod'][method]['initialization'],'D0 statistical')
        for name,(baseline,candidate) in PAIRS.items():
            pair, point, boot = produced['comparisons'][name], physical['comparisons'][name], statistical['comparisons'][name]
            require((pair['baselineMethod'],pair['candidateMethod']) == (baseline,candidate), 'Comparison identity differs')
            require(pair['completion']['gainedEpisodeCount'] == point['gainedEpisodes'] and pair['completion']['lostEpisodeCount'] == point['lostEpisodes'], 'Completion losses/gains differ')
            same_interval(pair['completion']['uncertainty']['meanScenarioRateGain95CI'],boot['completionRateGain95CI'],'completion CI',1e-12)
            require([pair['revisions']['baselineObservedSum'],pair['revisions']['candidateObservedSum']] == point['revisions'], 'Revision denominator/numerator differs')
            for key in ('total','incremental'):
                fee, checked, sampled = pair['fees'][key], point['fees'][key], boot['fees'][key]
                require(fee['pairedScenarioCount'] == checked['pairedScenarios'] == sampled['pairedScenarioCount'] and fee['pairedEpisodeCount'] == checked['pairedEpisodes'] == sampled['pairedEpisodeCount'], 'Conditional denominators differ')
                require(fee['scenarioWinTieLoss'] == checked['scenarioWinTieLoss'], 'Fee win/tie/loss differs')
                for field in ('baselineScenarioMeanSumKRW','candidateScenarioMeanSumKRW'): close(fee[field],checked[field],field)
                uncertainty = fee['uncertainty']; require(uncertainty['unit'] == 'scenario' and uncertainty['scenarioCount'] == fee['pairedScenarioCount'] and uncertainty['draws'] == declared['draws'], 'CI eligibility differs')
                same_interval(uncertainty['weightedSavingsPct95CI'],sampled['weightedSavingsPct95CI'],'fee CI',1e-9)
                same_interval(uncertainty['meanSavings95CI'],sampled['meanSavings95CI'],'KRW CI',1e-6)
        same_tree(produced['jointCompletePareto'],statistical['jointCompletePareto'],'all-seven common set')
        gates = {key:value for key,value in produced['observedDecisionGates'].items() if key!='nativeComparison'}
        same_tree(gates,statistical['observedDecisionGates'],'observed necessary gates')
        require(gates['assessmentState'] == 'complete-panel', 'Partial-panel gates cannot publish')
    adverse_count = validate_adverse_indices(cases, summary)
    return {'root':root,'wrapper':wrapper,'index':index,'protocol':protocol,'summary':summary,'review':review,
            'bootstrapReview':statistics,'release':release,'hashes':hashes,'adverseIdentifierOccurrences':adverse_count}


def validate_publication_header(release, wrapper):
    require(release.get('schemaVersion') == 'portplan-v14-release' and
            all(release.get(k) is True for k in ('formalExperimentComplete','independentReviewComplete','independentBootstrapReviewComplete')),
            'Complete formal V14 evidence and both reviews are required')
    require(wrapper.get('schemaVersion') == 'portplan-v14-published-results' and wrapper.get('executionState') == 'complete' and
            wrapper.get('fixture',False) is False and wrapper.get('scenarioCount') == 144 and wrapper.get('episodeCount') == 3024,
            'Only the completed 144-scenario / 3024-episode formal assessment can generate materials')


def validate_review_binding(value, schema, index_sha, manifest_sha):
    require(value.get('schemaVersion') == schema and value.get('valid') is True and value.get('fixture',False) is False and
            value.get('readOnly') is True and value.get('importsJavaScript') is False and type(value.get('solverRuns')) is int and value['solverRuns'] == 0,
            'Missing valid independent read-only review (fixtures cannot publish)')
    require(value.get('inputIndexSha256') == index_sha and value.get('frozenManifestSha256') == manifest_sha,
            'Independent review binds to another first-run index or frozen source')


def validate_natural_contract(protocol, index):
    require(protocol.get('schemaVersion') == 'portplan-protocol-v14' and protocol.get('methods') == list(METHODS), 'Unexpected V14 study methods')
    require(protocol.get('rollout') == {'startDay':0,'initialization':'natural','forcedEndpointImports':False,'freshBusinessScenarios':144},
            'Formal assessment must start naturally at D0 with 144 new scenarios')
    initialization=protocol.get('initialization',{})
    require(initialization.get('mode') == 'natural-from-day-zero' and initialization.get('forcedEndpointImports') is False and
            initialization.get('developmentEndpointImport') is False, 'Development endpoint import is forbidden')
    require(index.get('rootAuthorized') is True and index.get('studyMethod') == 'guardedHiGHSV13' and index.get('initialization') == initialization,
            'Missing separate root authorization or natural V13 method identity')
    analysis=protocol['analysis']
    require(analysis.get('comparisons') == {k:list(v) for k,v in PAIRS.items()} and
            analysis.get('comparisonGroups',{}).get('primary') == list(PRIMARY_PAIRS) and
            analysis.get('comparisonGroups',{}).get('secondary') == list(SECONDARY_PAIRS), 'Changed predeclared comparison families')
    require(analysis.get('bootstrap') == {'seed':800001,'draws':2000}, 'Changed scenario bootstrap declaration')
    gate=analysis.get('defaultGateDefinition',{})
    require(gate.get('candidateMethod') == 'guardedHiGHSV13' and gate.get('comparators') == list(METHODS[:-1]) and
            gate.get('completePanelRequired') is True and gate.get('maximumMatchingRepeatCompletionLossesPerComparator') == 0 and
            gate.get('feeOrRevisionConditionInDefaultGate') is False, 'Changed completion-only observed default gate')
    require(protocol['callBudgets']['guardedHiGHSV13'] == protocol['callBudgets']['seededHiGHSV12'] and
            all('enableDevelopmentEndpointImport' not in protocol['callBudgets'][m] for m in METHODS), 'Guarded V13 uses fixed V12 budgets without development configuration')


def validate_frozen_binding(root, index, manifest, frozen_dir, hashes):
    require(manifest.get('schemaVersion') == 'portplan-v14-frozen-source' and manifest.get('generatedProspectiveInputs') is False and
            manifest.get('protocolHash') == index.get('protocolHash') and manifest.get('protocol') == index.get('protocol'), 'Frozen protocol differs')
    for key in ('sourceGitCommit','sourceGitTree','gitObjectFormat','committedBlobHashes','implementationHashes'):
        require(manifest.get(key) == index.get(key), 'Source provenance differs: '+key)
    require(manifest.get('sourceCommitVerified') is True, 'Frozen source was not committed before the assessment')
    form=manifest.get('gitObjectFormat');require(form in ('sha1','sha256'), 'Missing Git object format')
    size=40 if form=='sha1' else 64
    require(all(isinstance(manifest.get(k),str) and len(manifest[k]) == size and all(c in '0123456789abcdef' for c in manifest[k])
                for k in ('sourceGitCommit','sourceGitTree')), 'Malformed recorded source commit/tree')
    bindings=manifest.get('implementationHashes');blobs=manifest.get('committedBlobHashes')
    require(isinstance(bindings,dict) and bindings and isinstance(blobs,dict) and set(bindings)==set(blobs), 'Incomplete frozen source inventory')
    for name,digest in bindings.items():
        path=relative_path(root,str((frozen_dir/'source'/name).relative_to(root)))
        raw=path.read_bytes();require(hashlib.sha256(raw).hexdigest() == digest,'Frozen source SHA differs: '+name)
        require(hashlib.new(form,b'blob '+str(len(raw)).encode()+b'\0'+raw).hexdigest() == blobs[name], 'Committed Git blob differs: '+name)
    script='scripts/build-defense-v14.py'
    require(bindings.get(script) == sha(Path(__file__)), 'Material generator must be frozen before prospective generation')
    for name in ('scripts/review-v14.py','scripts/review-bootstrap-v14.py','data/protocol-v14.json'):
        require(name in bindings, 'Unbound independent evidence source: '+name)
    hashes['sourceGitCommit']=manifest['sourceGitCommit']


def revision_adverses(cases, panel, baseline, candidate):
    scenarios=[];episodes=[]
    for case in cases:
        if case['panel'] != panel: continue
        differences=[]
        for repeat in case['repeats']:
            before,after=repeat['policies'][baseline],repeat['policies'][candidate]
            if before['complete'] and after['complete'] and finite(before.get('revisionCount')) and finite(after.get('revisionCount')):
                difference=after['revisionCount']-before['revisionCount'];differences.append(difference)
                if difference>0: episodes.append(case['id']+'::repeat-'+str(repeat['repeatIndex']))
        if differences and sum(differences)/len(differences)>0:scenarios.append(case['id'])
    return {'scenarioIds':scenarios,'episodeIds':episodes}


def build_slides(e):
    slides=[];summary=e['summary']
    def add(title,text,table=None,flow=None,note=''):
        slides.append({'title':title,'text':text,'table':table,'flow':flow,'chart':None,'note':note})
    add('더 싼 계획을 고를 때\n내일의 필수 창고 부담도 확인합니다.',
        'V14 연구는 기존 V13 guard를 자연 D0부터 평가합니다. 새로운 V14 알고리즘이나 미래 예측 모형을 만들었다는 주장이 아닙니다.',
        flow=['현재 공개 정보','요금 + 필요 창고 M','기한 내 검산·선택'],
        note='합성 일 단위 연구입니다. 전역 최적·미래 안전·기업 실현 절감·대회 순위를 보장하지 않습니다.')
    add('반출·연속 적출·반납을\n하나의 날짜 장부에서 검산합니다.',
        '모든 주문은 시작부터 알려집니다. 계획기는 당일 공개 관측과 실제 실행 잠금만 받고, 미래 미공지 충격은 받지 않습니다.',
        table=(['날짜','자원·제약'],[['반출 p / 반납 r','각각 운송 작업; 방출·납기·과거 잠금'],['연속 적출 u','연속 창고 일수; 순서와 완료 기한'],['체류 p부터 r 전','staging 용량; 비용과 실행 장부']]),
        note='경로·기사 시간·예약 슬롯·실제 계약 청구는 이 일 단위 모형의 검증 범위 밖입니다.')
    add('공개 자료와 재현 가능한 연구로\n근거와 가정을 구분합니다.',
        '공개 요율·관측·수학 기준과 합성 주문 실험은 서로 다른 증거입니다. 기업 주문·연락처·견적·파일럿은 선택 확장입니다.',
        table=(['증거','가능한 범위'],[['공개 요율','명시한 요금 함수 재계산'],['공개 기준·관측','그 자료에 한정한 검증'],['새 합성 주문·충격','고정 모형에서 방법 간 비교']]),
        note='합성 결과를 실단·실현 매출·기업 채택으로 바꾸어 쓰지 않습니다. 기업 자료는 이 연구의 선행 조건이 아닙니다.')
    add('V13은 엄격한 요금 개선에도\n필요 창고 지표의 비증가를 요구합니다.',
        '유효 초기해가 있으면 더 싼 native 후보도 완전한 M 비교가 기한 안에 끝나고 M이 늘지 않을 때만 교체합니다. 동률이면 M이 엄격히 줄어야 합니다.',
        table=(['상태','선택'],[['기한 내 완전 incumbent 있음','엄격 절감 + M 비증가 / 동률 + M 감소'],['점수 불완전·늦음·오류','이미 검산된 incumbent 유지'],['유효 incumbent 없음','유일한 완전 후보를 M만으로 거부하지 않음']]),
        note='M은 현재 정보의 필요 하한 대리량입니다. 낮은 M도 미공지 충격에서 완주하거나 복구할 충분조건은 아닙니다.')
    add('144개 새 업무 장면을\n동일한 일곱 방법으로 비교합니다.',
        '동일 분포와 스트레스는 각 72장면입니다. 장면별 세 시계 반복은 새 업무 표본 세 개가 아닙니다.',
        table=(['고정 항목','선언'],[['행렬','3 기작 × 2 용량 × 24/50/100 × 4'],['업무 시드','791001–791072 / 792001–792072'],['반복·bootstrap','79011/79023/79039; 800001 × 2000'],['전체','144 × 3 × 7 = 3024; 전부 최초 결과 보존']]),
        note='기존 주문·과거 강제 prefix·개발 endpoint import 없이 매회 자연 D0에서 시작합니다. 실패·중복·비싼 결과를 교체하지 않습니다.')
    add('200ms 협력 호출과\n5초 냉시작 수명주기를 함께 보존합니다.',
        '구성·native·현재 감사·M 점수·선택이 같은 200ms를 나눕니다. 부모 5초는 시작 전부터 정상 종료·전체 그룹 정리·최종 감사까지 포함합니다.',
        table=(['시계','포함·접수 조건'],[['일별 200ms','협력 예산; 함수 반환의 하드 보장 아님'],['부모 5000ms','import / ready / 전체 경로 / close / IPC'],['완료 접수','정상 exit0 + ESRCH 그룹 소멸 + 최종 부모 기한'],['증거 누락·기한 실패','fees·history는 null/unknown; 원문 별도 보존']]),
        note='먼저 final을 보냈다는 이유만으로 완주를 인정하지 않습니다. 정상 종료·정리·최종 접수가 같은 5초 안이어야 합니다.')
    for panel in PANELS:
        s=summary['panels'][panel]
        rows=[]
        for method in METHODS:
            value=s['byMethod'][method];initial=value['initialization']
            rows.append([LABELS[method],str(value['completeEpisodeCount'])+'/216',str(initial['knownInitialFailureCount']),str(initial['knownLaterFailureCount']),str(initial['unknownDayZeroCount'])])
        add(PANELS[panel]+'\n완주와 초기·후속 실패를 분리합니다.',
            '완료는 부모 수명주기까지 접수된 전체 실행입니다. D0와 후속 실패를 합쳐 초기해 개선으로 포장하지 않습니다.',
            table=(['방법','완료','D0 실패','후속 실패','D0 미확보'],rows),
            note='각 방법 72장면 × 3시계 = 216회. unknown은 물리 실패의 증명이 아니며, 초기 수락이 전체 완주를 보장하지 않습니다.')
        rows=[]
        for name,(baseline,_) in PRIMARY_PAIRS.items():
            pair=s['comparisons'][name]
            rows.append([LABELS[baseline],str(pair['completion']['gainedEpisodeCount']),str(pair['completion']['lostEpisodeCount']),str(pair['fees']['incremental']['pairedScenarioCount'])+'/'+str(pair['fees']['incremental']['pairedEpisodeCount'])])
        add(PANELS[panel]+'\nV13의 이득·손실을 각 기준별로 보존합니다.',
            '이득은 기준 미완료·V13 완료, 손실은 기준 완료·V13 미완료입니다. 비용 분모는 해당 쌍이 함께 완료한 반복만 사용합니다.',
            table=(['기준','완료 이득','완료 손실','비용 장면/회'],rows),
            note='다른 기준의 이득이나 순완료 증가는 손실 한 건을 상쇄하지 않습니다. 모든 손실 ID는 중·한 보고서 부록에 있습니다.')
        rows=[]
        for name,(baseline,_) in PRIMARY_PAIRS.items():
            pair=s['comparisons'][name];fees=pair['fees'];rev=pair['revisions']
            rows.append([LABELS[baseline],str(fees['incremental']['pairedScenarioCount'])+'/'+str(fees['incremental']['pairedEpisodeCount']),
                         ci(fees['total']['uncertainty']['weightedSavingsPct95CI']),ci(fees['incremental']['uncertainty']['weightedSavingsPct95CI']),
                         fnum(rev['baselineObservedSum'])+' → '+fnum(rev['candidateObservedSum'])+' / n='+str(rev['pairedEpisodeCount'])])
        add(PANELS[panel]+'\n같은 완료 집합에서 두 비용과 변경을 봅니다.',
            '절감률의 양수는 V13이 저렴함을 뜻합니다. 전부 음수인 구간은 조건부 비용 악화이며, 0 포함은 안정적 절감의 증거가 아닙니다.',
            table=(['기준','장면/회','총요금 95% CI','증분 95% CI','변경 원합 / 알려진 n'],rows),
            note='비용은 장면 등가중, 변경은 같은 완료 집합의 알려진 원이벤트 합과 자체 n입니다. 시계 반복을 독립 CI 표본으로 보지 않습니다. 구간은 기술적 주변구간이며 동시 우월 보장이 아닙니다.')
        pareto=s['jointCompletePareto'];rows=[]
        for method in METHODS:
            value=pareto['byMethod'][method]
            rows.append([LABELS[method],fnum(value['scenarioMeanTotalCostSumKRW']),fnum(value['scenarioMeanIncrementalCostSumKRW']),fnum(value['scenarioMeanRevisionCountSum'])])
        add(PANELS[panel]+'\n일곱 방법 공통 완료 부분집합을 따로 둡니다.',
            f"공통 {pareto['jointScenarioCount']}장면 / {pareto['jointEpisodeCount']}회. 장면 내 평균 뒤 장면을 등가중 합산한 조건부 수치입니다.",
            table=(['방법','총요금 합','증분요금 합','변경 합'],rows),
            note='전체 표본 완료 손실을 이 표로 덮지 않습니다. 변경을 임의의 KRW 가중치로 바꾸거나 서로 다른 분모를 혼합하지 않습니다.')
    rows=[]
    for panel in PANELS:
        gate=summary['panels'][panel]['observedDecisionGates']
        native=gate['nativeComparisons'];allowed=sum(value['stabilityImprovementClaimAllowed'] is True for value in native.values())
        rows.append([PANELS[panel],str(gate['defaultRecommendationEligible']),str(allowed)+'/4'])
    add('완료 기본문턱과\n비용·변경 안정성 주장을 구분합니다.',
        '기본문턱은 완전 panel과 여섯 기준 모두의 완료 손실 0입니다. native 안정성은 같은 쌍의 변경 엄격 감소와 두 비용 구간 조건을 별도로 요구합니다.',
        table=(['panel','관측 기본 자격','native 안정성 허용'],rows),
        note='통과도 이 모형·포함 기준의 필요 관측 조건일 뿐입니다. 비용 조건을 몰래 기본 자격에 더하거나 빼지 않으며 실무 기본값·우승 보장이 아닙니다.')
    add('불리한 식별자와 원문을\n보고서 부록에서 빠짐없이 공개합니다.',
        f"14개 사전 비교의 완료 손실·비용 악화 식별자 {e['adverseIdentifierOccurrences']}개 출현을 원영수증에서 재도출했습니다. 변경 악화 ID도 별도로 나열합니다.",
        table=(['보존 대상','보고 방식'],[['모든 비교의 완료 손실','중·한 보고서 부록; 반복 ID 전부'],['총·증분 비용 악화','장면 ID와 반복 ID; 조건부 분모'],['변경 악화·unknown','원영수증 재도출; null을 0으로 바꾸지 않음']]),
        note='동일 ID가 여러 비교에 나올 수 있어 출현 수는 독립 업무 표본 수가 아닙니다. 실패 경로의 null 비용을 성공 비용과 비교하지 않습니다.')
    add('최초 실행·물리·통계·소스를\n같은 SHA 사슬에 묶습니다.',
        '두 독립 Python 구현은 내부 검산입니다. 기록된 시계·OS·원격 push 관측은 외부 인증이나 제삼자 보증이 아닙니다.',
        table=(['증거','SHA-256'],[['정식 최초 index',e['wrapper']['formalIndex']['sha256']],['물리·선택 검산',e['hashes']['data/review-v14.json']],['통계 검산',e['hashes']['data/bootstrap-review-v14.json']]]),
        note='원입력·실행 물리 장부·선택 증거·실패·동결 source를 보존합니다. V12 정식 결과와 V13 개발 결과는 다시 쓰지 않습니다.')
    add('제출 준비와 본인 수행은\n각각 검토해야 합니다.',
        '이 자료는 AI 도움으로 만든 연구·발표 초안입니다. 본인의 자격·등록·당해 AI 규칙·기여·재현·발표 이해는 자동으로 인증되지 않습니다.',
        flow=['본인 근거 확인','새 출력에서 재현','질문에 본인 설명'],
        note='기업 연락처를 선행 조건으로 요구하지 않습니다. 모델 범위와 실제 실패를 설명하고 대회 공식 요구에 맞춰 본인이 최종 검토합니다.')
    require(len(slides)==18 and len(slides)<=20, 'Declared deck exceeds 20 pages')
    return slides


def detailed_results(e, language):
    zh=language=='zh';lines=[]
    for panel in PANELS:
        s=e['summary']['panels'][panel]
        lines+=['## '+(PANELS_ZH[panel] if zh else PANELS[panel]),'']
        rows=[]
        for method in METHODS:
            value=s['byMethod'][method];init=value['initialization']
            rows.append([method,value['completeEpisodeCount'],value['recordedEpisodeCount'],init['acceptedInitialPlanCount'],init['knownInitialFailureCount'],init['knownLaterFailureCount'],init['unknownDayZeroCount']])
        lines += [markdown_table(['method','complete','recorded','D0 accepted','D0 failed','later failed','D0 unknown'],rows),'']
        for name,(baseline,candidate) in PAIRS.items():
            pair=s['comparisons'][name];group='primary' if name in PRIMARY_PAIRS else 'secondary'
            lines += ['### '+name+' ['+group+']',f'{baseline} → {candidate}',
                      'Completion gained/lost: '+str(pair['completion']['gainedEpisodeCount'])+'/'+str(pair['completion']['lostEpisodeCount'])+
                      '; scenario rate gain 95% CI '+ci(pair['completion']['uncertainty']['meanScenarioRateGain95CI']),'']
            rows=[]
            for key in ('total','incremental'):
                fee=pair['fees'][key]
                rows.append([key,fee['pairedScenarioCount'],fee['pairedEpisodeCount'],fnum(fee['baselineScenarioMeanSumKRW'],6),fnum(fee['candidateScenarioMeanSumKRW'],6),
                             pct(fee['weightedSavingsPct']),ci(fee['uncertainty']['weightedSavingsPct95CI']),ci(fee['uncertainty']['meanSavings95CI']),
                             fnum(fee['baselineEpisodeSumKRW']),fnum(fee['candidateEpisodeSumKRW']),fee['scenarioWinTieLoss'],fnum(fee['worstScenarioLossKRW'])])
            lines += [markdown_table(['fee','scenarios','matching episodes','baseline equal-scenario KRW','candidate equal-scenario KRW','savings %','95% savings % CI','95% mean KRW CI','baseline raw episode KRW','candidate raw episode KRW','win/tie/loss','worst scenario loss KRW'],rows),
                      'Revisions on this exact pair: '+str(pair['revisions']['baselineObservedSum'])+' → '+str(pair['revisions']['candidateObservedSum'])+'; known matching episodes='+str(pair['revisions']['pairedEpisodeCount']),'']
        lines += ['### All-seven common-complete subset / observed necessary gates','','```json',
                  json.dumps({'jointCompletePareto':s['jointCompletePareto'],'observedDecisionGates':s['observedDecisionGates']},ensure_ascii=False,indent=2),'```','']
    return '\n'.join(lines)


def adverse_appendix(e):
    lines=['## 全部负例 / 전체 불리한 식별자 부록','',
           'Membership is re-derived from all formal scalar receipts. Repeated IDs across comparisons are occurrences, not independent cases. Empty sets are explicitly [].','']
    for panel in PANELS:
        for name,(baseline,candidate) in PAIRS.items():
            pair=e['summary']['panels'][panel]['comparisons'][name]
            revisions=revision_adverses(e['index']['cases'],panel,baseline,candidate)
            sets={'completionLostEpisodeIds':pair['completion']['lostEpisodeIds'],
                  'totalFeeWorseScenarioIds':pair['fees']['total']['negativeScenarioIds'],
                  'totalFeeWorseEpisodeIds':pair['fees']['total']['negativeEpisodeIds'],
                  'incrementalFeeWorseScenarioIds':pair['fees']['incremental']['negativeScenarioIds'],
                  'incrementalFeeWorseEpisodeIds':pair['fees']['incremental']['negativeEpisodeIds'],
                  'revisionWorseScenarioIds':revisions['scenarioIds'],'revisionWorseEpisodeIds':revisions['episodeIds']}
            lines += ['### '+panel+' / '+name,'','```json',json.dumps(sets,ensure_ascii=False,indent=2),'```','']
    return '\n'.join(lines)


def material_texts(e, slides):
    appendix=adverse_appendix(e);hashes=json.dumps(e['hashes'],ensure_ascii=False,indent=2)
    zh=['# PortPlan V14 正式研究：现有 V13 守卫策略的自然 D0 检验','',
        '这是既有 guardedHiGHSV13 的新研究批次，不是新 V14 排程算法。144 个新合成业务场景、两个独立报告面板、三次时钟重复、七方法共 3024 条；开发端点、旧前缀和旧样本不进入正式分母。','',
        '200ms 是构造、原生、审计和评分共享的合作预算。冷父时钟 5 秒包括启动、正常退出、ESRCH 进程组清理与最终父审核。及时 final 不足以接纳完整历史。未接纳或缺失历史保持 unknown/null，失败费用不填零。','',
        '默认观察资格仅为完整面板且对六个基线逐个零完成损失；净收益、费用、Pareto 不抵消损失。费用/改期的原生稳定性主张另列，不是默认完成门槛的一部分。','',
        '主比较六组、次比较八组事先固定。配对费用用该对照共同完成的匹配重复，先场景内平均再场景等权聚合，两个费用分子分母均披露。配对改期为同一共同完成集合中改期已知的匹配 episode 原始事件总数，另报自身已知条数分母；原生稳定性主张要求全部匹配条数的改期已知。只有七方法 Pareto 的改期轴采用场景内均值再等权聚合。时钟重复不作为独立 CI 样本。bootstrap 固定 seed800001、2000 次场景簇抽样；区间未经多重比较校正，是描述性边际区间，不证明同时优越；共同完成条件不消除选择偏差。全负节省区间为条件费用恶化，含零不证明稳定省费。','',
        'M 来自当天公开快照与实际锁，是明日必要仓承诺的代理量，不预测未公告事件，不保证未来可恢复。自然全路径与已见强制前缀开发分别保留；不同初始物理路径的成功不证明同前缀恢复。','',
        '公开费率、公开观测、数学基准与合成数据各自标注范围。企业订单、报价、联系人及试点为可选增强，不是本路线前提。不能据此声称企业采纳、实际收益或冠军。V1–V13 历史证据不改写。','']
    ko=['# PortPlan V14 연구 — 기존 V13 guard의 자연 D0 평가','',
        '새 V14 알고리즘이 아니라 기존 guardedHiGHSV13의 새 연구입니다. 새 합성 업무 144장면, 분리한 두 panel, 세 시계 반복과 일곱 방법으로 3024회를 보존합니다. 과거 개발 prefix와 endpoint import는 정식 표본에 포함하지 않습니다.','',
        '공통 200ms는 구성·native·감사·점수가 공유하는 협력 예산입니다. 부모 냉시작 5초는 정상 종료·ESRCH 전체 그룹 정리·최종 접수까지 포함합니다. timely final만으로는 충분하지 않습니다. 실패 비용과 미확보 장부는 null/unknown이며 0으로 채우지 않습니다.','',
        '기본 관측 자격은 완전 panel && 여섯 기준 각각 완료 손실 0입니다. 순이득·요금·Pareto가 손실을 상쇄하지 않습니다. native 비용·변경 안정성 주장은 별도 조건이며 기본 완료 gate를 바꾸지 않습니다.','',
        '주 비교6·부 비교8은 사전 고정했습니다. 쌍별 비용은 공동 완료 반복의 장면 내 평균을 장면 등가중 집계하고 두 비용 분자·분모를 제시합니다. 쌍별 변경은 같은 공동 완료 집합 중 값이 알려진 matching episode의 원이벤트 합과 자체 알려진 회수 분모입니다. native 안정성 주장은 모든 matching 변경 값이 알려져야 합니다. 일곱 방법 Pareto의 변경 축만 장면 내 평균 뒤 장면 등가중입니다. 시계 반복은 독립 CI 표본이 아닙니다. seed800001·2000회 장면-cluster bootstrap은 보정 없는 기술적 주변구간이며 동시 우월·familywise 확신이 아닙니다. 완료 선택 편향도 제거하지 않습니다. 전부 음수인 절감 구간은 조건부 악화이고 0 포함은 안정적 절감의 증거가 아닙니다.','',
        'M은 당일 공개 정보와 실행 잠금에서 계산한 필요 창고 대리량입니다. 미공지 미래·복구·강건성을 보장하지 않습니다. 개발 강제 prefix와 자연 전체 경로는 별도 증거입니다. 서로 다른 과거 경로의 완주는 같은 prefix 복구 증거가 아닙니다.','',
        '공개 요율·관측·수학 기준·합성 실험을 구분합니다. 기업 실단·연락처·견적·파일럿은 선택 확장이며 선행 조건이 아닙니다. 기업 실현 절감·채택·순위를 보장하지 않고 V1–V13 기록은 변경하지 않습니다.','']
    for panel in PANELS:
        gate=e['summary']['panels'][panel]['observedDecisionGates']
        zh += [PANELS_ZH[panel]+'观察默认资格：'+str(gate['defaultRecommendationEligible'])+'；各基线完成损失 '+json.dumps(gate['completionLossCountsByBaseline'],ensure_ascii=False),'']
        ko += [PANELS[panel]+' 관측 기본 자격: '+str(gate['defaultRecommendationEligible'])+'; 기준별 완료 손실 '+json.dumps(gate['completionLossCountsByBaseline'],ensure_ascii=False),'']
    zh += [detailed_results(e,'zh'),appendix,'## 证据 SHA / 原始源码身份','','```json',hashes,'```','',
           '材料为 AI 辅助草稿；本人资格、报名、当届 AI 规则、真实贡献、独立复现和答辩理解尚须本人确认，技术审核不替代这些事项。']
    ko += [detailed_results(e,'ko'),appendix,'## 원증거 SHA / source 신원','','```json',hashes,'```','',
           'AI 도움 초안입니다. 본인 자격·등록·당해 AI 규칙·실제 기여·독립 재현·발표 이해는 본인이 확인해야 하며 기술 검산으로 인증되지 않습니다.']
    en=['# PortPlan V14 study — natural-D0 assessment of the unchanged V13 guard','',
        '144 newly seeded synthetic business scenarios, two separate panels, three clock repetitions and seven methods: 3024 first-run episodes. No endpoint import or forced prefix. V14 identifies the study, not a new algorithm.','',
        'The completion-only observed default gate requires a complete panel and zero matching-repeat completion losses against each of six included baselines. Native fee/revision stability conditions are separate. Pairwise fees average matching complete repeats within scenario then weight scenarios equally. Pairwise revisions retain raw known matching-episode event sums and their own known denominator; clock repetitions are not independent CI samples. Only the all-seven Pareto revision axis uses equal-weight scenario means. Conditional scenario-cluster intervals have no multiplicity adjustment or simultaneous superiority claim. Failures and null costs remain visible.','',
        'A lower public necessary warehouse commitment proxy does not forecast hidden events or prove recovery. No real-business ROI, entrant verification or contest-rank guarantee.','']
    for panel in PANELS:
        s=e['summary']['panels'][panel]
        en += [panel+': V13 '+str(s['byMethod']['guardedHiGHSV13']['completeEpisodeCount'])+'/216; observed default eligibility '+str(s['observedDecisionGates']['defaultRecommendationEligible']), '']
    en += ['Full adverse identifiers and precise fee numerators, denominators and intervals are retained in report-v14.zh.md and submission-v14.ko.md.','```json',hashes,'```']
    contribution='# V14 material provenance\n\nExisting V13 algorithm; new preregistered formal assessment and internal independent Python reviews.\nAI assisted code, analysis and drafts. Entrant contribution/understanding remains unverified.\n\n```json\n'+hashes+'\n```\n'
    entrant='# 본인 확인 / 本人验收（空白）\n\n- 当届资格与报名证据：\n- 当届 AI 使用规则与披露：\n- 本人实际贡献：\n- 全新输出目录复现记录：\n- 解释当前模型、全部负例及局限：\n\n企业资料不是此公开研究路线的前提。此表不构成本人已通过验收。\n'
    return {'ko':'\n'.join(ko)+'\n','zh':'\n'.join(zh)+'\n','en':'\n'.join(en)+'\n','contribution':contribution,'entrant':entrant}


def validate_report_coverage(e, texts):
    for key in ('ko','zh'):
        report=texts[key]
        for panel in PANELS:
            for name,(baseline,candidate) in PAIRS.items():
                pair=e['summary']['panels'][panel]['comparisons'][name]
                identifiers=[*pair['completion']['lostEpisodeIds']]
                for kind in ('total','incremental'):
                    fee=pair['fees'][kind];identifiers += fee['negativeScenarioIds']+fee['negativeEpisodeIds']
                    require(ci(fee['uncertainty']['weightedSavingsPct95CI']) in report, 'Report drops precise signed fee interval')
                revisions=revision_adverses(e['index']['cases'],panel,baseline,candidate)
                identifiers += revisions['scenarioIds']+revisions['episodeIds']
                require(all(json.dumps(identifier,ensure_ascii=False) in report for identifier in identifiers), 'Report drops adverse ID')
                require('### '+panel+' / '+name in report,'Report drops comparison appendix')
    return True


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    action=parser.add_mutually_exclusive_group();action.add_argument('--check-only',action='store_true');action.add_argument('--generate',action='store_true')
    parser.add_argument('--schema-only',action='store_true');parser.add_argument('--root-authorized',action='store_true')
    parser.add_argument('--bootstrap-review',default='data/bootstrap-review-v14.json')
    parser.add_argument('--output-dir',default='docs')
    args=parser.parse_args(argv)
    if args.schema_only:
        require(not args.generate,'Schema-only cannot publish materials')
        print(json.dumps({'state':'schema-only-no-formal-result','writes':0,'candidateMethod':'guardedHiGHSV13','studyVersion':'v14',
                          'methods':METHODS,'primaryComparisons':list(PRIMARY_PAIRS),'secondaryComparisons':list(SECONDARY_PAIRS),
                          'plannedBusinessScenarios':144,'plannedEpisodes':3024,'maximumSlides':20}));return
    if args.generate: require(args.root_authorized,'Material generation requires separate root authorization after both formal reviews')
    evidence=verified_inputs(bootstrap_review=args.bootstrap_review);slides=build_slides(evidence);texts=material_texts(evidence,slides)
    validate_report_coverage(evidence,texts)
    if not args.generate:
        print(json.dumps({'valid':True,'checkOnly':True,'writes':0,'pages':len(slides),'evidenceHashes':evidence['hashes'],
                          'adverseIdentifierOccurrences':evidence['adverseIdentifierOccurrences']}));return
    output=(ROOT/args.output_dir).resolve();require(output.is_relative_to(ROOT) and output.is_dir(),'Output directory must be an existing project directory')
    targets=[output/name for name in OUTPUTS];audit_path=ROOT/MATERIAL_REVIEW
    require(all(not path.exists() for path in [*targets,audit_path]),'V14 outputs already exist; preserve first generated materials')
    pdf=pdf_document(slides);pptx=pptx_document(slides);audit=audit_rendered(pdf,pptx,slides)
    values=[html_document(slides).encode(),pdf,pptx,texts['ko'].encode(),texts['zh'].encode(),texts['en'].encode(),texts['contribution'].encode(),texts['entrant'].encode()]
    # Every input/binding/layout/text check completes before the first exclusive write.
    for path,contents in zip(targets,values):
        with path.open('xb') as stream:stream.write(contents)
    audit.update({'studyVersion':'v14','candidateMethod':'guardedHiGHSV13','sourceGitCommit':evidence['index']['sourceGitCommit'],
                  'inputIndexSha256':evidence['wrapper']['formalIndex']['sha256'],'evidenceHashes':evidence['hashes'],
                  'adverseIdentifierOccurrences':evidence['adverseIdentifierOccurrences'],'fullAdverseAppendixLanguages':['ko','zh'],
                  'outputHashes':{str(path.relative_to(ROOT)):sha(path) for path in targets},'rootAuthorized':True})
    with audit_path.open('x') as stream:json.dump(audit,stream,ensure_ascii=False,indent=2);stream.write('\n')
    print(json.dumps({'valid':True,'pages':len(slides),'written':[str(path.relative_to(ROOT)) for path in targets],'materialReview':MATERIAL_REVIEW},ensure_ascii=False))


if __name__=='__main__':main()


