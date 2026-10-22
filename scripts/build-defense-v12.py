#!/usr/bin/env python3
"""Generate result-driven Korean V12 materials only after complete reviewed evidence.

No solver, browser, experiment generation, network access or old-material writes.
The published wrapper, restored formal index and independent review must bind to
the same bytes. All validation and PDF layout finish in memory before any write.
"""
from __future__ import annotations

import argparse
from collections import Counter
from html import escape
from io import BytesIO
import hashlib
import json
import math
from pathlib import Path
import shlex

ROOT = Path(__file__).resolve().parents[1]
METHODS = ('costPolishV7', 'adaptiveV9', 'highsV8', 'stableHiGHSV10', 'commitmentHiGHSV11', 'seededHiGHSV12')
LABELS = dict(zip(METHODS, ('V7 비용 다듬기', 'V9 적시 후보', 'native HiGHS', 'V10 동일 요금 유지', 'V11 필요 창고 선택', 'V12 완전 초기해 + 창고')))
PANELS = {'same-distribution':'동일 생성분포 · 새 주문 시드','stress':'스트레스 · 새 주문 시드'}
PANELS_ZH = {'same-distribution':'同生成机制 · 新订单种子','stress':'压力扰动 · 新订单种子'}
PAIRS = {'v12VsV11':('commitmentHiGHSV11','seededHiGHSV12'), 'v12VsV10':('stableHiGHSV10','seededHiGHSV12'),
         'v12VsV9':('adaptiveV9','seededHiGHSV12'), 'v12VsHiGHS':('highsV8','seededHiGHSV12'),
         'v12VsV7':('costPolishV7','seededHiGHSV12'), 'v11VsV10':('stableHiGHSV10','commitmentHiGHSV11'),
         'v10VsHiGHS':('highsV8','stableHiGHSV10'), 'highsVsV7':('costPolishV7','highsV8')}
OUTPUTS = ('defense-v12.ko.html','defense-v12.ko.pdf','submission-v12.ko.md','defense-v12.en.md',
           'contribution-log-v12.md','defense-v12.ko.pptx','report-v12.zh.md',
           'optional-pilot-v12.blank.md','entrant-defense-v12.blank.md')


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def relative_path(root, value):
    require(isinstance(value, str), 'Evidence path must be a string')
    path = (root / value).resolve()
    require(path.is_relative_to(root.resolve()), 'Evidence path escapes project: ' + value)
    require(path.is_file(), 'Required completed evidence is missing: ' + value)
    return path


def load(root, name, hashes):
    path = relative_path(root, name)
    hashes[name] = sha(path)
    return json.loads(path.read_text())


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def close(actual, expected, label):
    require(finite(actual) and finite(expected) and abs(actual - expected) <= max(1e-5, abs(expected) * 1e-10),
            'Independent point estimate differs: ' + label)


def same_interval(actual, expected, label, tolerance):
    if expected is None:
        require(actual is None, 'Independent interval null differs: ' + label)
        return
    require(isinstance(actual, list) and isinstance(expected, list) and len(actual) == len(expected) == 2,
            'Malformed independent interval: ' + label)
    for a, b in zip(actual, expected):
        require(finite(a) and finite(b) and abs(a - b) <= tolerance,
                'Independent interval differs: ' + label)


def same_tree(actual, expected, label):
    if isinstance(expected, dict):
        require(isinstance(actual, dict) and set(actual) == set(expected), 'Independent tree keys differ: '+label)
        for key in expected: same_tree(actual[key], expected[key], label+'/'+key)
    elif isinstance(expected, list):
        require(isinstance(actual, list) and len(actual) == len(expected), 'Independent tree length differs: '+label)
        for i, (left, right) in enumerate(zip(actual, expected)): same_tree(left, right, label+'/'+str(i))
    elif finite(expected):
        require(finite(actual) and abs(actual-expected) <= 1e-6, 'Independent tree number differs: '+label)
    else:
        require(type(actual) is type(expected) and actual == expected, 'Independent tree value differs: '+label)


def fnum(value, digits=0):
    if not finite(value):
        return 'null'
    return f'{value:,.{digits}f}'


def pct(value):
    return fnum(value, 2) + '%' if finite(value) else 'null'


def ci(value):
    # Keep the raw sign of small endpoints; rounding must not turn a strictly
    # positive/negative bound into an apparent zero for the declared gates.
    return '[' + ', '.join(format(v, '.10g') if finite(v) else 'null' for v in value) + ']' if isinstance(value, list) and len(value) == 2 else 'null'


def validate_adverse_indices(cases, summary):
    """Derive every negative/lost ID from the complete original scalar receipts."""
    occurrences = 0
    for panel in PANELS:
        selected = [case for case in cases if case['panel'] == panel]
        for name, (baseline, candidate) in PAIRS.items():
            pair = summary['panels'][panel]['comparisons'][name]
            gains, losses = [], []
            for case in selected:
                for repeat in case['repeats']:
                    before, after = repeat['policies'][baseline], repeat['policies'][candidate]
                    identifier = case['id'] + '::repeat-' + str(repeat['repeatIndex'])
                    if before['complete'] and not after['complete']: losses.append(identifier)
                    if after['complete'] and not before['complete']: gains.append(identifier)
            require(pair['completion'].get('lostEpisodeIds') == losses and pair['completion'].get('gainedEpisodeIds') == gains,
                    'Missing or changed complete gain/loss identifiers: '+panel+'/'+name)
            require(pair['completion']['lostEpisodeCount'] == len(losses) and pair['completion']['gainedEpisodeCount'] == len(gains),
                    'Completion loss/gain counts disagree with receipts')
            occurrences += len(losses)
            for key, field in (('total', 'totalCost'), ('incremental', 'incrementalCost')):
                negative_scenarios, negative_episodes = [], []
                for case in selected:
                    differences = []
                    for repeat in case['repeats']:
                        before, after = repeat['policies'][baseline], repeat['policies'][candidate]
                        if before['complete'] and after['complete']:
                            require(finite(before[field]) and finite(after[field]), 'Completed receipt fee is missing')
                            difference = after[field]-before[field]; differences.append(difference)
                            if difference > 1e-6: negative_episodes.append(case['id']+'::repeat-'+str(repeat['repeatIndex']))
                    if differences and sum(differences)/len(differences) > 1e-6: negative_scenarios.append(case['id'])
                fee = pair['fees'][key]
                require(fee.get('negativeScenarioIds') == negative_scenarios and fee.get('negativeEpisodeIds') == negative_episodes,
                        'Missing or changed all negative identifiers: '+panel+'/'+name+'/'+key)
                require(fee['scenarioWinTieLoss']['loss'] == len(negative_scenarios), 'Negative scenario denominator differs')
                occurrences += len(negative_scenarios)+len(negative_episodes)
    return occurrences


def validate_parent_lifecycle(result):
    require(result.get('endToEndCompletedWithinBudget') is True and result.get('episodeCompletedWithinBudget') is True and
            result.get('terminalReceiptAcceptedAsHistory') is True and result.get('normalWorkerExit') is True and
            result.get('exitCode') == 0 and result.get('exitSignal') is None and result.get('deadlineReached') is False,
            'Completed episode lacks normal accepted parent lifecycle')
    cleanup=result.get('processGroupCleanup',{})
    require(cleanup.get('successful') is True and cleanup.get('verifiedProcessGroupGone') is True and
            cleanup.get('verificationResult') == 'ESRCH' and cleanup.get('verificationError') is None,
            'Completed episode lacks verified process-group cleanup')
    for field in ('wallTimeMs','endToEndWallTimeMs','processExitWallTimeMs','processCloseWallTimeMs','cleanupFinishedMs','receiptValidatedMs'):
        require(finite(result.get(field)) and 0<=result[field]<=5000,'Completed lifecycle exceeds 5s or lacks '+field)


def verified_seed_development(root, hashes):
    """Bind the corrected public wrapper and independent review to raw summary bytes.

    Post-first-formal-run presentation-reader correction: the review's
    inputSummarySha256 names the original development summary, not the later
    public scope-correction wrapper. Frozen evidence and its hashes stay intact.
    """
    review_name = 'data/development-seed-review-v12.json'
    development = load(root, review_name, hashes)
    require(development.get('schemaVersion') == 'portplan-independent-v12-seed-development-review' and
            development.get('valid') is True and development.get('developmentOnly') is True and development.get('readOnly') is True and
            development.get('importsJavaScript') is False and development.get('solverRuns') == 0 and
            development.get('newHeldOutExperiments') is False and development.get('fullRollingCompletionClaim') is False,
            'D0 development is not separate audited development')
    published_name = 'data/seeded-initial-development-v12.json'
    published = load(root, published_name, hashes)
    require(published.get('schemaVersion') == 'portplan-seeded-initial-development-v12' and published.get('valid') is True and
            published.get('dataRole') == 'predeclared-seen-V10-subset' and published.get('newHeldOutExperiments') == 0,
            'D0 public wrapper is not the corrected seen-input development scope')
    review_reference = published.get('independentReview', {})
    require(review_reference.get('path') == review_name and review_reference.get('sha256') == hashes[review_name],
            'D0 public wrapper binds to different independent review bytes')
    reference = published.get('originalSummary', {})
    original_name = reference.get('path')
    require(original_name == 'artifacts/seeded-v12-d0-first-development/summary.json', 'Unexpected D0 original summary reference')
    raw_name = original_name
    if not (root/original_name).is_file():
        restore_hint = ('D0 raw development summary missing; restore it with: python scripts/restore-evidence-v12.py '
                        '--manifest data/archives-v12/seed-development/manifest.json --destination data/evidence-v12')
        manifest_name = 'data/archives-v12/seed-development/manifest.json'
        raw_name = 'data/evidence-v12/v12-seed-development/'+original_name
        require((root/manifest_name).is_file() and (root/raw_name).is_file(), restore_hint)
        manifest = load(root, manifest_name, hashes)
        require(manifest.get('schemaVersion') == 'portplan-v12-lossless-tar-chunks' and manifest.get('kind') == 'development' and
                manifest.get('archiveFormat') == 'tar.xz' and manifest.get('compression') == 'xz-lzma2-preset6' and
                manifest.get('archiveRoot') == 'v12-seed-development' and manifest.get('entrypoint') == published_name and
                manifest.get('entrypointSha256') == hashes[published_name], 'D0 development archive/wrapper binding differs')
        files = manifest.get('files', [])
        require(isinstance(files, list) and all(isinstance(item, dict) and isinstance(item.get('path'), str) for item in files),
                'Malformed D0 development archive file receipts')
        entries = {item['path']:item for item in files}
        require(len(entries) == len(files) == manifest.get('fileCount'), 'Duplicate or incomplete D0 development archive file receipts')
        for name, expected in ((original_name, reference.get('sha256')), (review_name, hashes[review_name]), (published_name, hashes[published_name])):
            record = entries.get(name, {})
            actual_path = relative_path(root, raw_name if name == original_name else name)
            require(record.get('sha256') == expected == sha(actual_path) and record.get('bytes') == actual_path.stat().st_size,
                    'D0 development archive file binding differs: '+name)
        chunks = manifest.get('chunks', [])
        require(isinstance(chunks, list) and chunks, 'D0 development archive chunks missing')
        archive_digest = hashlib.sha256(); archive_bytes = 0
        for item in chunks:
            name = item.get('file') if isinstance(item, dict) else None
            require(isinstance(name, str) and Path(name).name == name and name not in ('.','..'), 'Unsafe D0 development chunk name')
            chunk_name = str(Path(manifest_name).parent/name); path = relative_path(root, chunk_name)
            require(path.stat().st_size == item.get('bytes') and sha(path) == item.get('sha256'), 'D0 development archive chunk SHA differs')
            hashes[chunk_name] = item['sha256']
            with path.open('rb') as stream:
                for block in iter(lambda:stream.read(1024*1024), b''): archive_digest.update(block);archive_bytes += len(block)
        require(archive_bytes == manifest.get('archiveBytes') and archive_digest.hexdigest() == manifest.get('archiveSha256'),
                'Combined D0 development archive SHA differs')
    raw = load(root, raw_name, hashes)
    require(hashes[raw_name] == reference.get('sha256') == development.get('inputSummarySha256'),
            'D0 original summary SHA differs from public wrapper or independent review')
    require(raw.get('schemaVersion') == 'portplan-v12-D0-development-summary-1' and raw.get('developmentOnly') is True and
            raw.get('newHeldOutExperiments') is False and raw.get('fullRollingCompletionClaim') is False,
            'D0 original summary is not bounded development evidence')
    for reviewed, original, public in (('endpointCount','endpointCount','endpoints'),
                                      ('reusedBusinessScenarioCount','independentBusinessScenarioCount','independentBusinessScenarios'),
                                      ('timelyCompleteSeedCount','timelyCompleteCandidateCount','timelyCompleteConstructedSeeds')):
        require(development['summary'][reviewed] == raw.get(original) == published.get(public),
                'D0 public/original/reviewed development count differs: '+reviewed)
    return development


def verified_inputs(root=ROOT, bootstrap_review='data/bootstrap-review-v12.json'):
    hashes = {}
    release = load(root, 'data/v12-release.json', hashes)
    require(all(release.get(key) is True for key in ('formalExperimentComplete','independentReviewComplete','independentBootstrapReviewComplete')),
            'V12 formal evidence and both reviews must complete before materials')
    wrapper = load(root, 'data/v12-results.json', hashes)
    require(wrapper.get('schemaVersion') == 'portplan-v12-published-results' and wrapper.get('executionState') == 'complete' and
            not wrapper.get('fixture'), 'V12 results are not completed formal evidence')
    reference = wrapper.get('formalIndex', {})
    formal_path = relative_path(root, reference.get('path'))
    require(sha(formal_path) == reference.get('sha256'), 'Formal index SHA differs')
    hashes[reference['path']] = reference['sha256']; index = json.loads(formal_path.read_text())
    require(index.get('schemaVersion') == 'portplan-v12-streamed-experiment' and index.get('executionState') == 'complete' and
            not index.get('fixture'), 'Restored formal run is incomplete or a fixture')
    require(wrapper.get('summary') == index.get('summary') and wrapper.get('protocol') == index.get('protocol'), 'Published summary/protocol differ from formal bytes')
    protocol = load(root, 'data/protocol-v12.json', hashes)
    require(protocol == index['protocol'] and hashes['data/protocol-v12.json'] == index.get('protocolHash'), 'Protocol bytes differ')
    require(tuple(protocol.get('methods', [])) == METHODS and protocol['budget']['hardWallMs'] == 5000, 'Unexpected methods or hard episode clock')
    require(protocol['budget'].get('parentAcceptance') == {'normalWorkerExitRequired':True,'confirmedProcessGroupGoneRequired':True,
            'processGroupVerificationResult':'ESRCH','cleanupWithinHardWallRequired':True,'endToEndWithinHardWallRequired':True,
            'timelyFinalAloneIsCompletion':False}, 'Missing complete parent lifecycle budget declaration')
    require(all(protocol['callBudgets'][method]['timeBudgetMs'] == 200 for method in METHODS), 'Unexpected cooperative call budget')
    require(protocol.get('executionAuthorization', {}).get('status') in ('ready-for-source-freeze','frozen'), 'Draft development declaration cannot generate formal materials')
    frozen_dir = (formal_path.parent/index.get('frozenSource','..')).resolve()
    manifest_path = relative_path(root, str((frozen_dir/'manifest.json').relative_to(root.resolve())))
    manifest_hash = sha(manifest_path); hashes[str(manifest_path.relative_to(root))] = manifest_hash
    review = load(root, 'data/review-v12.json', hashes); statistics = load(root, bootstrap_review, hashes)
    for value, schema in ((review,'portplan-independent-v12-run-review'),(statistics,'portplan-independent-v12-bootstrap-review')):
        require(value.get('schemaVersion') == schema and value.get('valid') is True and value.get('fixture',False) is False and
                value.get('readOnly') is True and value.get('solverRuns') == 0 and value.get('importsJavaScript') is False,
                'Incomplete, fixture or non-independent review')
        require(value.get('inputIndexSha256') == reference['sha256'] and value.get('frozenManifestSha256') == manifest_hash,
                'Independent review belongs to different result/frozen bytes')
    cases = index.get('cases', []); repeats = protocol.get('algorithmRepeats', [])
    require(len(cases) == len({case['id'] for case in cases}) == 144 and Counter(case['panel'] for case in cases) == Counter({'same-distribution':72,'stress':72}), 'Missing/duplicate business scenarios')
    require(len(repeats) == len({repeat['seed'] for repeat in repeats}) == 3, 'Missing fixed timing repeats')
    reviewed = {item['id']:item for item in review.get('episodes', [])}; count = 0; consumed=set()
    require(len(reviewed) == 2592, 'Physical review lacks unique episode coverage')
    for case in cases:
        artifact = case['inputArtifact']; path = relative_path(root, str((formal_path.parent/artifact['path']).relative_to(root)))
        require(sha(path) == artifact['sha256'], 'Scenario input SHA differs')
        require(len(case['repeats']) == 3, 'Missing repeat')
        for number, repeat in enumerate(case['repeats'],1):
            require(repeat['repeatIndex'] == number and repeat['algorithmSeed'] == repeats[number-1]['seed'] and set(repeat['policies']) == set(METHODS), 'Missing/reordered repeat/method')
            for method, receipt in repeat['policies'].items():
                ref = receipt['artifact']; path = relative_path(root, str((formal_path.parent/ref['path']).relative_to(root)))
                require(sha(path) == ref['sha256'], 'Episode artifact SHA differs')
                envelope = json.loads(path.read_text()); actual = envelope['result']; audited = reviewed.get(envelope['episodeId'])
                require(audited and audited['artifactSha256'] == ref['sha256'] and envelope['caseId'] == case['id'] and envelope['method'] == method, 'Episode/review identity binding differs')
                require(envelope.get('repeatIndex') == number and envelope.get('algorithmSeed') == repeats[number-1]['seed'] and
                        envelope['episodeId'] not in consumed,'Episode repeat/seed is mismatched or duplicated')
                consumed.add(envelope['episodeId'])
                require(type(receipt.get('complete')) is bool and receipt['complete'] is actual['complete'], 'Receipt completion differs')
                if receipt['complete']:
                    validate_parent_lifecycle(receipt);validate_parent_lifecycle(actual)
                for field in ('totalCost','incrementalCost'):
                    if receipt['complete']: close(receipt[field],actual[field],'receipt/'+field)
                    else: require(receipt[field] is None and actual[field] is None, 'Failed full-path fees must remain null')
                count += 1
    require(count == len(consumed) == 2592 and wrapper.get('scenarioCount') == 144 and wrapper.get('episodeCount') == count and
            review.get('summary',{}).get('episodes') == count and statistics.get('episodeCount') == count and statistics.get('businessScenarioCount') == 144,
            'Formal/independent episode coverage differs')
    summary = index['summary']; require(summary.get('inferentialState') == 'complete-scenario-cluster-bootstrap' and set(summary['panels']) == set(PANELS), 'Missing final scenario intervals')
    declared = protocol['analysis']['bootstrap']; procedure = statistics.get('procedure',{})
    require(procedure.get('samplingUnit') == 'whole business scenario' and procedure.get('seed') == declared['seed'] and procedure.get('draws') == declared['draws'], 'Bootstrap unit/seed/draws differ')
    for panel in PANELS:
        produced, physical, statistical = summary['panels'][panel], review['pointEstimates'][panel], statistics['panels'][panel]
        require(produced['scenarioCount'] == 72 and set(produced['byMethod']) == set(METHODS) and set(produced['comparisons']) == set(PAIRS), 'Unexpected panel/method/comparison coverage')
        for method in METHODS:
            for field in ('recordedEpisodeCount','completeEpisodeCount','fullyCompletedScenarioCount','anyCompletedScenarioCount'):
                require(produced['byMethod'][method][field] == physical['byMethod'][method][field] == statistical['byMethod'][method][field], 'Reviewed completion differs')
            require(produced['byMethod'][method]['recordedEpisodeCount'] == 216, 'Missing full-sample method receipts')
            same_tree(produced['byMethod'][method]['initialization'],physical['byMethod'][method]['initialization'],'D0 physical')
            same_tree(produced['byMethod'][method]['initialization'],statistical['byMethod'][method]['initialization'],'D0 statistical')
        for name,(baseline,candidate) in PAIRS.items():
            pair, point, boot = produced['comparisons'][name], physical['comparisons'][name], statistical['comparisons'][name]
            require((pair['baselineMethod'],pair['candidateMethod']) == (baseline,candidate), 'Comparison identity differs')
            require(pair['completion']['gainedEpisodeCount'] == point['gainedEpisodes'] and pair['completion']['lostEpisodeCount'] == point['lostEpisodes'], 'Completion losses/gains differ')
            same_interval(pair['completion']['uncertainty']['meanScenarioRateGain95CI'],boot['completionRateGain95CI'],'completion CI',1e-12)
            require([pair['revisions']['baselineObservedSum'],pair['revisions']['candidateObservedSum']] == point['revisions'], 'Revision denominator/numerator differs')
            for key in ('total','incremental'):
                fee, checked, sampled = pair['fees'][key], point['fees'][key], boot['fees'][key]
                require(fee['pairedScenarioCount'] == checked['pairedScenarios'] == sampled['pairedScenarioCount'] and fee['pairedEpisodeCount'] == checked['pairedEpisodes'] == sampled['pairedEpisodeCount'], 'Conditional denominators differ')
                require(fee['scenarioWinTieLoss'] == checked['scenarioWinTieLoss'], 'Fee win/tie/loss differs')
                for field in ('baselineScenarioMeanSumKRW','candidateScenarioMeanSumKRW'): close(fee[field],checked[field],field)
                uncertainty = fee['uncertainty']; require(uncertainty['unit'] == 'scenario' and uncertainty['scenarioCount'] == fee['pairedScenarioCount'] and uncertainty['draws'] == declared['draws'], 'CI eligibility differs')
                same_interval(uncertainty['weightedSavingsPct95CI'],sampled['weightedSavingsPct95CI'],'fee CI',1e-9)
                same_interval(uncertainty['meanSavings95CI'],sampled['meanSavings95CI'],'KRW CI',1e-6)
        same_tree(produced['jointCompletePareto'],statistical['jointCompletePareto'],'all-six common set')
        gates = {key:value for key,value in produced['observedDecisionGates'].items() if key!='nativeComparison'}
        same_tree(gates,statistical['observedDecisionGates'],'observed necessary gates')
        require(gates['assessmentState'] == 'complete-panel', 'Partial-panel gates cannot publish')
    adverse_count = validate_adverse_indices(cases, summary)
    development = None
    if (root/'data/development-seed-review-v12.json').exists():
        development = verified_seed_development(root, hashes)
    diagnostic=None
    if (root/'data/native-diagnostic-review-v12-final.json').exists():
        diagnostic=load(root,'data/native-diagnostic-review-v12-final.json',hashes)
        require(diagnostic.get('schemaVersion') == 'portplan-independent-native-diagnostic-review-v12' and diagnostic.get('valid') is True and
                diagnostic.get('readOnly') is True and diagnostic.get('solverRuns') == 0 and diagnostic.get('snapshotCount') == 9 and
                diagnostic.get('taskCount') == 36 and diagnostic.get('newHeldOutScenarios') == 0, 'CP-SAT is not independently reviewed old-input diagnosis')
        original=relative_path(root,'data/native-diagnostic-v12.json');hashes['data/native-diagnostic-v12.json']=sha(original)
        require(diagnostic.get('publishedSummarySha256') == hashes['data/native-diagnostic-v12.json'],'Independent CP diagnostic review binds to different summary bytes')
    for name in ('data/public-source-field-coverage-v12.json','docs/public-source-field-coverage-v12.zh.md'):
        if (root/name).exists():hashes[name]=sha(relative_path(root,name))
    return {'root':root,'wrapper':wrapper,'index':index,'protocol':protocol,'summary':summary,'review':review,'bootstrapReview':statistics,
            'release':release,'development':development,'diagnostic':diagnostic,'hashes':hashes,'adverseIdentifierOccurrences':adverse_count}



def presentation_table(source):
    """Short CI display for slides only; reports retain the source table precision."""
    if not source['table']:
        return None
    headers, rows = source['table']
    if '조건부 95% CI' not in headers:
        return headers, rows
    column = headers.index('조건부 95% CI')
    def endpoint(value):
        if not finite(value): return 'null'
        if value == 0: return '0'
        # Do not round a tiny strictly positive/negative endpoint into zero.
        return format(value, '.2g') if abs(value) < 0.00005 else f'{value:.4f}'
    displayed=[]
    for row in rows:
        cells=list(row); value=json.loads(cells[column])
        cells[column]='['+', '.join(endpoint(v) for v in value)+']' if isinstance(value,list) else 'null'
        displayed.append(cells)
    return headers, displayed


def presentation_note(source):
    return source['note']+(' CI는 약 소수 4자리로 표시하고 작은 값의 부호는 보존합니다. 전체 정밀도는 보고서·JSON을 참조합니다.'
                           if source['table'] and '조건부 95% CI' in source['table'][0] else '')


def presentation_table_layout(headers, rows):
    dense = len(rows)>7 or len(headers)>=5
    size = 16 if len(rows)>7 else 18
    widths = [1156/len(headers)]*len(headers)
    if len(headers)==2: widths=[550,606]
    elif len(headers)==6:
        widths=[255,145,145,120,250,241] if '조건부 95% CI' in headers else [260,179.2,179.2,179.2,179.2,179.2]
    elif len(headers)==5: widths=[180,280,140,270,286]
    return size, size*1.35, 4 if dense else 9, widths


def html_document(slides):
    parts=['<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PortPlan V12</title><style>',
        'body{margin:0;background:#dce8e0;color:#173c40;font-family:"Noto Sans CJK KR","Noto Sans KR",sans-serif}.slide{box-sizing:border-box;width:1280px;min-height:720px;margin:24px auto;padding:44px 62px;background:#f7faf5;position:relative;page-break-after:always}h1{font-size:35px;line-height:1.4;margin:10px 0 22px;white-space:pre-line}.kicker{font-size:13px;letter-spacing:3px;color:#39826d}p{font-size:20px;line-height:1.7}table{width:100%;border-collapse:collapse;font-size:15px;margin:22px 0}th,td{text-align:left;border-bottom:1px solid #d3e0d3;padding:12px 10px;overflow-wrap:anywhere}th{background:#e6f0e4}.note{font-size:14px;background:#e8f2e5;padding:14px 18px;line-height:1.8}.flow{display:flex;gap:24px;align-items:stretch;margin:38px 0}.flow div{flex:1;background:#e0efe0;border-left:5px solid #36836c;padding:24px;font-size:23px;white-space:pre-line}.chart{margin:22px 0}.row{display:grid;grid-template-columns:240px 1fr 250px;align-items:center;gap:16px;margin:19px 0;font-size:17px}.track{height:21px;background:#dfebde}.fill{height:100%;background:#35826d}.footer{font-size:12px;letter-spacing:.5px;margin-top:25px;color:#638176}@media(max-width:1300px){.slide{width:100%;min-height:0;margin:12px auto;padding:30px}h1{font-size:29px}.row{grid-template-columns:180px 1fr 190px}}@media print{@page{size:1280px 720px;margin:0}body{background:white}.slide{width:1280px;height:720px;margin:0;overflow:hidden}}',
        '</style><body>']
    for i,slide in enumerate(slides,1):
        parts.extend([f'<section class="slide" id="slide-{i}"><div class="kicker">{i:02d} / PORTPLAN V12</div>',
                      '<h1>'+escape(slide['title'])+'</h1><p>'+escape(slide['text'])+'</p>'])
        if slide['flow']:parts.append('<div class="flow">'+''.join('<div>'+escape(text)+'</div>' for text in slide['flow'])+'</div>')
        if slide['table']:
            headers,rows=presentation_table(slide);size,_,_,_=presentation_table_layout(headers,rows);parts.append(f'<table style="font-size:{size}px"><thead><tr>'+''.join('<th>'+escape(str(h))+'</th>' for h in headers)+'</tr></thead><tbody>')
            parts.extend('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows);parts.append('</tbody></table>')
        if slide['chart']:
            parts.append('<div class="chart">')
            for c in slide['chart']:
                width=100*c['value']/c['denominator'] if c['denominator'] else 0
                parts.append(f'<div class="row"><span>{escape(c["label"])}</span><div class="track"><div class="fill" style="width:{width:.4f}%"></div></div><span>{c["value"]}/{c["denominator"]} · 3/3 {c["all3"]} · 미확보 {c["unknown"]}</span></div>')
            parts.append('</div>')
        parts.extend(['<div class="note">'+escape(presentation_note(slide))+'</div>',f'<div class="footer">AI 작성 준비 자료 · 본인 검토·현장 검증 미확인 / {i:02d} / {len(slides):02d}</div></section>'])
    parts.append('</body></html>');return '\n'.join(parts)+'\n'


def register_noto(text):
    """Embed Noto KR; convert only the used CFF glyphs to quadratic outlines in memory."""
    from fontTools import subset
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.cu2quPen import Cu2QuPen
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as PDFont
    path=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    require(path.is_file(), 'Noto Sans CJK font is required for a readable Korean PDF')
    font=TTFont(path,fontNumber=1)
    require('KR' in font['name'].getDebugName(1), 'Expected Noto Korean font face')
    chars=set(ord(c) for c in text if not c.isspace());cmap=font.getBestCmap()
    require(not (chars-set(cmap)), 'Noto is missing required PDF glyphs: '+repr(chars-set(cmap)))
    subsetter=subset.Subsetter();subsetter.populate(text=text);subsetter.subset(font)
    order=font.getGlyphOrder();glyphset=font.getGlyphSet();glyf={}
    for name in order:
        pen=TTGlyphPen(glyphset);glyphset[name].draw(Cu2QuPen(pen,1.0,reverse_direction=True));glyf[name]=pen.glyph()
    builder=FontBuilder(font['head'].unitsPerEm,isTTF=True);builder.setupGlyphOrder(order);builder.setupCharacterMap(font.getBestCmap());builder.setupGlyf(glyf)
    builder.setupHorizontalMetrics(font['hmtx'].metrics);builder.setupHorizontalHeader(ascent=font['hhea'].ascent,descent=font['hhea'].descent)
    builder.setupNameTable({'familyName':'Noto Sans CJK KR PortPlan subset','styleName':'Regular','uniqueFontIdentifier':'PortPlanNotoKR','fullName':'PortPlan Noto KR subset','psName':'PortPlanNotoKR',
                          'copyright':'Noto Sans CJK. Copyright Adobe and Google. SIL Open Font License 1.1. Subset outlines converted for embedding.'})
    builder.setupOS2(sTypoAscender=font['hhea'].ascent,sTypoDescender=font['hhea'].descent,usWinAscent=font['hhea'].ascent,usWinDescent=-font['hhea'].descent)
    builder.setupPost();builder.setupMaxp();buffer=BytesIO();builder.save(buffer);buffer.seek(0);pdfmetrics.registerFont(PDFont('PortPlanNoto',buffer))


def pdf_document(slides):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle
    all_text=json.dumps(slides,ensure_ascii=False)+' '.join(presentation_note(slide) for slide in slides)+' PORTPLAN V12 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789'
    register_noto(all_text)
    buffer=BytesIO();c=canvas.Canvas(buffer,pagesize=(1280,720),invariant=1,pageCompression=1)
    c.setTitle('PortPlan V12 · Korean competition preparation');c.setAuthor('AI-assisted PortPlan research; entrant verification pending')
    bg=HexColor('#f7faf5');ink=HexColor('#173c40');green=HexColor('#35826d');light=HexColor('#e6f0e4')
    def paragraph(text,x,y,w,size=20,leading=None):
        style=ParagraphStyle('p',fontName='PortPlanNoto',fontSize=size,leading=leading or size*1.55,textColor=ink,wordWrap='CJK')
        obj=Paragraph(escape(str(text)).replace('\n','<br/>'),style);_,h=obj.wrap(w,700)
        obj.drawOn(c,x,y-h);return y-h
    for index,slide in enumerate(slides,1):
        c.setFillColor(bg);c.rect(0,0,1280,720,stroke=0,fill=1)
        c.setFillColor(green);c.setFont('PortPlanNoto',13);c.drawString(62,672,f'{index:02d} / PORTPLAN V12')
        y=paragraph(slide['title'],62,642,1156,33,43)-15
        y=paragraph(slide['text'],62,y,1156,18,27)-19
        if slide['flow']:
            w=(1156-44)/3;h=118
            for j,text in enumerate(slide['flow']):
                x=62+j*(w+22);c.setFillColor(light);c.roundRect(x,y-h,w,h,7,stroke=0,fill=1);paragraph(text,x+18,y-22,w-36,23,34)
                if j<2:c.setStrokeColor(green);c.setLineWidth(2);c.line(x+w+4,y-h/2,x+w+18,y-h/2);c.line(x+w+13,y-h/2+5,x+w+18,y-h/2);c.line(x+w+13,y-h/2-5,x+w+18,y-h/2)
            y-=h+24
        if slide['chart']:
            for row in slide['chart']:
                paragraph(row['label'],62,y-5,252,17,23)
                c.setFillColor(light);c.rect(322,y-24,580,20,stroke=0,fill=1)
                c.setFillColor(green);c.rect(322,y-24,580*row['value']/row['denominator'] if row['denominator'] else 0,20,stroke=0,fill=1)
                paragraph(f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y-4,294,15,22);y-=53
            y-=8
        if slide['table']:
            headers,rows=presentation_table(slide);font_size,leading,padding,widths=presentation_table_layout(headers,rows)
            style=ParagraphStyle('cell',fontName='PortPlanNoto',fontSize=font_size,leading=leading,textColor=ink,wordWrap='CJK')
            data=[[Paragraph(escape(str(v)),style) for v in row] for row in [headers,*rows]]
            table=Table(data,colWidths=widths,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),light),('VALIGN',(0,0),(-1,-1),'TOP'),
                ('BOTTOMPADDING',(0,0),(-1,-1),padding),('TOPPADDING',(0,0),(-1,-1),padding),('LEFTPADDING',(0,0),(-1,-1),10),
                ('LINEBELOW',(0,0),(-1,-1),0.5,HexColor('#d3e0d3'))]))
            _,h=table.wrap(1156,600);require(y-h>100,f'PDF table overflows page {index}: {y-h:.1f}');table.drawOn(c,62,y-h);y-=h+18
        note_style=ParagraphStyle('note',fontName='PortPlanNoto',fontSize=12,leading=19,textColor=ink,wordWrap='CJK')
        note=Paragraph(escape(presentation_note(slide)),note_style);_,h=note.wrap(1116,300)
        require(y-h-28>64,f'PDF note overflows page {index}: {y-h-28:.1f}')
        c.setFillColor(light);c.roundRect(62,y-h-28,1156,h+25,5,stroke=0,fill=1);note.drawOn(c,82,y-h-14)
        c.setFillColor(HexColor('#638176'));c.setFont('PortPlanNoto',10)
        c.drawString(62,32,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인');c.drawRightString(1218,32,f'PORTPLAN V12 / {index:02d} / {len(slides):02d}');c.showPage()
    c.save();return buffer.getvalue()


def pptx_document(slides):
    """Editable native text, tables and vector shapes; no raster page pictures."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from reportlab.pdfbase import pdfmetrics

    require(1 <= len(slides) <= 20, 'Editable deck must contain at most 20 slides')
    # pdf_document normally registers the same outline metrics first; direct
    # renderer use remains possible without rendering a PDF or writing a font.
    if 'PortPlanNoto' not in pdfmetrics.getRegisteredFontNames():
        register_noto(json.dumps(slides,ensure_ascii=False)+' '.join(presentation_note(slide) for slide in slides)+' PORTPLAN V12 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789')
    presentation=Presentation();presentation.slide_width=Inches(1280/96);presentation.slide_height=Inches(720/96)
    presentation.core_properties.title='PortPlan V12 · Korean competition preparation'
    presentation.core_properties.author='AI-assisted PortPlan research; entrant verification pending'
    presentation.core_properties.subject='Editable result-driven research deck; synthetic daily model'
    colors={'background':'F7FAF5','ink':'173C40','green':'35826D','light':'E6F0E4','muted':'638176'}
    unit=lambda px: Inches(px/96)

    def wrapped(value,width,size):
        lines=[]
        for raw in str(value).split('\n'):
            line='';used=0
            for char in raw:
                advance=pdfmetrics.stringWidth(char,'PortPlanNoto',size)
                if line and used+advance>width:
                    lines.append(line.rstrip());line='';used=0
                line+=char;used+=advance
            lines.append(line.rstrip())
        return '\n'.join(lines),len(lines)

    def style_text(frame,text,size,leading,color='ink'):
        frame.clear();frame.text=text;frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        frame.vertical_anchor=MSO_ANCHOR.TOP
        for paragraph in frame.paragraphs:
            paragraph.font.name='Noto Sans CJK KR';paragraph.font.size=Pt(size*.75)
            paragraph.font.color.rgb=RGBColor.from_string(colors[color])
            paragraph.line_spacing=Pt(leading*.75);paragraph.space_before=Pt(0);paragraph.space_after=Pt(0)
            properties=paragraph._p.get_or_add_pPr().get_or_add_defRPr()
            east_asian=properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
            if east_asian is None:
                east_asian=OxmlElement('a:ea');properties.append(east_asian)
            east_asian.set('typeface','Noto Sans CJK KR')

    def text(slide,value,x,y,width,size=18,leading=27,color='ink',height=None,name='Editable text'):
        content,line_count=wrapped(value,width-2,size)
        measured=line_count*leading+3
        box=slide.shapes.add_textbox(unit(x),unit(y),unit(width),unit(height or measured));box.name=name
        style_text(box.text_frame,content,size,leading,color)
        return y+(height or measured)

    def rect(slide,x,y,width,height,color='light',kind=MSO_SHAPE.RECTANGLE,name='Editable vector'):
        shape=slide.shapes.add_shape(kind,unit(x),unit(y),unit(width),unit(height));shape.name=name
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(colors[color]);shape.line.fill.background()
        return shape

    for index,source in enumerate(slides,1):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string(colors['background'])
        text(slide,f'{index:02d} / PORTPLAN V12',62,31,1156,13,20,'green',name='Slide index')
        y=text(slide,source['title'],62,72,1156,33,43,name='Slide title')+15
        y=text(slide,source['text'],62,y,1156,18,27,name='Main explanation')+19
        if source['flow']:
            width=(1156-44)/3
            for j,label in enumerate(source['flow']):
                x=62+j*(width+22);rect(slide,x,y,width,118,name='Flow node '+str(j+1))
                rect(slide,x,y,5,118,'green',name='Flow node accent')
                text(slide,label,x+18,y+22,width-36,23,34,height=80,name='Editable flow label')
                if j<2:rect(slide,x+width+4,y+52,14,14,'green',MSO_SHAPE.CHEVRON,'Editable flow arrow')
            y+=142
        if source['chart']:
            for row in source['chart']:
                text(slide,row['label'],62,y+5,252,17,23,name='Chart method label')
                rect(slide,322,y+4,580,20,name='Editable completion bar track')
                width=580*row['value']/row['denominator'] if row['denominator'] else 0
                if width:rect(slide,322,y+4,width,20,'green',name='Editable completion bar')
                text(slide,f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y+4,294,15,22,name='Chart exact value')
                y+=53
            y+=8
        if source['table']:
            headers,rows=presentation_table(source);size,leading,padding,widths=presentation_table_layout(headers,rows)
            formatted=[];heights=[]
            for row in [headers,*rows]:
                cells=[wrapped(value,widths[j]-20,size) for j,value in enumerate(row)]
                formatted.append([value for value,_ in cells]);heights.append(max(count for _,count in cells)*leading+2*padding+3)
            height=sum(heights);require(y+height<620,f'Editable table overflows slide {index}')
            native=slide.shapes.add_table(len(formatted),len(headers),unit(62),unit(y),unit(1156),unit(height));native.name='Editable evidence table'
            table=native.table;table.first_row=True
            for j,width in enumerate(widths):table.columns[j].width=unit(width)
            for i,row in enumerate(formatted):
                table.rows[i].height=unit(heights[i])
                for j,value in enumerate(row):
                    cell=table.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(colors['light' if i==0 else 'background'])
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
                    style_text(cell.text_frame,value,size,leading)
                    cell.text_frame.auto_size=MSO_AUTO_SIZE.NONE
                    # Keep table padding after generic text-frame styling.
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
            y+=height+18
        note,note_lines=wrapped(presentation_note(source),1116,12);note_height=note_lines*19+3
        require(y+note_height+28<656,f'Editable note overflows slide {index}')
        rect(slide,62,y,1156,note_height+25,name='Editable boundary note background')
        text(slide,note,82,y+11,1116,12,19,height=note_height+3,name='Evidence boundary note')
        text(slide,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인',62,678,900,10,15,'muted',name='Contribution footer')
        text(slide,f'PORTPLAN V12 / {index:02d} / {len(slides):02d}',1040,678,178,10,15,'muted',name='Page footer')
        for shape in slide.shapes:
            require(shape.shape_type!=MSO_SHAPE_TYPE.PICTURE,'Raster slide pictures are prohibited in the editable deck')
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=presentation.slide_width+1 and
                    shape.top+shape.height<=presentation.slide_height+1,f'Editable shape escapes slide {index}: {shape.name}')
    require(len(presentation.slides)==len(slides),'Editable deck page count changed')
    output=BytesIO();presentation.save(output);return output.getvalue()


def markdown_table(headers, rows):
    cell=lambda value:str(value).replace('|','\\|').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(map(cell,headers))+' |','| '+' | '.join(['---']*len(headers))+' |',
                      *['| '+' | '.join(map(cell,row))+' |' for row in rows]])


def build_slides(e):
    slides=[]; summary=e['summary']; protocol=e['protocol']
    def add(title,text,table=None,flow=None,note='',english=''):
        slides.append({'title':title,'text':text,'table':table,'flow':flow,'chart':None,'note':note,'english':english})
    add('완전한 초기해를 먼저 확보하면\n이후 날짜 선택도 달라질까요?',
        'PortPlan V12는 빠른 완전 초기해와 동일 요금의 필요 창고 선택을 결합합니다. 초기화와 이후 실행 실패를 분리해 검증합니다.',
        flow=['빠른 완전 seed','물리·요금·잠금 검산','같은 예산의 native 선택'],
        note='공개 자료 기반 합성 일 단위 연구입니다. 전역 최적·미래 안전·기업 실현 절감·우승을 보장하지 않습니다.',
        english='V12 combines an early complete incumbent with equal-fee necessary-warehouse selection. Initialization and later execution failures are reported separately, with no claim of universal robustness or contest rank.')
    add('하나의 날짜 장부에서\n반출·연속 적출·반납을 연결합니다.',
        '반출과 반납은 각각 운송 작업입니다. 적출은 연속 창고 일수를, 반출부터 반납 전까지는 체류 공간을 사용합니다.',
        table=(['모형','검사','범위 밖'],[['p / u / r','방출·납기·과거 잠금','도로·예약·기사 시간'],['일별 자원','truck / warehouse / staging','시간대별 현장'],['두 요금','공개 PNC·Maersk 해석','계약·세금·실제 청구']]),
        note='모든 주문은 시작부터 알려집니다. 미래 미공지 충격 지속시간은 계획기에 제공하지 않습니다.',
        english='The unchanged known-order daily model checks precedence, consecutive unloading, deadlines, resources and exact executed locks; it does not model routes, appointments or private invoice terms.')
    add('공개 증거로 개발하며\n기업 자료를 선행 조건으로 두지 않습니다.',
        '공개 출처·공개 요율·공개 수학 기준·재현 가능한 합성 실험을 각각의 범위로 검증합니다.',
        table=(['증거','가능한 주장'],[['공개 요율','명시한 요금 함수의 재계산'],['공개 관측·수학 기준','해당 자료에 한정한 검증'],['합성 주문·충격','고정 일별 모형의 방법 비교'],['기업 실단·견적·시험','선택 확장; 본 경로의 조건 아님']]),
        note='공개 관측을 PNC 실단으로, 교육 가격을 실제 견적으로, 시뮬레이션을 실현 절감으로 바꾸어 쓰지 않습니다.',
        english='Public sources and reproducible studies are the main route. Enterprise orders, quotes, contacts and pilots are optional; provenance and scope remain explicit.')
    add('V12: 초기 완전 후보도\n원래 호출 시계 안에서 확보합니다.',
        f"유효 incumbent가 없으면 기존 시계의 앞 {protocol['callBudgets']['seededHiGHSV12']['constructorBudgetMs']}ms 안에서 완전 seed를 구성·검산합니다. native는 남은 예산만 받습니다.",
        flow=['cache 검산 / 없거나 무효','완전 seed를 먼저 수락','남은 native 예산'],
        note='constructor 반환이나 내부 최초 발견 시각만으로는 수락 증거가 아닙니다. 실제 후보·검산·선택 기록을 보존합니다.',
        english='Without a valid incumbent, the constructor and its audited acceptance share the original daily clock. Native optimization gets only the remaining budget, rather than an additional call allowance.')
    add('같은 요금일 때의 창고 지표는\n미래 예측이 아닌 필요 하한입니다.',
        '동일 관측의 같은 요금 후보끼리, 다음 날 반드시 필요한 창고 점유 하한이 줄어드는 적시 완전 후보를 비교합니다.',
        table=(['선택 기준','보장하지 않는 것'],[['현재 요금·물리·과거 잠금','미래 미공지 용량'],['동일 요금의 필요 창고 하한','모든 충격에서 완주'],['완전한 점수·기한 기록','전역 최소 변경·전역 최적']]),
        note='더 낮은 하한은 미래 강건성의 충분조건이 아닙니다. 변경 횟수를 임의의 원화로 환산하지 않습니다.',
        english='The equal-fee score is a necessary next-day warehouse commitment bound under current public information, not a prediction or sufficient robustness guarantee.')
    development=e.get('development')
    if development:
        d=development['summary']; development_text=(f"이전 D0 개발 {d['timelyCompleteSeedCount']}/{d['endpointCount']}개 초기 seed 수락. 전체 호출 반환은 {fnum(min(d['wholeCallReturnRangeMs']),2)}–{fnum(max(d['wholeCallReturnRangeMs']),2)}ms였습니다.")
    else: development_text='이전 D0 개발 자료는 별도 증거입니다. 완전 초기해 수락과 호출 반환 및 전체 경로 완주는 서로 다른 측정입니다.'
    add('200ms는 협력식 목표입니다.\n반환 시간을 숨기지 않습니다.',development_text,
        table=(['시계','해석'],[['일 호출 200ms','구성·native·IPC·검산·선택 공유'],['냉 episode 5초','부모 검산·정상 exit·자손그룹 정리 확인 포함'],['이전 D0 개발','새 정식 성과·전체 완주가 아님']]),
        note='적시 수신만으로 완료가 아닙니다. 이전 15개 D0는 전체 63개 실패가 아니며 약 207–220ms 반환을 hard 200ms 보장으로 쓰지 않습니다.',
        english='The 200ms target is cooperative, not a hard call interruption. The 5-second parent boundary includes receipt audit, normal worker exit and verified process-group cleanup; timely receipt alone does not count. Old 15-endpoint D0 acceptance and 207–220ms return are not full completion.')
    add('새 업무 시나리오와 반복을\n구분한 최초 고정 실험입니다.',
        f"{len(e['index']['cases'])}개 업무 시나리오 × {len(protocol['algorithmRepeats'])}회 고정 반복 × {len(METHODS)}방법 = {e['wrapper']['episodeCount']:,} episode.",
        table=(['패널','입력','추론 단위'],[['동일 분포','72개 새 주문 시드','업무 시나리오'],['스트레스','72개 새 주문 시드','업무 시나리오'],['반복','3회 알고리즘/시계 반복','새 독립 업무 표본 아님']]),
        note='생성기·충격 목록은 이전 규칙을 유지합니다. V12 constructor와 native seed0의 반복을 독립 무작위 알고리즘 표본으로 주장하지 않습니다.',
        english='The frozen first run has 144 distinct business scenarios, three fixed algorithm/timing repeats and six methods. The generator and shock catalogs are retained; repeats are not additional independent business observations.')
    for panel in PANELS:
        s=summary['panels'][panel]
        rows=[]
        for method in METHODS:
            p=s['byMethod'][method]; init=p['initialization']
            rows.append([LABELS[method],str(p['completeEpisodeCount'])+'/216',init['acceptedInitialPlanCount'],init['knownInitialFailureCount'],init['knownLaterFailureCount'],init['unknownDayZeroCount']])
        add(PANELS[panel]+'\n초기화와 이후 실패를 나눕니다.',
            '완전 초기 계획을 얻은 수와 전체 실행 완주 수는 같지 않습니다. 확인하지 못한 D0 이력은 실패 0으로 바꾸지 않습니다.',
            table=(['방법','완주','초기 수락','D0 실패','후속 실패','D0 미확보'],rows),
            note='표의 숫자는 같은 정식 입력에서 수집한 원기록과 독립 검산에서 가져옵니다. 실패 전체 요금은 null입니다.',
            english='Initialization acceptance, known D0 failures, known later failures and unavailable history are separate from complete-path coverage on the same formal inputs.')
    for panel in PANELS:
        s=summary['panels'][panel]; rows=[]
        for name in ('v12VsV11','v12VsV10','v12VsHiGHS','v12VsV9','v12VsV7'):
            p=s['comparisons'][name]; f=p['fees']['incremental']; r=p['revisions']
            rows.append([LABELS[p['baselineMethod']],f"{p['completion']['gainedEpisodeCount']}/{p['completion']['lostEpisodeCount']}",f"{f['pairedScenarioCount']}/{f['pairedEpisodeCount']}",pct(f['weightedSavingsPct']),ci(f['uncertainty']['weightedSavingsPct95CI']),f"{r['baselineObservedSum']}→{r['candidateObservedSum']}"])
        add(PANELS[panel]+'\nV12의 각 기준별 공동 완주 비교',
            '추가 두 요금의 절감과 날짜 변경은 각 비교의 같은 공동 완주 집합에서만 해석합니다.',
            table=(['V12 / 기준','완주 얻음/잃음','시나리오/쌍','추가 두 요금 절감','조건부 95% CI','변경 기준→V12'],rows),
            note='절감 양수는 후보가 저렴함, 음수는 더 비쌈입니다. 0을 포함한 구간은 안정적 비용 우위가 아닙니다. 분모가 다른 비교를 순위로 합치지 않습니다.',
            english='Each comparator has its own jointly completed set and scenario-cluster interval. Negative savings mean higher fees; zero-inclusive intervals do not establish a stable fee advantage.')
    for panel in PANELS:
        common=summary['panels'][panel]['jointCompletePareto']; rows=[]
        for method in METHODS:
            p=common['byMethod'][method];rows.append([LABELS[method],fnum(p['scenarioMeanTotalCostSumKRW']),fnum(p['scenarioMeanIncrementalCostSumKRW']),fnum(p['scenarioMeanRevisionCountSum'],2)])
        add(PANELS[panel]+'\n여섯 방법 공통 집합도 따로 봅니다.',
            f"모든 여섯 방법이 같은 반복에 완주한 {common['jointScenarioCount']}개 시나리오 / {common['jointEpisodeCount']}개 일치 실행입니다.",
            table=(['방법','총요금 시나리오 평균 합','추가요금 평균 합','변경 평균 합'],rows),
            note='먼저 시나리오 안의 공통 반복을 평균합니다. 이 조건 집합의 Pareto는 전체 완주 손실을 상쇄하지 않고, 가중 종합 점수도 아닙니다.',
            english='The all-six joint subset has a different denominator from pairwise comparisons. Equal scenario weighting and separate fee/revision axes do not erase full-sample completion losses.')
    gate_rows=[]
    for panel in PANELS:
        g=summary['panels'][panel]['observedDecisionGates']; native=g['nativeComparisons']
        gate_rows.append([PANELS[panel],str(g['completionNonRegressionAllBaselines']),str(g['defaultRecommendationEligible']),'/'.join(str(native[key]['stabilityImprovementClaimAllowed']) for key in ('highsV8','stableHiGHSV10','commitmentHiGHSV11'))])
    add('사전 문턱을 통과했는지\n기준별로 그대로 공개합니다.',
        '포함된 다섯 기준 중 하나라도 완주 손실이 있으면 기본값 필요조건을 통과하지 않습니다.',
        table=(['패널','5기준 손실 0','관측 기본값 필요조건','native/V10/V11 안정성 조건'],gate_rows),
        note='순증가·조건부 절감·Pareto 위치가 손실을 취소하지 않습니다. 관측 문턱 통과도 현장 기본값이나 보편적 비퇴보 증명이 아닙니다. V5는 미포함입니다.',
        english='Every included comparator must have zero full-sample completion losses for default eligibility. Conditional native stability is separately assessed for native/V10/V11; finite observed gates are not production guarantees.')
    adverse=[]
    for panel in PANELS:
        for name in ('v12VsV11','v12VsV10','v12VsHiGHS','v12VsV9','v12VsV7'):
            p=summary['panels'][panel]['comparisons'][name];adverse.append([panel,LABELS[p['baselineMethod']],p['completion']['lostEpisodeCount'],len(p['fees']['total']['negativeScenarioIds']),len(p['fees']['incremental']['negativeScenarioIds'])])
    add('불리한 사례를 제외하지 않습니다.','모든 완료 손실과 요금 악화 식별자는 중·한 보고서에 전체 목록으로 보존합니다.',
        table=(['패널','V12 / 기준','완주 손실','총요금 악화 시나리오','추가요금 악화 시나리오'],adverse),
        note='두 요금의 모든 negativeScenarioIds/negativeEpisodeIds와 완료 lostEpisodeIds를 원기록에서 다시 도출해 누락을 거부합니다. 최악 손실도 함께 보고합니다.',
        english='All completion losses and adverse fee scenario/episode identifiers are preserved. Their membership is independently re-derived from complete scalar receipts before material generation.')
    add('개발 진단은 새 정식 분모에\n합치지 않습니다.',
        '기존 D0 초기화, 고정 실행 전후의 작은 재생, CP-SAT 진단은 각각의 오래된 입력 범위를 유지합니다.',
        table=(['별도 자료','해석 범위'],[['이전 D0 seed 개발','15 endpoints; 전체 D0 실패·전체 경로 아님'],['동일 prefix 개발 재생','동일 과거 잠금에서의 소규모 기전'],['CP-SAT','이전 9개 snapshot 독립 진단; 정식 6법 아님']]),
        note='별도 개발 결과로 새 주문 144개 또는 2,592개 실행 분모를 늘리거나, 다른 예산·정보 범위를 직접 순위로 만들지 않습니다.',
        english='Old initialization, controlled-prefix and CP-SAT studies are motivation/diagnostics only. CP-SAT is an independent old nine-snapshot diagnostic, not a seventh formal method.')
    add('원본·두 검산·자료를\n같은 SHA로 전달합니다.',
        '원본 입력·반복·실패·구성 후보·현재 관측·과거 잠금·선택 증거와 고정 source를 모두 남깁니다.',
        table=(['증거','SHA-256'],[['정식 index',e['wrapper']['formalIndex']['sha256']],['물리·선택 검산',e['hashes']['data/review-v12.json']],['통계 검산',e['hashes']['data/bootstrap-review-v12.json']]]),
        note='다운로드 청크를 같은 프로젝트 root로 합친 뒤 SHA를 검증합니다. 재실행은 새 디렉터리에 저장하고 최초 결과를 덮어쓰지 않습니다.',
        english='Formal inputs, first-run episodes, frozen source and two independent read-only reviews are bound by SHA. Downloads merge into one root; replays use fresh output directories.')
    add('본인이 재현하고 설명할 수 있는지\n별도로 확인합니다.',
        '연구 자료 생성은 본인 이해·등록·현장 조작을 대신 인증하지 않습니다. 빈 기록으로 실제 재현과 질문을 준비합니다.',
        table=(['항목','구분'],[['초기화 / 후속 실패 / 전체 완주','각 범위와 숫자를 설명'],['협력 200ms / 부모 5초','반환 초과를 숨기지 않음'],['공동 분모 / 조건부 CI / 음의 결과','실제 원기록을 찾아 설명'],['Browser / 본인 실연 / 당해 규정','실제 기록으로만 확인; 기업 시험은 선택']]),
        note='발표자료는 20쪽 이내입니다. 당해 발표/질의 시간, 자격·등록·AI/서약은 확인된 내용만 쓰며 우승을 보장하지 않습니다.',
        english='Entrant reproduction, presentation, eligibility, registration and current-year AI terms need actual records. Browser and personal acceptance remain separate; enterprise trials are optional.')
    require(1 <= len(slides) <= 20,'Competition deck exceeds 20 pages')
    return slides


def detailed_results(e, language):
    lines=[]
    for panel, s in e['summary']['panels'].items():
        lines.extend(['## '+(PANELS_ZH if language=='zh' else PANELS)[panel],''])
        for name,pair in s['comparisons'].items():
            lines.extend(['### '+name+' · '+pair['candidateMethod']+' / '+pair['baselineMethod'],''])
            rows=[]
            for key in ('total','incremental'):
                fee=pair['fees'][key];u=fee['uncertainty'];rows.append([key,f"{fee['pairedScenarioCount']}/{fee['pairedEpisodeCount']}",fnum(fee['baselineScenarioMeanSumKRW'],6),fnum(fee['candidateScenarioMeanSumKRW'],6),pct(fee['weightedSavingsPct']),ci(u['weightedSavingsPct95CI']),fnum(fee['worstScenarioLossKRW'],6),json.dumps(fee['scenarioWinTieLoss'])])
                rows.append([key+' raw episode sums','',fnum(fee['baselineEpisodeSumKRW'],6),fnum(fee['candidateEpisodeSumKRW'],6),'','KRW CI '+ci(u['meanSavings95CI']),'',''])
            lines.extend([markdown_table(['fee','scenario / repeat','baseline KRW','candidate KRW','savings %','conditional scenario CI','worst increase KRW','win/tie/loss'],rows),'',
                          'Completion gained/lost: '+str(pair['completion']['gainedEpisodeCount'])+'/'+str(pair['completion']['lostEpisodeCount'])+'; completion scenario CI '+ci(pair['completion']['uncertainty']['meanScenarioRateGain95CI']),
                          'Revisions on same matching complete set: '+str(pair['revisions']['baselineObservedSum'])+' → '+str(pair['revisions']['candidateObservedSum'])+'; n='+str(pair['revisions']['pairedEpisodeCount']),''])
            for key in ('total','incremental'):
                for field in ('negativeScenarioIds','negativeEpisodeIds'):
                    lines.extend(['**'+key+'/'+field+' — ALL retained**','',*(['- `'+identifier+'`' for identifier in pair['fees'][key][field]] or ['none / 없음 / 无']),''])
            for field in ('gainedEpisodeIds','lostEpisodeIds'):
                lines.extend(['**'+field+' — ALL retained**','',*(['- `'+identifier+'`' for identifier in pair['completion'][field]] or ['none / 없음 / 无']),''])
        lines.extend(['### All-six common-complete subset and observed necessary gates','','```json',json.dumps({'jointCompletePareto':s['jointCompletePareto'],'observedDecisionGates':s['observedDecisionGates']},ensure_ascii=False,indent=2),'```',''])
    return lines


def material_texts(e,slides):
    zh=['# PortPlan V12：及时完整初解与同费必要仓位','',
        'AI 编写的准备材料；不代表本人已理解、亲自完成、已经报名或通过浏览器验收。公开资料主路线无需企业实单、私人报价或联系人；企业试点为可选增强。','',
        '200ms 是调用合作式目标，不是硬返回保证；完整冷 episode 5 秒须包含父进程终态审核、worker正常退出及进程组清理验证，仅及时收件不算完成。旧 15/15 D0 初解开发的调用返回约 207–220ms，不是全部 63 个 D0 失败，更不是完整路径完成率。',
        '本轮 144 个新业务场景、3 次预固定算法/时钟重复、6 方法，共 2,592 次 episode。重复不是独立业务样本；失败费用为 null。全路径结果包含不同初始计划，不自动归因为相同历史下恢复。',
        '费用和改期使用各对共同完成集合，先场景内平均、再等权；场景聚类区间不能消除完成选择偏差。任何基准完成损失不能由净增、费用或 Pareto 抵消。V5 未列入本轮，不推断对 V5 不退步。',
        '必要仓位是当前快照下的必要下界，不预测未来。CP-SAT 只作旧 9 快照独立诊断，不混入正式六法分母，也不据此宣称胜过 CP-SAT。','']
    ko=['# PortPlan V12 · 적시 완전 초기해와 동일 요금 필요 창고 선택','',
        'AI 작성 준비 자료. 본인 창작·이해·등록·Browser 조작을 대신 인증하지 않습니다. 공개 자료가 기본 경로이고 기업 실단·견적·연락처·시험은 선택 확장입니다.','',
        '200ms는 협력식 호출 목표이며 hard 반환 보장이 아닙니다. 부모 5초는 수신 검산·worker 정상 종료·자손 프로세스 그룹 정리 확인까지 포함하며 적시 수신만으로 완료가 아닙니다. 이전 15/15 D0의 약 207–220ms 반환은 전체 63 D0 실패나 전체 경로 완주가 아닙니다.',
        '새 업무 시나리오 144개 × 고정 반복 3회 × 6방법 = 2,592 episode. 반복은 추가 독립 업무 표본이 아닙니다. 실패 비용은 null이고 초기 계획 차이를 동일 prefix 복구로 해석하지 않습니다.',
        '시나리오별 공동 완주 반복을 먼저 평균합니다. 조건부 CI는 완주 선택 편향을 제거하지 않으며 완주 손실을 비용·순증가·Pareto로 상쇄하지 않습니다. V5 및 CP-SAT은 정식 6법에 포함되지 않습니다.','']
    en=['# PortPlan V12 — timely complete initialization and equal-fee necessary warehouse commitment','',
        'AI-created preparation; no entrant-understanding, authorship, registration, field or browser attestation. Public evidence is the main route; enterprise orders, quotes, contacts and pilots are optional.',
        'The 200ms call budget is cooperative. Old 15-endpoint D0 development returned in approximately 207–220ms and is neither all 63 known D0 failures nor full rolling completion. The parent 5-second deadline includes receipt audit, normal worker exit and verified process-group cleanup; timely receipt alone is insufficient.',
        'Formal first-run counts use 144 new business scenarios, three fixed algorithm/timing repeats and six methods: 2,592 episodes. Repeats are not additional independent business samples. Failure fees remain null.',
        'Every fee/revision comparison uses its own matching jointly completed set with equal scenario weighting. Conditional scenario-cluster intervals do not remove completion-selection bias. Completion losses cannot be cancelled by net gains, fees or Pareto membership. CP-SAT is separate old nine-snapshot diagnosis; V5 is omitted.','']
    zh_sections=[
        ('先交付完整初解，再比较同费必要仓位','V12将及时接受的完整初解与同费候选必要仓位选择结合；初始化、后续失败与整段完成分别检验。'),
        ('日级约束和两费模型','提箱和还空分别消耗运输箱次；卸货连续占用仓库，在场占用为[p,r)。未建模路线、日内预约、私有合同和全部企业成本。'),
        ('公开证据主路线','公开费率、公开观测、数学基准和合成实验各有范围。企业实单、私人报价、联系人和试点为可选增强，不是继续研发的条件。'),
        ('完整初解共享原有调用时钟','没有有效incumbent时，构造与独立审核也计入原200ms合作式调用；原生求解只使用剩余预算。内部发现或函数返回不能代替实际接受记录。'),
        ('必要仓位是下界','同费候选的明日必需仓位得分来自当前快照的必要下界，不预测未公告容量，也不保证未来完整执行或全局最少改期。'),
        ('合作式预算不等于硬返回保证','旧D0开发的15个及时初解与约207–220ms调用返回必须同时披露；它不是全部63个D0失败，也不是15次完整滚动完成。'),
        ('场景、重复与六方法','正式记录使用144个新业务场景、三次固定算法/时钟重复和六方法；生成器与扰动目录沿用既有规则。重复不能当成新的独立业务样本。'),
        ('同生成机制：初始化与后续失败','按同一正式输入列出完整执行、初始计划接受、已知D0失败、已知后续失败和未知D0历史。未知不能改填为零。'),
        ('压力扰动：初始化与后续失败','压力面板单列，不与同分布合并排名。完整初解接受不等于完成整条真实冲击路径。'),
        ('同生成机制：各基准的共同完成比较','每对比较使用自己的共同完成场景及重复。正节省表示候选更便宜，负数表示更贵；区间跨零不能称稳健费用优势。'),
        ('压力扰动：各基准的共同完成比较','费用和改期必须在完全相同的配对集合中解释，先场景内平均再等权。不能跨分母宣布综合领先。'),
        ('同生成机制：六法共同集合','六法共同完成集合与两两配对集合不同；费用和改期分别作为轴，不将一次改期任意换算为韩元。'),
        ('压力扰动：六法共同集合','共同集合内的Pareto位置不抵消全样本完成损失。无共同样本或未知改期保持null。'),
        ('预声明观察门槛','对任意包含基准有完成损失即不通过默认候选必要门槛；相对三个原生基准的条件稳定性各自判断。V5未纳入，不推断对V5不退步。'),
        ('全部不利结果','总两费与新增两费的全部负场景、负运行和全部完成损失均在下方逐条保留；不删除昂贵、无解、迟交或未知案例。'),
        ('开发诊断不混入正式分母','旧D0、同前缀控制延续和CP-SAT旧9快照诊断分别作为开发动机，不是新正式输入或第七种正式方法。'),
        ('证据和跨电脑复现','输入、原始episode、首次冻结源码和两份独立复核由SHA绑定。下载清单全部包合并同根，重跑另建目录，不覆盖首次结果。'),
        ('本人复现与答辩独立记录','材料生成不等于本人理解、资格报名或真实浏览器验收。使用空表记录实际复现与问答，不代签或代讲，官方时长未知时不杜撰。')]
    translated_headers={'방법':'方法','완주':'完整','초기 수락':'初始接受','D0 실패':'D0失败','후속 실패':'后续失败','D0 미확보':'D0未知','V12 / 기준':'V12 / 基准','완주 얻음/잃음':'完成增/损','시나리오/쌍':'场景/配对','추가 두 요금 절감':'新增两费节省','조건부 95% CI':'条件95%区间','변경 기준→V12':'改期基准→V12','총요금 시나리오 평균 합':'场景均值总两费之和','추가요금 평균 합':'场景均值新增两费之和','변경 평균 합':'场景均值改期之和'}
    for i,slide in enumerate(slides,1):
        ko.extend([f"## {i:02d}. {slide['title'].replace(chr(10),' ')}",'',slide['text'],''])
        if slide['table']:ko.extend([markdown_table(*slide['table']),''])
        if slide['flow']:ko.extend([' → '.join(slide['flow']),''])
        ko.extend([slide['note'],''])
        title,paragraph=zh_sections[i-1];zh.extend([f'## {i:02d}. {title}','',paragraph,''])
        if slide['table'] and (8<=i<=15 or i==17):
            headers,rows=slide['table'];reverse={value:key for key,value in LABELS.items()}
            zh.extend([markdown_table([translated_headers.get(value,value) for value in headers],[[reverse.get(str(value),value) for value in row] for row in rows]),''])
        en.extend([f'## {i:02d}','',''+slide['english'],''])
        if slide['table']:
            headers,rows=slide['table'];en.extend([markdown_table(['incremental fee savings' if h=='추가 두 요금 절감' else h for h in headers],rows),''])
    zh.extend(detailed_results(e,'zh'));ko.extend(detailed_results(e,'ko'))
    evidence_text='\n```json\n'+json.dumps(e['hashes'],ensure_ascii=False,indent=2)+'\n```\n'
    zh.extend(['## 原始证据哈希',evidence_text]);ko.extend(['## 원증거 SHA',evidence_text]);en.extend(['## Evidence hashes',evidence_text,'All adverse identifiers are fully listed in both reports and bound formal JSON; no filtering.'])
    contribution='# PortPlan V12 · 贡献记录与本人确认空表\n\nAI 已协助实现、检查脚本和整理材料；开源求解器/依赖与原数据来源按项目记录列示。此记录不把 AI 工作认作本人直接完成。\n\n| 本人实际工作 | 证据引用 | 实际日期KST | 本人核对 |\n| --- | --- | --- | --- |\n|  |  |  |  |\n\n本人资格、当届 AI/誓约、理解、签名与实际答辩状态：未验证。无需公开个人信息。\n'
    optional='# V12 可选企业试点空表 / 선택 현장 시험 빈 기록\n\n可选增强，不是公开研究路线的条件；不索要企业实单、报价或联系人。未发生则保持空白，不生成访谈、采购或收益。\n\n| 实际发生事项 | 去标识内部证据 | 日期KST | 实际观察者角色 |\n| --- | --- | --- | --- |\n|  |  |  |  |\n'
    entrant='# V12 本人复现与答辩空表 / 본인 재현·발표 빈 기록\n\n尚未填写，不表示本人已理解、报名、签署誓约或完成现场演示。时长以当届官方为准，不编造。\n\n| 实际任务/问题 | 本人实际操作/回答证据 | 观察者角色 | 日期KST | 结果与待改进 |\n| --- | --- | --- | --- | --- |\n|  |  |  |  |  |\n\n解释：D0与完整路径；合作200ms与返回超时；共同分母/条件CI；负例；必要仓位下界；CP-SAT旧诊断；真实贡献。\n'
    return ['\n'.join(ko)+'\n','\n'.join(en)+'\n',contribution,'\n'.join(zh)+'\n',optional,entrant]


def audit_rendered(pdf,pptx,slides):
    import fitz
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    import unicodedata
    normalize=lambda text:''.join(unicodedata.normalize('NFKC',str(text)).split())
    document=fitz.open(stream=pdf,filetype='pdf');deck=Presentation(BytesIO(pptx));require(len(document)==len(deck.slides)==len(slides)<=20,'Rendered material page count differs or exceeds 20')
    text_shapes=tables=0; table_font_sizes=[]
    for number,(page,slide,source) in enumerate(zip(document,deck.slides,slides),1):
        require(not page.get_images(),'PDF must use vector/text content, not slide screenshots')
        for block in page.get_text('dict')['blocks']:
            if 'lines' in block:
                x0,y0,x1,y1=block['bbox'];require(x0>=-1 and y0>=-1 and x1<=page.rect.width+1 and y1<=page.rect.height+1,'PDF text escapes page '+str(number))
        ppt_strings=[]
        for shape in slide.shapes:
            require(shape.shape_type!=MSO_SHAPE_TYPE.PICTURE,'PPT must remain native editable')
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=deck.slide_width+1 and shape.top+shape.height<=deck.slide_height+1,'PPT shape escapes slide')
            if shape.has_text_frame:ppt_strings.append(shape.text);text_shapes+=bool(shape.text.strip())
            if shape.has_table:
                tables+=1;ppt_strings.extend(cell.text for row in shape.table.rows for cell in row.cells)
                for row in shape.table.rows:
                    for cell in row.cells:
                        for paragraph in cell.text_frame.paragraphs:
                            if paragraph.text:
                                require(paragraph.font.size is not None and paragraph.font.size.pt>=12,'Editable table text below 16px/12pt minimum')
                                table_font_sizes.append(paragraph.font.size.pt)
        expected=[source['title'],source['text'],presentation_note(source),*(source['flow'] or [])]
        if source['table']:
            headers,rows=presentation_table(source);expected.extend(value for row in [headers,*rows] for value in row)
        pdf_text=normalize(page.get_text());ppt_text=normalize(' '.join(ppt_strings))
        for value in expected:require(normalize(value) in pdf_text and normalize(value) in ppt_text,'Rendered source text missing on page '+str(number)+': '+str(value)[:70])
    require(text_shapes>=len(slides)*3 and tables>=1,'Editable native objects missing')
    return {'valid':True,'schemaVersion':'portplan-v12-material-review','pdfPageCount':len(document),'pptxSlideCount':len(deck.slides),'editableTextShapes':text_shapes,'editableTables':tables,'minimumEditableTableFontPt':min(table_font_sizes),'sameSlideSource':True,'sourceTextCoverage':True,'pdfAndPptBoundsValid':True,'solverRuns':0,'browserInteractionCertified':False,'entrantUnderstandingCertified':False}


def validate_report_coverage(e,ko,zh):
    for panel in e['summary']['panels'].values():
        for pair in panel['comparisons'].values():
            ids=pair['completion']['lostEpisodeIds']
            for key in ('total','incremental'):
                ids=ids+pair['fees'][key]['negativeScenarioIds']+pair['fees'][key]['negativeEpisodeIds']
            for identifier in ids:require('`'+identifier+'`' in ko and '`'+identifier+'`' in zh,'Report omitted adverse identifier')


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__,epilog='Rendering requires requirements-materials-v12.txt and the system Noto CJK font. Schema/check-only do not render.')
    action=parser.add_mutually_exclusive_group();action.add_argument('--check-only',action='store_true');action.add_argument('--generate',action='store_true',help='After root review, render and write V12 outputs only.')
    parser.add_argument('--schema-only',action='store_true',help='Inspect declared layout/schema without claiming any formal result.')
    parser.add_argument('--bootstrap-review',default='data/bootstrap-review-v12.json')
    parser.add_argument('--replace-generated',action='store_true')
    args=parser.parse_args(argv)
    if args.schema_only:
        require(not args.generate,'Schema-only cannot generate materials')
        print(json.dumps({'state':'schema-only-no-formal-result','writes':0,'methods':METHODS,'comparisons':list(PAIRS),'plannedBusinessScenarios':144,'plannedEpisodes':2592,'maximumSlides':20}));return
    e=verified_inputs(bootstrap_review=args.bootstrap_review);slides=build_slides(e);ko,en,contribution,zh,optional,entrant=material_texts(e,slides);validate_report_coverage(e,ko,zh)
    if not args.generate:
        print(json.dumps({'valid':True,'checkOnly':True,'writes':0,'pages':len(slides),'evidenceHashes':e['hashes'],'adverseIdentifierOccurrences':e['adverseIdentifierOccurrences']}));return
    targets=[ROOT/'docs'/name for name in OUTPUTS];require(args.replace_generated or not any(path.exists() for path in targets),'V12 materials already exist; explicit --replace-generated required')
    pdf=pdf_document(slides);pptx=pptx_document(slides);audit=audit_rendered(pdf,pptx,slides)
    audit.update({'formalIndexSha256':e['wrapper']['formalIndex']['sha256'],'evidenceHashes':e['hashes'],'adverseIdentifierOccurrences':e['adverseIdentifierOccurrences'],'generatorSha256':sha(Path(__file__))})
    values=[html_document(slides).encode(),pdf,ko.encode(),en.encode(),contribution.encode(),pptx,zh.encode(),optional.encode(),entrant.encode()]
    # Every guard, report check, renderer and in-memory saved-object audit finishes before the first write.
    for path,contents in zip(targets,values):path.write_bytes(contents)
    audit['outputHashes']={str(path.relative_to(ROOT)):sha(path) for path in targets};(ROOT/'data/material-review-v12.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'valid':True,'pages':len(slides),'written':[str(path.relative_to(ROOT)) for path in targets],'materialReview':'data/material-review-v12.json'},ensure_ascii=False))


if __name__=='__main__':main()



