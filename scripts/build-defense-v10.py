#!/usr/bin/env python3
"""Generate result-driven Korean V10 materials only after complete reviewed evidence.

No solver, browser, experiment generation, network access or old-material writes.
The published wrapper, restored formal index and independent review must bind to
the same bytes. All validation and PDF layout finish in memory before any write.
"""
from __future__ import annotations

import argparse
from collections import Counter
from html import escape
from io import BytesIO
import hashlib
import json
import math
from pathlib import Path
import shlex

ROOT = Path(__file__).resolve().parents[1]
METHODS = ('stableRetainV5', 'costPolishV7', 'adaptiveV9', 'highsV8', 'stableHiGHSV10')
LABELS = dict(zip(METHODS, ('V5 유효 계획 유지', 'V7 비용 다듬기', 'V9 적시 후보 전달', '독립 native HiGHS', 'V10 동일 요금 계획 유지')))
PANELS = {'same-distribution': '동일 생성분포 · 새 주문 시드', 'stress': '스트레스 · 새 주문 시드'}
PANELS_ZH = {'same-distribution': '同生成机制·新订单种子', 'stress': '压力扰动·新订单种子'}
PAIRS = {'v10VsHiGHS': ('highsV8', 'stableHiGHSV10'), 'v10VsV9': ('adaptiveV9', 'stableHiGHSV10'),
         'v10VsV7': ('costPolishV7', 'stableHiGHSV10'), 'v10VsV5': ('stableRetainV5', 'stableHiGHSV10'),
         'highsVsV7': ('costPolishV7', 'highsV8')}
OUTPUTS = ('defense-v10.ko.html', 'defense-v10.ko.pdf', 'submission-v10.ko.md', 'defense-v10.en.md',
           'contribution-log-v10.md', 'defense-v10.ko.pptx', 'report-v10.zh.md',
           'field-acceptance-v10.blank.md', 'entrant-defense-v10.blank.md')


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def relative_path(root, value):
    require(isinstance(value, str), 'Evidence path must be a string')
    path = (root / value).resolve()
    require(path.is_relative_to(root.resolve()), 'Evidence path escapes project: ' + value)
    require(path.is_file(), 'Required completed evidence is missing: ' + value)
    return path


def load(root, name, hashes):
    path = relative_path(root, name)
    hashes[name] = sha(path)
    return json.loads(path.read_text())


def finite(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def close(actual, expected, label):
    require(finite(actual) and finite(expected) and abs(actual - expected) <= max(1e-5, abs(expected) * 1e-10),
            'Independent point estimate differs: ' + label)


def same_interval(actual, expected, label, tolerance):
    if expected is None:
        require(actual is None, 'Independent interval null differs: ' + label)
        return
    require(isinstance(actual, list) and isinstance(expected, list) and len(actual) == len(expected) == 2,
            'Malformed independent interval: ' + label)
    for a, b in zip(actual, expected):
        require(finite(a) and finite(b) and abs(a - b) <= tolerance,
                'Independent interval differs: ' + label)


def same_tree(actual, expected, label):
    if isinstance(expected, dict):
        require(isinstance(actual, dict) and set(actual) == set(expected), 'Independent tree keys differ: '+label)
        for key in expected: same_tree(actual[key], expected[key], label+'/'+key)
    elif isinstance(expected, list):
        require(isinstance(actual, list) and len(actual) == len(expected), 'Independent tree length differs: '+label)
        for i, (left, right) in enumerate(zip(actual, expected)): same_tree(left, right, label+'/'+str(i))
    elif finite(expected):
        require(finite(actual) and abs(actual-expected) <= 1e-6, 'Independent tree number differs: '+label)
    else:
        require(type(actual) is type(expected) and actual == expected, 'Independent tree value differs: '+label)


def fnum(value, digits=0):
    if not finite(value):
        return 'null'
    return f'{value:,.{digits}f}'


def pct(value):
    return fnum(value, 2) + '%' if finite(value) else 'null'


def ci(value):
    # Keep the raw sign of small endpoints; rounding must not turn a strictly
    # positive/negative bound into an apparent zero for the declared gates.
    return '[' + ', '.join(format(v, '.10g') if finite(v) else 'null' for v in value) + ']' if isinstance(value, list) and len(value) == 2 else 'null'


def verified_inputs(root=ROOT, bootstrap_review='data/bootstrap-review-v10.json'):
    """No fixture, partial run, unbound summary or unreviewed interval may publish."""
    hashes = {}
    wrapper = load(root, 'data/v10-results.json', hashes)
    require(wrapper.get('schemaVersion') == 'portplan-v10-published-results' and
            wrapper.get('executionState') == 'complete', 'Published V10 results are not a completed formal wrapper')
    reference = wrapper.get('formalIndex', {})
    formal_path = relative_path(root, reference.get('path'))
    require(sha(formal_path) == reference.get('sha256'), 'Restored formal-index SHA differs; restore exact evidence first')
    hashes[reference['path']] = reference['sha256']
    index = json.loads(formal_path.read_text())
    require(index.get('schemaVersion') == 'portplan-v10-streamed-experiment' and
            index.get('executionState') == 'complete', 'Formal V10 index is not complete')
    require(wrapper.get('summary') == index.get('summary') and wrapper.get('protocol') == index.get('protocol'),
            'Published summary/protocol differs from the formal index')
    protocol = load(root, 'data/protocol-v10.json', hashes)
    require(protocol == index['protocol'] and hashes['data/protocol-v10.json'] == index.get('protocolHash'),
            'Protocol bytes differ from frozen formal protocol')
    require(tuple(protocol.get('methods', [])) == METHODS, 'Unexpected formal method set/order')
    require(protocol['budget']['hardWallMs'] == 5000, 'Unexpected whole-episode budget')
    require(protocol['callBudgets']['stableHiGHSV10'] == protocol['callBudgets']['highsV8'] and
            protocol['callBudgets']['stableHiGHSV10']['timeBudgetMs'] == 200, 'V10 changed the original native numeric budget')
    require(protocol['nativeSelectionRule']['fixedNumericalEpsilon'] == 1e-6, 'Unexpected strict fee replacement epsilon')
    repeats = protocol.get('algorithmRepeats', [])
    require(len(repeats) == 3 and len({r['seed'] for r in repeats}) == 3, 'Missing predeclared algorithm repeats')
    cases = index.get('cases', [])
    require(len(cases) == 144 and len({c['id'] for c in cases}) == 144, 'Missing/duplicate formal business scenarios')
    require(Counter(c['panel'] for c in cases) == Counter({'same-distribution': 72, 'stress': 72}), 'Panel coverage differs')
    episode_count = 0
    for case in cases:
        require(len(case.get('repeats', [])) == len(repeats), 'Missing repeat in ' + case['id'])
        for i, repetition in enumerate(case['repeats']):
            require(repetition.get('repeatIndex') == i + 1 and repetition.get('algorithmSeed') == repeats[i]['seed'], 'Reordered repeat seed')
            require(set(repetition.get('policies', {})) == set(METHODS), 'Missing policy receipt in ' + case['id'])
            episode_count += len(repetition['policies'])
    require(episode_count == 2160 and wrapper.get('scenarioCount') == len(cases) and wrapper.get('episodeCount') == episode_count,
            'Published scenario/episode counts differ from complete receipts')
    review = load(root, 'data/review-v10.json', hashes)
    require(review.get('schemaVersion') == 'portplan-independent-v10-run-review' and review.get('valid') is True and
            review.get('fixture') is False and review.get('inputIndexSha256') == reference['sha256'],
            'Independent physical review is invalid, a fixture or bound to different result bytes')
    require(review.get('summary', {}).get('episodes') == episode_count, 'Independent review did not cover every episode')
    require(review.get('solverRuns') == 0 and review.get('importsJavaScript') is False,
            'Review scope is not the independent read-only ledger')
    summary = index['summary']
    require(summary.get('inferentialState') == 'complete-scenario-cluster-bootstrap', 'Inferential summaries are pending')
    require(set(summary.get('panels', {})) == set(PANELS), 'Missing/extra formal panel')
    for panel in PANELS:
        published = summary['panels'][panel]
        checked = review.get('pointEstimates', {}).get(panel)
        require(checked is not None and published['scenarioCount'] == 72, 'Missing independently reviewed panel')
        require(set(published.get('byMethod', {})) == set(METHODS) and
                set(published.get('comparisons', {})) == set(PAIRS), 'Unexpected method/comparison coverage')
        for method in METHODS:
            produced, independent = published['byMethod'][method], checked['byMethod'][method]
            for key in ('recordedEpisodeCount', 'completeEpisodeCount', 'fullyCompletedScenarioCount', 'anyCompletedScenarioCount'):
                require(produced[key] == independent[key], 'Reviewed method count differs: ' + panel + '/' + method + '/' + key)
        for name, (baseline, candidate) in PAIRS.items():
            pair, independent = published['comparisons'][name], checked['comparisons'][name]
            require((pair['baselineMethod'], pair['candidateMethod']) == (baseline, candidate), 'Comparison identity differs')
            require(pair['completion']['gainedEpisodeCount'] == independent['gainedEpisodes'] and
                    pair['completion']['lostEpisodeCount'] == independent['lostEpisodes'], 'Completion gain/loss differs')
            for key in ('total', 'incremental'):
                fee, validated = pair['fees'][key], independent['fees'][key]
                require(fee['pairedScenarioCount'] == validated['pairedScenarios'] and fee['pairedEpisodeCount'] == validated['pairedEpisodes'], 'Fee denominator differs')
                require(fee['scenarioWinTieLoss'] == validated['scenarioWinTieLoss'], 'Fee win/tie/loss differs')
                close(fee['baselineScenarioMeanSumKRW'], validated['baselineScenarioMeanSumKRW'], panel + '/' + name + '/baseline')
                close(fee['candidateScenarioMeanSumKRW'], validated['candidateScenarioMeanSumKRW'], panel + '/' + name + '/candidate')
            require([pair['revisions']['baselineObservedSum'], pair['revisions']['candidateObservedSum']] == independent['revisions'], 'Revision totals differ')
    require(bool(bootstrap_review), 'A separate independent bootstrap review is required for V10 materials')
    statistics = load(root, bootstrap_review, hashes)
    require(statistics.get('schemaVersion') == 'portplan-independent-v10-bootstrap-review' and
            statistics.get('valid') is True and statistics.get('fixture', False) is False and
            statistics.get('readOnly') is True and statistics.get('importsJavaScript') is False and
            statistics.get('solverRuns') == 0, 'Bootstrap review is not a valid independent read-only formal review')
    require(statistics.get('inputIndexSha256') == reference['sha256'], 'Bootstrap review belongs to different index bytes')
    require(statistics.get('frozenManifestSha256') == review.get('frozenManifestSha256') and
            bool(statistics.get('frozenManifestSha256')), 'Statistical and physical reviews bind to different frozen manifests')
    require(statistics.get('episodeCount') == episode_count and statistics.get('businessScenarioCount') == len(cases),
            'Bootstrap review coverage differs')
    procedure = statistics.get('procedure', {})
    declared = protocol['analysis']['bootstrap']
    require(procedure.get('samplingUnit') == 'whole business scenario' and procedure.get('seed') == declared['seed'] and
            procedure.get('draws') == declared['draws'], 'Bootstrap sampling procedure differs from frozen protocol')
    require(procedure.get('tolerances') == {'feeKRW': 1e-6, 'percentagePoints': 1e-9, 'completionRate': 1e-12},
            'Unexpected bootstrap numerical tolerances')
    require(set(statistics.get('panels', {})) == set(PANELS), 'Missing/extra bootstrap panel')
    for panel in PANELS:
        produced, checked = summary['panels'][panel], statistics['panels'][panel]
        require(checked.get('scenarioCount') == produced['scenarioCount'] == 72, 'Bootstrap panel scenario count differs')
        require(set(checked.get('byMethod', {})) == set(METHODS) and set(checked.get('comparisons', {})) == set(PAIRS),
                'Bootstrap method/comparison coverage differs')
        for method in METHODS:
            for key in ('recordedEpisodeCount', 'completeEpisodeCount', 'fullyCompletedScenarioCount', 'anyCompletedScenarioCount'):
                require(checked['byMethod'][method][key] == produced['byMethod'][method][key], 'Bootstrap completion eligibility differs')
            close(checked['byMethod'][method]['meanScenarioCompletionRate'], produced['byMethod'][method]['meanScenarioCompletionRate'],
                  panel + '/' + method + '/scenario-completion-rate')
        for name in PAIRS:
            pair, validated = produced['comparisons'][name], checked['comparisons'][name]
            same_interval(validated['completionRateGain95CI'], pair['completion']['uncertainty']['meanScenarioRateGain95CI'],
                          panel + '/' + name + '/completion', 1e-12)
            for key in ('total', 'incremental'):
                original, audited = pair['fees'][key], validated['fees'][key]
                uncertainty = original['uncertainty']
                require(audited['pairedScenarioCount'] == original['pairedScenarioCount'] and
                        audited['pairedEpisodeCount'] == original['pairedEpisodeCount'], 'Bootstrap paired denominator differs')
                require(uncertainty['unit'] == 'scenario' and audited['scenarioCount'] == uncertainty['scenarioCount'] == original['pairedScenarioCount'] and
                        audited['draws'] == uncertainty['draws'] == declared['draws'], 'Bootstrap draw/unit counts differ')
                same_interval(audited['meanSavings95CI'], uncertainty['meanSavings95CI'], panel + '/' + name + '/' + key + '/KRW', 1e-6)
                same_interval(audited['weightedSavingsPct95CI'], uncertainty['weightedSavingsPct95CI'], panel + '/' + name + '/' + key + '/percent', 1e-9)
        for field in ('jointCompletePareto', 'observedDecisionGates'):
            require(field in produced and field in checked, 'Independent Pareto/decision-gate review is missing')
            same_tree(produced[field], checked[field], panel+'/'+field)
        require(produced['observedDecisionGates']['assessmentState'] == 'complete-panel', 'Decision gates are not from a complete panel')
    native_counts = review.get('nativeStabilitySelection', {})
    require(all(type(native_counts.get(key)) is int and native_counts[key] >= 0 for key in
                ('calls', 'cacheAcceptedCalls', 'nativeAcceptedCalls', 'retainedCacheCalls', 'completeNativeProposals', 'strictReplacements', 'feeTiesRetained')),
            'Independent native original-date selection proof coverage is missing')
    # Historical V9 results are a separate new-order experiment, not V10 pairs.
    previous = load(root, 'data/v9-results.json', hashes)
    require(previous.get('schemaVersion') == 'portplan-v9-published-results' and previous.get('executionState') == 'complete',
            'Historical V9 wrapper is not complete')
    old_reference = previous.get('formalIndex', {})
    old_path = relative_path(root, old_reference.get('path'))
    require(sha(old_path) == old_reference.get('sha256'), 'Historical V9 index hash differs')
    hashes[old_reference['path']] = old_reference['sha256']
    old_index = json.loads(old_path.read_text())
    require(old_index.get('executionState') == 'complete' and previous.get('summary') == old_index.get('summary'),
            'Historical V9 summary differs from its complete first run')
    old_review = load(root, 'data/review-v9.json', hashes)
    require(old_review.get('valid') is True and old_review.get('fixture') is False and
            old_review.get('inputIndexSha256') == old_reference['sha256'], 'Historical V9 review is not bound to its run')
    acceptance = load(root, 'data/championship-acceptance-v10.json', hashes)
    require(acceptance.get('schemaVersion') == 'portplan-championship-acceptance-v10' and
            acceptance.get('businessPilot', {}).get('requiredForPublicEvidenceRoute') is False and
            acceptance['businessPilot'].get('passed') is None,
            'Public-evidence materials require the optional, unverified business-pilot state; do not imply field acceptance')
    diagnosis = None
    if (root/'data/v10-loss-diagnosis.json').exists():
        diagnosis = load(root, 'data/v10-loss-diagnosis.json', hashes)
        require(diagnosis.get('valid') is True and diagnosis.get('readOnlyOriginalEvidence') is True and
                diagnosis.get('newSolverCalls') == diagnosis.get('onlinePolicyReruns') == diagnosis.get('newHeldOutExperiments') == 0 and
                reference['sha256'] in diagnosis.get('sourceHashes', {}).values(),
                'Post-hoc loss diagnosis is not read-only or belongs to different formal bytes')
    return {'root': root, 'wrapper': wrapper, 'index': index, 'protocol': protocol, 'summary': summary,
            'review': review, 'bootstrapReview': statistics, 'previousV9': previous, 'acceptance': acceptance, 'diagnosis': diagnosis, 'hashes': hashes,
            'materialOptions': {'bootstrap-review': bootstrap_review}}


def pair_label(pair):
    return LABELS[pair['candidateMethod']] + ' / ' + LABELS[pair['baselineMethod']]


def comparison_rows(panel):
    return [[pair_label(p), f"{p['fees']['total']['pairedScenarioCount']} / {p['fees']['total']['pairedEpisodeCount']}",
             pct(p['fees']['total']['weightedSavingsPct']), pct(p['fees']['incremental']['weightedSavingsPct']),
             '/'.join(str(p['fees']['total']['scenarioWinTieLoss'][k]) for k in ('win', 'tie', 'loss'))]
            for p in panel['comparisons'].values()]


def completion_text(summary, method, language='ko'):
    labels = PANELS_ZH if language == 'zh' else PANELS
    return '; '.join(f"{labels[panel]} {summary['panels'][panel]['byMethod'][method]['completeEpisodeCount']}/"
                     f"{summary['panels'][panel]['byMethod'][method]['recordedEpisodeCount']}" for panel in PANELS)


def panel_gates(panel):
    """Observed necessary gates, never a production or universal guarantee."""
    no_losses = all(panel['comparisons'][name]['completion']['lostEpisodeCount'] == 0
                    for name in ('v10VsHiGHS', 'v10VsV9', 'v10VsV7', 'v10VsV5'))
    native = panel['comparisons']['v10VsHiGHS']
    intervals = {key: native['fees'][key]['uncertainty']['weightedSavingsPct95CI'] for key in ('total', 'incremental')}
    fee_safe = all(isinstance(value, list) and len(value) == 2 and all(finite(v) for v in value) and
                   value[1] >= 0 for value in intervals.values())
    revisions = native['revisions']
    known_revisions = native['fees']['incremental']['pairedEpisodeCount'] > 0 and revisions['pairedEpisodeCount'] == native['fees']['incremental']['pairedEpisodeCount']
    fewer = (known_revisions and finite(revisions['candidateObservedSum']) and finite(revisions['baselineObservedSum']) and
             revisions['candidateObservedSum'] < revisions['baselineObservedSum'])
    native_stability = native['completion']['lostEpisodeCount'] == 0 and fewer and fee_safe
    return {'zeroCompletionLossAgainstEveryBaseline': no_losses, 'zeroNativeCompletionLoss': native['completion']['lostEpisodeCount'] == 0,
            'strictlyFewerNativePairedRevisions': fewer, 'bothFeeIntervalsAvailableAndNotWhollyNegative': fee_safe,
            'conditionalNativeStabilityGate': native_stability,
            'conditionalFeeAdvantage': {key: isinstance(value, list) and len(value) == 2 and value[0] > 0
                                       for key, value in intervals.items()}}


def conclusion(e):
    gates = [panel_gates(e['summary']['panels'][panel]) for panel in PANELS]
    if all(gate['zeroCompletionLossAgainstEveryBaseline'] and gate['conditionalNativeStabilityGate'] for gate in gates):
        return '이 유한 평가에서 선언된 관측 문턱을 충족했습니다. 실제 운영 기본값 적합성은 별도 검증 대상입니다.'
    return '선언된 관측 문턱을 모두 충족하지 못했습니다. V10은 실험 선택 항목이며 기본값 권고 조건을 충족하지 않았습니다.'


def pareto_rows(e):
    """Final producer field contract is verified before rendering."""
    rows = []
    for panel in PANELS:
        common = e['summary']['panels'][panel]['jointCompletePareto']
        for method in METHODS:
            values = common['byMethod'][method]
            rows.append([panel, LABELS[method], f"{common['jointScenarioCount']} / {common['jointEpisodeCount']}",
                         fnum(values['scenarioMeanTotalCostSumKRW']), fnum(values['scenarioMeanIncrementalCostSumKRW']),
                         fnum(values['scenarioMeanRevisionCountSum'])])
    return rows


def build_slides(e):
    p, s, r, previous = e['protocol'], e['summary'], e['review'], e['previousV9']['summary']
    n, repeats, methods = len(e['index']['cases']), len(p['algorithmRepeats']), len(p['methods'])
    episodes = r['summary']['episodes']
    observed = ('V10 ' + completion_text(s, 'stableHiGHSV10') + ' / native ' + completion_text(s, 'highsV8'))
    slides = []
    def add(title, text, *, table=None, chart=None, flow=None, note='', english=''):
        slides.append({'title': title, 'text': text, 'table': table, 'chart': chart, 'flow': flow, 'note': note, 'english': english})
    add('같은 요금이라면 유효한 날짜를\n그대로 유지할 수 있을까요?',
        'PortPlan V10은 원래 native HiGHS가 계산한 새 후보와 현재 유효 계획을 같은 관측 정보로 재검사합니다. 실제 요금이 엄격히 낮을 때만 계획을 교체합니다.',
        flow=['현재 유효 계획', 'native 후보·같은 요금?', '동률 유지 / 엄격 절감 교체'],
        note=conclusion(e)+' '+observed+'. 알려진 합성 주문의 일 단위 모형이며 실제 기업 절감·우승 확률이 아닙니다.',
        english='V10 retains a currently feasible cached schedule unless one native HiGHS candidate has strictly lower independently reconstructed current-snapshot joint fees. '+observed+'. The declared observed gates determine whether V10 stays experimental; they are not a production recommendation or first-place probability.')
    add('반출·적출·반납과 요금을\n하나의 날짜 장부로 연결합니다.',
        '반출·반납은 각각 운송 작업을 사용합니다. 적출은 연속 창고 일수를 차지하고, 반출부터 반납 전까지 체류 공간을 점유합니다.',
        table=(['모형', '계산할 항목', '미포함 범위'], [
            ['반출 p / 적출 u / 반납 r', '방출·납기·선후행·실행 잠금', '도로 경로·예약·기사 근무시간'],
            ['일별 운송·창고·체류 [p,r)', '용량·연속 점유·container/TEU', '시간대별 실운영'],
            ['PNC 보관료 / Maersk D&D', '공개 요율·명시한 기산일 해석', '계약·실제 청구·세금 검증']]),
        note='기존 주문은 시작 시 모두 알려집니다. 비용 목표는 선언한 두 요금이며 신규 주문 도착이나 전체 물류 원가가 아닙니다.',
        english='The unchanged daily model covers pickup, consecutive unloading, return, exact past locks, transport and staging, plus explicitly interpreted PNC/Maersk tariffs. It does not model routes, appointments, driver hours, incoming orders, private contracts, invoices, taxes or all logistics costs.')
    add('V9의 한계도\n수정하지 않은 이전 결과로 남깁니다.',
        'V9는 V8의 늦은 후보 전달을 개선했지만 V7에 대한 일관된 향상을 보이지 않았습니다. 이를 V10의 다른 새 주문 입력과 빼서 개선량으로 만들지 않습니다.',
        table=(['이전 V9 평가 · 별도 입력', 'V9 완주', 'V7 완주'], [
            [PANELS[panel], f"{previous['panels'][panel]['byMethod']['adaptiveV9']['completeEpisodeCount']}/216",
             f"{previous['panels'][panel]['byMethod']['costPolishV7']['completeEpisodeCount']}/216"] for panel in PANELS]),
        note='이전 V8/V9 최초 입력·반복·실패·비용 악화는 그대로 보존합니다. 같은 새 V10 입력에서 함께 실행한 정책끼리만 이번 차이를 계산합니다.',
        english='The prior V9 first run remains unchanged: '+completion_text(previous, 'adaptiveV9')+'; prior V7 '+completion_text(previous, 'costPolishV7')+'. These are separate older input cohorts. V10 effects must compare methods run on the same new V10 input and repeat, not subtract counts across versions or cohorts.')
    add('V10: 단 한 번의 native MILP와\n엄격한 비용 교체 규칙입니다.',
        '현재 cache를 재가격·물리/잠금 검사한 뒤 원래 native 서비스를 한 번 호출합니다. 현재 유효 cache보다 추가 요금이 엄격히 낮은 적시 후보만 교체합니다.',
        flow=['cache 재가격·검사', 'native MILP 1회', '엄격 절감만 교체'],
        note='비교 epsilon은 고정 0.000001원입니다. 동일/더 높은 요금은 유지합니다. cache가 없거나 무효이면 기존 복구 방식으로 완전 후보를 구합니다.',
        english='One native MILP call uses the original same-model service and budget. A valid cache is replaced only when candidate joint incremental fees are lower by more than the fixed KRW 1e-6 numerical epsilon. Ties and costlier proposals retain the cache; absent or invalid cache follows original recovery.')
    add('날짜 유지의 의미와\n보장하지 않는 것을 구분합니다.',
        '같은 비용의 후보를 거부하면 그 호출의 불필요한 변경을 피할 수 있습니다. 엄격히 더 저렴한 후보는 여러 날짜를 바꿀 수 있고 이후 경로도 달라집니다.',
        table=(['구현한 것', '구현/증명하지 않은 것'], [
            ['동일 요금의 유효 cache 유지', '전체 경로 최소 변경'],
            ['strict fee replacement', '두 번째 단계 사전식/lexicographic 최적화'],
            ['같은 native 요금 목표·1회 solve', '변경 1회를 임의 원화 가중치로 환산'],
            ['현재 관측·실행 잠금 검사', '미래 충격의 완주·비용 보장']]),
        note='더 적은 날짜 변경과 비용·완주의 교환관계는 실험에서 확인합니다. 지역 선택 규칙을 전역 안정성 정리로 표현하지 않습니다.',
        english='Retaining an equal-fee valid cache avoids unproductive changes at that call, but a strictly cheaper candidate may change many dates and the future trajectory. V10 is not a global minimum-revision algorithm, a second-stage lexicographic solve, or a KRW-weighted revision objective.')
    add('native 기준과 웹 시연을\n정확히 구분합니다.',
        'V10과 원래 HiGHS는 같은 native Python 서비스·모형·수치 예산을 사용합니다. 초기화, IPC, 날짜 복사·재가격·검사는 전체 기한에 포함합니다.',
        table=(['항목', '실행 범위'], [
            ['V10 / 원래 native HiGHS', '오프라인 연구 worker와 Python/HiGHS'],
            ['웹 브라우저 계획기', '기존 GLPK와 합성 입력 시뮬레이터'],
            ['일별 / 전체 냉 기한', '200ms / 부모 5,000ms'],
            ['검산', '기록 시각 일관성·독립 날짜/요율 재구성']]),
        note='웹에서 native HiGHS를 실행한다고 주장하지 않습니다. 내부 기록 검산은 외부 시계 인증·같은 CPU 시간·절대 실시간 보장이 아닙니다.',
        english='V10 and original native HiGHS share one Python service, model and numeric configuration. Native startup, IPC and outer copy/repricing/audits are charged to the cold deadline. The browser continues to use GLPK; native experiment results do not establish a browser-native deployment or external real-time certification.')
    add('기존 생성·충격 메커니즘에\n새 주문 시드를 사전 고정했습니다.',
        f"{n}개 업무 시나리오 × {methods}개 정책 × {repeats}개 알고리즘/타이밍 반복 = {episodes:,}개 episode. 반복은 새 독립 업무 표본이 아닙니다.",
        table=(['패널 / 반복', '범위'], [
            ['동일 생성분포 72개', '기존 생성기·기존 충격; 새 주문 시드'],
            ['스트레스 72개', '기존 V8/V9 충격 catalog; 새 주문 시드'],
            ['모든 정책·반복', '순차 냉 worker; 순서 순환; 모든 실패 보존'],
            ['해석', '패널별 보고; 입력 교체·실패 재추첨 없음']]),
        note='구현·예산·분석을 새 입력 생성 전에 고정합니다. 144개 업무 시나리오를 432개 독립 사례로 세지 않습니다.',
        english=f'{n} new order scenarios x {methods} fixed methods x {repeats} retained repeats = {episodes} episodes. Generator and shock mechanisms are unchanged; only the order seeds are new. Cases, repeats, negative outcomes and deadlines are retained. Separate panels and business-scenario bootstrap avoid counting 432 repeats as independent scenarios.')
    for panel in PANELS:
        ps = s['panels'][panel]
        add(PANELS[panel]+'\n조건부 비용보다 전체 완주를 먼저 봅니다.',
            '완주는 물리 행동과 최종 검사를 부모 기한 내 수락한 경우입니다. 미확보 이력과 비용은 null이며 물리 불가능이나 0원으로 바꾸지 않습니다.',
            chart=[{'label':LABELS[m],'value':ps['byMethod'][m]['completeEpisodeCount'],
                    'denominator':ps['byMethod'][m]['recordedEpisodeCount'],'unknown':ps['byMethod'][m]['executionHistoryUnavailableCount'],
                    'all3':ps['byMethod'][m]['fullyCompletedScenarioCount']} for m in METHODS],
            note='episode 완주 수, 세 반복 모두 완주한 시나리오, 이력 미확보를 분리합니다. 조건부 Pareto가 전체 완주 손실을 덮지 않습니다.',
            english='Full-sample completion precedes conditional cost analysis. Completed episodes, business scenarios completing all three repeats, and unavailable histories remain separate. A favorable conditional Pareto position cannot override completion losses.')
        add(PANELS[panel]+'\n같은 공동 완주 반복의 요금만 비교합니다.',
            '양수 절감은 기준 − 후보입니다. 시나리오 내 일치하는 공동 완주 반복을 평균한 뒤 시나리오를 동일 가중합니다.',
            table=(['후보 / 기준','시나리오 / 쌍','총요금 절감','추가요금 절감','낮음/동일/높음'],comparison_rows(ps)),
            note='비교별 분모가 다릅니다. 구간이 0을 포함하면 안정적인 비용 우위를 주장하지 않습니다. 모든 악화 ID와 조건부 구간은 보고서에 보존합니다.',
            english='Fees use matching jointly complete repeats, averaged inside each scenario then equally weighted by scenario. Pair-specific denominators differ. A fee interval containing zero supports no stable fee advantage; all adverse fee identifiers and conditional intervals remain disclosed.')
    losses=[];changes=[]
    for panel, ps in s['panels'].items():
        for name,pair in ps['comparisons'].items():
            f=pair['fees']['total'];rev=pair['revisions']
            losses.append([panel,name,len(f['negativeScenarioIds']),fnum(f['worstScenarioLossKRW']),
                           f"{pair['completion']['gainedEpisodeCount']} / {pair['completion']['lostEpisodeCount']}"])
            changes.append([panel,name,rev['pairedEpisodeCount'],fnum(rev['baselineObservedSum']),fnum(rev['candidateObservedSum'])])
    add('완주 손실·최악 비용 악화를\n모든 비교에서 보존합니다.',
        '완주 얻음/잃음은 전체 일치 반복의 episode 기준입니다. 최악 비용 증가는 비교 가능한 시나리오 평균 기준입니다.',
        table=(['패널','비교','비용 악화 시나리오','최악 증가 ₩','완주 얻음/잃음'],losses),
        note='어느 이전 기준에든 완주 손실이 있으면 기본값 관측 자격을 얻지 못합니다. 모든 총/추가 요금 악화·완주 손실 ID를 보고서에 남깁니다.',
        english='Every comparison retains completion gains/losses and worst scenario-average fee increases. Any completion loss to any previous baseline fails the necessary observed default-eligibility gate. Every negative total/incremental fee scenario/repeat ID and lost completion ID remains in the report.')
    add('변경 수는 같은 공동 완주 집합에서\n별도 단위로 비교합니다.',
        '변경 1회는 아직 실행하지 않은 반출·적출 시작·반납 날짜 하나의 수정입니다. 비용과 합친 점수나 임의 원화 가중치를 만들지 않습니다.',
        table=(['패널','비교','공동 완주 반복','기준 변경 합','후보 변경 합'],changes),
        note='변경 수의 미확보 값은 0이 아닙니다. native 안정성 주장은 같은 native 공동 완주 쌍에서 엄격히 적은 변경과 비용 구간 조건을 함께 요구합니다.',
        english='Revisions count edits of pending pickup, unload-start or return dates. Compare them on exactly the same jointly complete repeats, separately from money. Unknown counts remain unknown. No invented KRW weight, unified score or cross-denominator ranking is used.')
    add('다섯 방법이 모두 완주한 집합에서\n요금·변경의 Pareto를 따로 봅니다.',
        '다섯 방법의 동일 반복이 모두 완주한 부분집합만 사용합니다. 시나리오별 공동 반복을 평균한 뒤 시나리오를 동일 가중합니다.',
        table=(['패널','정책','시나리오/반복','총요금 평균합 ₩','추가요금 평균합 ₩','변경 평균합'],pareto_rows(e)),
        note='이 선택된 집합의 요금과 변경은 별도 축입니다. 우위는 양쪽 비악화와 한쪽 엄격 개선일 때만이며, 전체 완주 손실·선택 편향은 남습니다.',
        english='The all-five-method jointly complete matching-repeat subset has one disclosed common denominator per panel. Equal-scenario-weight total/incremental fees and revision counts remain separate axes. Dominance requires no worse on both and strictly better on one; no weighted total score or pooled panel leaderboard is constructed.')
    native=r['nativeStabilitySelection'];rc=r['summary']
    add('원 날짜·요금·선택 근거와\n통계를 독립 검산합니다.',
        f"Python이 {rc['episodes']:,}개 episode를 확인했습니다. V10 선택 호출 {native['calls']:,}개, 동률 유지 {native['feeTiesRetained']:,}개, 엄격 교체 {native['strictReplacements']:,}개를 원 날짜에서 재구성합니다.",
        table=(['검산','경계'],[
            ['cache / proposal / selected 원 날짜','현재 요율·물리 자원·실행 잠금 재계산'],
            ['동률 유지 / 엄격 교체','완전 proof와 기록된 기한 일관성'],
            ['bootstrap / 공동 Pareto / 관측 문턱','별도 Python 통계 검산; 같은 index SHA'],
            ['현장·전역 최적·외부 시계 인증','확인하지 않음']]),
        note='조건부 비용 구간은 공동 완주 선택 편향을 제거하지 않습니다. 내부 독립 검산은 제3자 인증이나 운영 보장이 아닙니다.',
        english='Separate read-only Python implementations reconstruct original cache/proposal/selected dates, tariffs, resources, locks, strict replacement, intervals, common-subset Pareto and observed gates. They verify recorded timing consistency, not external clock authenticity, future policy optimality or operational suitability.')
    add('실제 기능과 연구 기준을\n같은 성과로 합치지 않습니다.',
        '웹은 합성 주문·용량 사건의 날짜별 관측·잠금·실패·내보내기를 제공합니다. 명시적 가격·권한의 자원 메뉴는 별도 현재 관측 복구 도구입니다.',
        flow=['합성 웹 시뮬레이션','native 연구 증거','공개 출처·독립 검산'],
        note='작은 교육 예제는 공식 144개 입력이나 실제 현장 성과가 아닙니다. 자원 제안은 구매·배차가 아니며 V10 자체가 전역 최소 변경을 보장하지 않습니다.',
        english='Browser simulation, finite-menu counterfactual recovery and native research comparisons have separate scopes. Teaching fixtures are not the 144 formal business scenarios, supplier commitments or field savings. V10 does not purchase resources, dispatch vehicles or guarantee globally minimal revisions.')
    add('공개 자료로 개발을 계속하고\n남은 검증은 범위별로 기록합니다.',
        '공개 출처·공개 요율·재현 가능한 실험이 기본 개발 경로입니다. 기업 자료·연락처·현장 시험은 선택 확장이며 이 경로의 선행 조건이 아닙니다.',
        table=(['검증 기록','현재 상태'],[
            ['알고리즘 연구','완결 실험·두 검산; 선언된 문턱은 별도 판정'],
            ['브라우저 실제 조작','미검증 · 지원 연결 입구 없음'],
            ['현장 적합성 / 견적 / 효과','선택 확장 · 미실시; 공개 연구의 선행 조건 아님'],
            ['본인 설명 / 질의 / 재현','빈 양식 제공 · 미확인'],
            ['등록 / 당해 AI 규정 / 서약','명확히 확인되지 않음; 자료 제출 11/6 18:00 KST 해석']]),
        note='공개 등록은 정원 소진 시 종료하며 사용자의 자리·자격은 미확인입니다. 당해 AI 규정과 본인 이해를 대신 인증하지 않습니다.',
        english='Public sources, public tariffs and reproducible experiments are the default development route. Enterprise records, contacts and a business pilot are optional extensions, not prerequisites. Browser interaction lacks a supported connection; algorithm replacement gates and personal readiness remain separate unresolved items. The materials deadline is November 6 at 18:00 interpreted as KST; capacity-based registration, entrant eligibility and current-year AI conditions are not verified.')
    gate_rows=[]
    for panel in PANELS:
        g=panel_gates(s['panels'][panel])
        gate_rows.append([PANELS[panel],str(g['zeroCompletionLossAgainstEveryBaseline']),str(g['zeroNativeCompletionLoss']),
                          str(g['strictlyFewerNativePairedRevisions']),str(g['bothFeeIntervalsAvailableAndNotWhollyNegative'])])
    add('사전 문턱과 전체 원본을\n함께 보고 다음 판단을 내립니다.',
        conclusion(e),
        table=(['패널','모든 기준 완주 손실 0','native 완주 손실 0','동일 쌍 변경 감소','두 비용 구간 조건'],gate_rows),
        note='관측 문턱 통과도 운영 기본값 보장이 아닙니다. 전체 입력·모든 반복·실패·proof·고정 source·두 검산을 SHA로 전달하며 최초 결과를 덮어쓰지 않습니다.',
        english='The declared observed gate results determine whether V10 remains experimental. Passing finite-sample gates still does not establish a production default or universal non-regression. Complete inputs, every repeat, failures, native proofs, frozen source and both independent reviews are bound by SHA and first-run results remain immutable.')
    require(len(slides)==18, 'Main deck must contain exactly 18 pages')
    return slides


def markdown_table(headers, rows):
    cell=lambda value:str(value).replace('|', '\\|').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(map(cell,headers))+' |','| '+' | '.join(['---']*len(headers))+' |',
                      *['| '+' | '.join(map(cell,row))+' |' for row in rows]])


def detailed_pairs(e, language='ko'):
    zh = language == 'zh'
    lines = ['## 完整配对费用、改期和不利结果' if zh else '## 전체 비교 요금·변경·불리한 결과', '',
             '节省 = 基准 − 候选。每个情景内先平均共同完成的匹配重复，再对情景等权。各比较分母不同，失败费用为 null。费用区间以共同完成集合为条件，独立重算不能消除完成选择偏差。'
             if zh else '절감 = 기준 − 후보. 시나리오 내 공동 완주 반복을 먼저 평균한 뒤 시나리오를 동일 가중합니다. 비교별 분모가 다르고 실패 비용은 null입니다. 비용 구간은 선택된 공동 완주 집합에 조건부이며 독립 검산이 완주 선택 편향을 제거하지 않습니다.', '']
    for panel, ps in e['summary']['panels'].items():
        lines.extend(['### '+(PANELS_ZH if zh else PANELS)[panel], '',
            markdown_table(['method','complete / all','all 3 repeats','unknown history','observed revisions / unknown','parent mean ms'],
                [[method,f"{ps['byMethod'][method]['completeEpisodeCount']} / {ps['byMethod'][method]['recordedEpisodeCount']}",
                  ps['byMethod'][method]['fullyCompletedScenarioCount'],ps['byMethod'][method]['executionHistoryUnavailableCount'],
                  f"{ps['byMethod'][method]['revisionCount']['observedSum']} / {ps['byMethod'][method]['revisionCount']['unknownEpisodeCount']}",
                  fnum(ps['byMethod'][method]['wallTimeMs']['mean'],2)] for method in METHODS]),''])
        for name,pair in ps['comparisons'].items():
            lines.extend(['#### '+name+' · '+pair_label(pair),''])
            fee_rows=[]
            for key in ('total','incremental'):
                f=pair['fees'][key]
                fee_rows.append([key,f"{f['pairedScenarioCount']} / {f['pairedEpisodeCount']}",fnum(f['baselineScenarioMeanSumKRW'],2),
                    fnum(f['candidateScenarioMeanSumKRW'],2),fnum(f['scenarioWeightedSavingsKRW'],2),pct(f['weightedSavingsPct']),
                    '/'.join(str(f['scenarioWinTieLoss'][k]) for k in ('win','tie','loss')),fnum(f['worstScenarioLossKRW'],2),
                    f"{f['zeroBaselineScenarioCount']} / {f['zeroBaselineEpisodeCount']}",ci(f['uncertainty']['weightedSavingsPct95CI'])])
            lines.extend([markdown_table(['fee','scenario / repeat','baseline scenario-mean sum KRW','candidate scenario-mean sum KRW',
                'savings KRW','savings %','lower/same/higher','worst increase KRW','zero baseline scenario/repeat','conditional 95% CI %'],fee_rows),''])
            c=pair['completion'];rev=pair['revisions']
            lines.extend([f"{'完成增加/损失' if zh else '완주 얻음/잃음'}: {c['gainedEpisodeCount']} / {c['lostEpisodeCount']} episodes. "
                f"{'情景平均完成率差' if zh else '시나리오 평균 완주율 차이'}: {fnum(c['meanScenarioRateGain'],6)}; 95% CI {ci(c['uncertainty']['meanScenarioRateGain95CI'])}.",
                f"{'同一配对完成重复的改期合计' if zh else '같은 공동 완주 반복의 변경 합'}: {rev['baselineObservedSum']} → {rev['candidateObservedSum']}; n={rev['pairedEpisodeCount']}.", ''])
            for key in ('total','incremental'):
                f=pair['fees'][key];u=f['uncertainty']
                lines.extend([f"{key}: raw matched episode sums KRW {fnum(f['baselineEpisodeSumKRW'],2)} → {fnum(f['candidateEpisodeSumKRW'],2)}; "
                    f"mean scenario saving KRW 95% CI {ci(u['meanSavings95CI'])}; bootstrap scenario n={u['scenarioCount']}, draws={u['draws']}.",''])
                for field in ('negativeScenarioIds','negativeEpisodeIds'):
                    lines.extend([f'**{key} / {field} — all retained**',''])
                    lines.extend(['- `'+identifier+'`' for identifier in f[field]] or ['none / 없음 / 无']);lines.append('')
            for field in ('gainedEpisodeIds','lostEpisodeIds','gainedFromKnownFailureEpisodeIds','gainedFromUnknownHistoryEpisodeIds'):
                lines.extend(['**'+field+' — all retained**',''])
                lines.extend(['- `'+identifier+'`' for identifier in c[field]] or ['none / 없음 / 无']);lines.append('')
        common=ps['jointCompletePareto']
        lines.extend(['#### '+('五方法共同完成集：独立费用/改期轴' if zh else '다섯 방법 공동 완주: 요금/변경 별도 축'),' ',
            f"jointScenarioCount={common['jointScenarioCount']}; jointEpisodeCount={common['jointEpisodeCount']}. "
            + ('先平均每情景共同重复，再将情景均值求和；改期均值轴与原始 episode 改期合计分开。无共同样本或未知改期为 null。'
               if zh else '시나리오 내 공동 반복을 평균한 뒤 시나리오 평균을 합합니다. 변경 평균 축과 원 episode 변경 합은 별도입니다. 공동 표본 없음/미확보 변경은 null입니다.'),'',
            markdown_table(['method','total scenario-mean sum KRW','incremental scenario-mean sum KRW','revision scenario-mean sum',
                            'raw episode revisions','full-sample complete / all'],
                [[method,fnum(common['byMethod'][method]['scenarioMeanTotalCostSumKRW'],2),
                  fnum(common['byMethod'][method]['scenarioMeanIncrementalCostSumKRW'],2),
                  fnum(common['byMethod'][method]['scenarioMeanRevisionCountSum'],3),
                  fnum(common['byMethod'][method]['episodeRevisionCountSum']),
                  f"{common['byMethod'][method]['fullSampleCompleteEpisodeCount']} / {common['byMethod'][method]['fullSampleEpisodeCount']}"] for method in METHODS]),'',
            '```json',json.dumps(common['frontiers'],ensure_ascii=False,indent=2),'```','',
            '#### '+('预声明观察门槛' if zh else '사전 선언한 관측 문턱'),'',
            '```json',json.dumps(ps['observedDecisionGates'],ensure_ascii=False,indent=2),'```',''])
    return lines


def blank_acceptance_forms(e):
    index_hash=e['wrapper']['formalIndex']['sha256']
    field=['# PortPlan V10 · 现场验收空白记录 / 현장 검증 빈 기록','',
        '状态：尚未填写，未验收。此模板不构成企业背书、访谈、采购许可或实际节省证据。', '',
        '| 记录项 | 实际填写内容 | 证据文件/来源 | 实际日期 KST |','| --- | --- | --- | --- |',
        '| 业务单位及有权受访人 | | | |','| 当前人工流程与实际痛点 | | | |',
        '| 脱敏真实订单范围及数据使用依据 | | | |','| 实际合同/报价/费率差异 | | | |',
        '| 同一输入的人工基准与系统建议 | | | |','| 实际资源可用性、权限及执行限制 | | | |',
        '| 失败、费用恶化、额外改期及人工处理 | | | |','| 实际账单/结果及未观察部分 | | | |',
        '| 验收人结论、保留意见及签署 | | | |','',
        '不得把合成结果、AI输出或模板空栏标为现场通过；不得将采购建议当作实际购买。',
        '', '研究参考 formal index SHA-256：`'+index_hash+'`。该绑定不代替现场证据。']
    entrant=['# PortPlan V10 · 本人答辩验收空白记录 / 본인 발표 빈 기록','',
        '状态：尚未填写，未验收。AI生成资料不证明参赛者本人理解、开发贡献、资格或答辩熟练度。','',
        '| 实际检查项 | 本人回答/演示记录 | 复核者及证据 | 实际日期 KST |','| --- | --- | --- | --- |',
        '| 独立说明 p/u/r、连续卸箱与体积占用 | | | |','| 解释公开费率、起算与沉没费用 | | | |',
        '| 说明 V10 单次 MILP 与严格省费替换 | | | |','| 解释同费保留并非全局最少改期 | | | |',
        '| 读出完成损失、不同配对分母与条件 CI | | | |','| 解释共同完成 Pareto 与选择偏差 | | | |',
        '| 现场复现输入、失败、proof 与导出 | | | |','| 具体列出本人实际贡献与 AI 工作 | | | |',
        '| 核对比赛 AI 规则、报名资格和声明 | | | |','| 实际问答盲测、未能回答问题和补练 | | | |',
        '| 本人/复核者结论与保留意见 | | | |','',
        '不得预填通过、虚构第三方复核或替本人签名。比赛提交与声明应由有权者完成。',
        '', '研究参考 formal index SHA-256：`'+index_hash+'`。该绑定不证明本人能力。']
    return '\n'.join(field)+'\n','\n'.join(entrant)+'\n'


def native_result_paragraphs(e, language):
    paragraphs=[]
    for panel,ps in e['summary']['panels'].items():
        pair=ps['comparisons']['v10VsHiGHS'];r=pair['revisions'];c=pair['completion']
        reduction = 100*(r['baselineObservedSum']-r['candidateObservedSum'])/r['baselineObservedSum'] if r['baselineObservedSum'] else None
        total,inc=pair['fees']['total'],pair['fees']['incremental']
        lost_scenarios=len({identifier.rsplit('::repeat-',1)[0] for identifier in c['lostEpisodeIds']})
        complete=f"{ps['byMethod']['stableHiGHSV10']['completeEpisodeCount']}/{ps['byMethod']['stableHiGHSV10']['recordedEpisodeCount']}"
        baseline=f"{ps['byMethod']['highsV8']['completeEpisodeCount']}/{ps['byMethod']['highsV8']['recordedEpisodeCount']}"
        if language=='zh':
            text=(f"{PANELS_ZH[panel]}：V10完成{complete}，原生HiGHS完成{baseline}；匹配重复完成增加/损失为{c['gainedEpisodeCount']}/{c['lostEpisodeCount']}。"
                  f"损失涉及{lost_scenarios}个业务情景，不能把同一情景的多次重复当独立业务失败。共同完成{r['pairedEpisodeCount']}对中，改期{r['baselineObservedSum']}→{r['candidateObservedSum']}，下降{fnum(reduction,2)}%。"
                  f"总费用节省{fnum(total['weightedSavingsPct'],8)}%，条件95%区间{ci(total['uncertainty']['weightedSavingsPct95CI'])}；"
                  f"增量费用节省{fnum(inc['weightedSavingsPct'],8)}%，条件95%区间{ci(inc['uncertainty']['weightedSavingsPct95CI'])}。"
                  '区间跨零时只能说未观察到稳健费用优势；负节省点估计代表候选更贵。')
        elif language=='ko':
            text=(f"{PANELS[panel]}: V10 완주 {complete}, 원래 native {baseline}; 완주 얻음/잃음 {c['gainedEpisodeCount']}/{c['lostEpisodeCount']}. "
                  f"손실 업무 시나리오 {lost_scenarios}개; 같은 시나리오 반복을 독립 사례로 세지 않습니다. 공동 완주 {r['pairedEpisodeCount']}쌍의 날짜 변경 "
                  f"{r['baselineObservedSum']}→{r['candidateObservedSum']} ({fnum(reduction,2)}% 감소). 총요금 절감 {fnum(total['weightedSavingsPct'],8)}%, "
                  f"조건부 95% CI {ci(total['uncertainty']['weightedSavingsPct95CI'])}; 추가요금 절감 {fnum(inc['weightedSavingsPct'],8)}%, "
                  f"CI {ci(inc['uncertainty']['weightedSavingsPct95CI'])}. 0을 포함한 구간은 안정적 비용 우위가 아니며 음의 절감 점추정은 더 비쌈을 뜻합니다.")
        else:
            text=(f"{panel}: V10 completed {complete} versus original native {baseline}; gained/lost completion repeats {c['gainedEpisodeCount']}/{c['lostEpisodeCount']}, "
                  f"with losses in {lost_scenarios} distinct business scenarios. On the same {r['pairedEpisodeCount']} jointly completed repeats, revisions "
                  f"fell {r['baselineObservedSum']} to {r['candidateObservedSum']} ({fnum(reduction,2)}%). Total-fee savings {fnum(total['weightedSavingsPct'],8)}%, "
                  f"conditional 95% CI {ci(total['uncertainty']['weightedSavingsPct95CI'])}; incremental-fee savings {fnum(inc['weightedSavingsPct'],8)}%, "
                  f"CI {ci(inc['uncertainty']['weightedSavingsPct95CI'])}. An interval spanning zero supports no stable fee advantage; a negative savings point estimate means higher candidate fees.")
        paragraphs.extend([text,''])
        v7=ps['comparisons']['v10VsV7'];vf=v7['fees']['incremental'];vc=v7['completion']
        if language=='zh':
            text=(f"相对V7的另一个配对集：增量两费条件节省{fnum(vf['weightedSavingsPct'],3)}%，"
                  f"95%区间{ci(vf['uncertainty']['weightedSavingsPct95CI'])}；分母{vf['pairedScenarioCount']}个情景/{vf['pairedEpisodeCount']}对。"
                  f"全样本完成增加/损失{vc['gainedEpisodeCount']}/{vc['lostEpisodeCount']}，后者是V7完成而V10失败的重复。"
                  '这项条件节省与相对原生HiGHS无稳健费用优势属于不同基准和分母，不能混为统一排名或抵消完成损失。')
        elif language=='ko':
            text=(f"V7 대비 별도 공동 완주 집합의 추가 두 요금 절감은 {fnum(vf['weightedSavingsPct'],3)}%, "
                  f"95% CI {ci(vf['uncertainty']['weightedSavingsPct95CI'])}; 분모 {vf['pairedScenarioCount']}개 시나리오/{vf['pairedEpisodeCount']}쌍입니다. "
                  f"전체 완주 얻음/잃음 {vc['gainedEpisodeCount']}/{vc['lostEpisodeCount']}; 손실은 V7이 완주했지만 V10이 실패한 반복입니다. "
                  '이는 native 대비 비용 구간과 기준·분모가 다른 조건부 결과이며 전체 순위나 완주 손실 상쇄로 쓰지 않습니다.')
        else:
            text=(f"On the separate V7 paired set, incremental two-fee savings are {fnum(vf['weightedSavingsPct'],3)}%, "
                  f"95% CI {ci(vf['uncertainty']['weightedSavingsPct95CI'])}, over {vf['pairedScenarioCount']} scenarios/{vf['pairedEpisodeCount']} repeats. "
                  f"Full-sample gained/lost completions are {vc['gainedEpisodeCount']}/{vc['lostEpisodeCount']}; losses are repeats where V7 completed and V10 did not. "
                  'These conditional savings have a different baseline and denominator from the native comparison; they do not create a unified ranking or cancel completion losses.')
        paragraphs.extend([text,''])
    return paragraphs


def material_texts(e, slides):
    provenance='AI 작성 준비 자료. 참가자 본인의 창작·검증·이해·참가자격 확인서가 아닙니다.'
    report=['# PortPlan V10 · 동일 요금 유지와 native 계획 안정성','',provenance,'',conclusion(e),'',
            '완결된 새 주문 실험과 독립 물리/통계 검산에서 생성했습니다. 합성 일 단위 모델이며 실현 절감·기업 검증·전역 최소 변경이 아닙니다.','']
    english=['# PortPlan V10 — native strict-fee cache retention','',
        'AI-created preparation. No entrant-authorship, understanding, field-acceptance, eligibility or submission attestation.', '',
        'V10 makes one native MILP call and replaces a valid audited cache only for strictly lower current-snapshot joint fees by more than the fixed KRW 1e-6 epsilon. Ties retain the cache. This is not a second-stage lexicographic solve, globally minimum revision algorithm, or revision-to-KRW weighted objective.','']
    report.extend(native_result_paragraphs(e,'ko'))
    english.extend(native_result_paragraphs(e,'en'))
    for i,slide in enumerate(slides,1):
        report.extend([f"## {i:02d}. {slide['title'].replace(chr(10),' ')}",'',slide['text'],''])
        if slide['flow']:report.extend([' → '.join(slide['flow']).replace('\n',' / '),''])
        if slide['table']:report.extend([markdown_table(*slide['table']),''])
        if slide['chart']:report.extend([markdown_table(['정책','완주/전체','3회 모두 완주 시나리오','미확보 이력'],
            [[row['label'],f"{row['value']}/{row['denominator']}",row['all3'],row['unknown']] for row in slide['chart']]),''])
        report.extend([slide['note'],''])
        english.extend([f"## {i:02d}. {slide['title'].replace(chr(10),' ')}",'',slide['english'],''])
    report.extend(detailed_pairs(e,'ko'))
    zh=['# PortPlan V10：同费保留与原生排程稳定性','',
        'AI 编写的准备资料；不代表参赛者本人创作、理解、报名资格、现场验证或已提交。','',
        'V10 的改动是：原生 HiGHS 每日只求解一次；先对当前缓存按当天观测重新计费并检查物理约束与已执行锁定，只在新方案严格节省费用时替换。同费保留当前有效日期。数值容差固定为 0.000001 韩元。',
        '这不是全局最少改期算法、第二阶段词典序优化，也没有把一次改期乘以任意韩元权重。严格省费方案仍可能改变许多日期及未来轨迹。','',
        '## 当前结果的边界','',
        '两组各 72 个新订单情景，沿用旧生成器和扰动机制；每情景 3 次算法/计时重复、5 个方法，共 2,160 次完整冷启动。日预算 200ms，整段父进程 5s；初始化、IPC、检查及终端接收计入预算。重复不是新独立业务样本。',
        '对任何旧基准出现完成损失，均不满足预声明默认候选必要门槛。对原生 HiGHS 的条件稳定性结论还要求零完成损失、完全相同配对集内改期严格减少，以及总费/增量费两个可用区间均不完全落在负侧。区间含零只能说没有稳健费用优势。', '',
        '```json',json.dumps({panel:e['summary']['panels'][panel]['observedDecisionGates'] for panel in PANELS},ensure_ascii=False,indent=2),'```','',
        '未达到全部门槛时，V10 应保持实验选项。即使观察门槛通过，也不证明真实运营适合作为默认方案、未来不退步或夺冠概率。','',
        '## 保留既有负面历史','',
        '上一轮 V9：'+completion_text(e['previousV9']['summary'],'adaptiveV9','zh')+'；同轮 V7：'+completion_text(e['previousV9']['summary'],'costPolishV7','zh')+'。',
        '以上属于上一轮不同订单输入，不与 V10 的新输入计数相减作为提升。旧输入、失败和不利结果均保留。','',
        '## 公开资料路线与分别记录的验收范围','',
        markdown_table(['验收范围','当前证据/状态'],[
            ['算法研究','本报告的2,160次新订单实验、独立物理与统计复核；观察门槛逐项报告'],
            ['浏览器真实交互','尚未执行；当前缺少受支持浏览器连接入口。构建、DOM、Worker检查不能替代'],
            ['企业业务（可选）','尚未开展；实单、报价、反馈和账单不属于公开资料路线的必需前置条件，不反复索要'],
            ['本人复现与答辩','未验收；保留空白记录，AI不替本人填写理解、贡献、报名或通过']]),'',
        '材料截止为2026-11-06 18:00，按韩国大学语境解释为KST，官方页面没有独立时区标记；报名为满额即止，剩余名额与本人占位未知。当届完整AI规定、资格、誓约和本人准备仍须确认。',
        '默认以可核验公开资料、公开数据和明确标注的可复现实验继续开发。范围说明与可选空CSV见 `docs/acceptance-handoff-v10.zh.md`；浏览器范围见 `docs/browser-acceptance-v10.zh.md`；本轮官方复查见 `docs/rules-update-v10.md`。这些是准备与只读核查，不构成报名、提交、签署或现场验收通过。','']
    zh[4:4] = ['## 本轮直接结果','',*native_result_paragraphs(e,'zh')]
    if e.get('diagnosis'):
        zh.extend(['## 压力完成损失的事后只读诊断','',
            '相对原生HiGHS的3次压力完成损失来自同一个业务场景的3次重复。D5原生服务约4.3–4.5ms即返回严格不可行，未触及日或整段时限；这是已执行前缀的不可行性，不是迟到候选。',
            '独立Python从物理记录重建：C017已于D4开始连续2日卸货，因此D5必占1仓位；C010在D5仍未提卸、截止D5且未执行事件不能回到过去，因此D5也必占1。必要需求2大于当日公开容量1。',
            '证明仅针对V10走到D5的已执行前缀，不证明D0原问题无解、不代表基准另一前缀不可行，也不排除真实业务加仓位等补救。D1已经产生路径分歧；D4同费保留展示改期/未来可恢复性的取舍，未受控重放“只替换D4”，不能声称换一次就必定完成。',
            '完整原始绑定与必要条件见[只读损失诊断](v10-loss-diagnosis.zh.md)。该诊断未新增求解、在线重跑或新留出，也未改变首次结果。',''])
        report.extend(['## 압력 완주 손실의 사후 읽기 전용 진단','',
            'native 대비 압력 손실 3회는 한 업무 시나리오의 세 반복입니다. D5 native는 약 4.3–4.5ms에 엄격 불가능을 반환했으며 초과 시간이 아닙니다. C017의 D4 시작 2일 연속 적출과 C010의 D5 납기·미실행 하한이 D5 최소 창고 수요 2를 요구하지만 알려진 용량은 1입니다.',
            '이 증명은 V10의 이미 실행한 물리 prefix에 한정됩니다. D0부터 불가능하다는 뜻이나 D4 후보로 한 번 바꾸면 반드시 완주한다는 뜻이 아닙니다. 원 경로는 D1부터 달랐습니다. 새 solve·온라인 재실행·새 holdout 없이 원 최초 결과를 보존합니다.',
            '원 증거와 범위: [v10-loss-diagnosis.zh.md](v10-loss-diagnosis.zh.md).',''])
        english.extend(['## Read-only post-hoc pressure-loss diagnosis','',
            'The three native-comparison pressure losses are three repeats of one business scenario. At D5, native returned strict infeasibility in about 4.3–4.5ms, not a deadline loss. A Python necessary-condition proof reconstructs C017 already unloading across D4–D5 and C010 still pending with due D5: warehouse demand at D5 is at least two against capacity one.',
            'The proof applies only to the V10 executed prefix. Paths already diverged at D1; retaining equal-fee dates at D4 exposes a stability/recoverability tradeoff but no controlled replay proves that replacing D4 alone completes the path. The diagnosis adds no solver calls, online reruns or held-out cases and changes no first-run outcomes. See v10-loss-diagnosis.zh.md.',''])
    zh.extend(detailed_pairs(e,'zh'))
    hashes=markdown_table(['evidence input','SHA-256'],list(e['hashes'].items()))
    native_proof=json.dumps(e['review']['nativeStabilitySelection'],ensure_ascii=False,indent=2)
    formal_dir=Path(e['wrapper']['formalIndex']['path']).parent.as_posix()
    command='python scripts/build-defense-v10.py --require-results --bootstrap-review '+shlex.quote(e['materialOptions']['bootstrap-review'])
    reproduction=['```bash','python scripts/restore-evidence-v10.py',
        'python scripts/review-v10.py --run '+shlex.quote(formal_dir)+' --output artifacts/review-v10-reproduction.json',
        'python scripts/review-bootstrap-v10.py --run '+shlex.quote(formal_dir)+' --output artifacts/bootstrap-v10-reproduction.json',command,'```']
    for lines,title in ((report,'## 원본·재현·독립 선택 검산'),(zh,'## 原始证据、复现与独立选择复核')):
        lines.extend([title,'',hashes,'','```json',native_proof,'```','',*reproduction,'',
            'Native HiGHS 연구 기준은 브라우저 GLPK 실행과 별도입니다. / Native HiGHS 研究基准与浏览器 GLPK 分开。',
            '首次结果不可覆盖；新复现写入新路径。现场验收和本人答辩记录保持空白，不替用户填写通过。',''])
    english.extend(['## Exact result and proof bindings','',hashes,'','```json',native_proof,'```','',
        'See the Korean or Chinese full report for every adverse identifier, pair-specific fee denominator, paired revisions, conditional CI, common-subset Pareto and observed gate. No cross-denominator ranking, universal dominance, global minimum revisions, external clock certification, field ROI or first-place probability is claimed.','',*reproduction])
    contribution=['# PortPlan V10 · 실제 기여 / actual contribution','',provenance,'',
        markdown_table(['work','actual contributor','scope/status'],[
            ['Native strict-fee cache retention','AI','One unchanged native solve; no lexicographic stage or invented revision KRW'],
            ['Frozen prospective run and retained adverse outcomes','AI','Exact formal index and source bindings'],
            ['Original-date physical/native-proof review','AI','Internal independent Python; no solver replay'],
            ['Intervals/Pareto/observed-gate review','AI','Separate read-only Python reconstruction'],
            ['Korean HTML/PDF/editable PPTX and Chinese/English notes','AI','Same source slide objects and reviewed data'],
            ['Entrant personal creation/understanding/defense','pending','Blank form; never transferred from AI work'],
            ['Business interviews/invoices/field acceptance','not performed','Blank form; no fabricated validation or realized savings'],
            ['Competition registration/AI-rule eligibility/declaration','pending','No submission, signature or attestation performed']]),'',
        'GLPK, HiGHS, SciPy, original research/code/data notices and copyright remain attributed to their owners. Noto Sans CJK KR subset embedding retains SIL Open Font License attribution. Internal independent implementation is not third-party certification.','',hashes]
    field,entrant=blank_acceptance_forms(e)
    return '\n'.join(report)+'\n','\n'.join(english)+'\n','\n'.join(contribution)+'\n','\n'.join(zh)+'\n',field,entrant


def html_document(slides):
    parts=['<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>PortPlan V10</title><style>',
        'body{margin:0;background:#dce8e0;color:#173c40;font-family:"Noto Sans CJK KR","Noto Sans KR",sans-serif}.slide{box-sizing:border-box;width:1280px;min-height:720px;margin:24px auto;padding:44px 62px;background:#f7faf5;position:relative;page-break-after:always}h1{font-size:35px;line-height:1.4;margin:10px 0 22px;white-space:pre-line}.kicker{font-size:13px;letter-spacing:3px;color:#39826d}p{font-size:20px;line-height:1.7}table{width:100%;border-collapse:collapse;font-size:15px;margin:22px 0}th,td{text-align:left;border-bottom:1px solid #d3e0d3;padding:12px 10px;overflow-wrap:anywhere}th{background:#e6f0e4}.note{font-size:14px;background:#e8f2e5;padding:14px 18px;line-height:1.8}.flow{display:flex;gap:24px;align-items:stretch;margin:38px 0}.flow div{flex:1;background:#e0efe0;border-left:5px solid #36836c;padding:24px;font-size:23px;white-space:pre-line}.chart{margin:22px 0}.row{display:grid;grid-template-columns:240px 1fr 250px;align-items:center;gap:16px;margin:19px 0;font-size:17px}.track{height:21px;background:#dfebde}.fill{height:100%;background:#35826d}.footer{font-size:12px;letter-spacing:.5px;margin-top:25px;color:#638176}@media(max-width:1300px){.slide{width:100%;min-height:0;margin:12px auto;padding:30px}h1{font-size:29px}.row{grid-template-columns:180px 1fr 190px}}@media print{@page{size:1280px 720px;margin:0}body{background:white}.slide{width:1280px;height:720px;margin:0;overflow:hidden}}',
        '</style><body>']
    for i,slide in enumerate(slides,1):
        parts.extend([f'<section class="slide" id="slide-{i}"><div class="kicker">{i:02d} / PORTPLAN V10</div>',
                      '<h1>'+escape(slide['title'])+'</h1><p>'+escape(slide['text'])+'</p>'])
        if slide['flow']:parts.append('<div class="flow">'+''.join('<div>'+escape(text)+'</div>' for text in slide['flow'])+'</div>')
        if slide['table']:
            headers,rows=slide['table'];parts.append('<table><thead><tr>'+''.join('<th>'+escape(str(h))+'</th>' for h in headers)+'</tr></thead><tbody>')
            parts.extend('<tr>'+''.join('<td>'+escape(str(v))+'</td>' for v in row)+'</tr>' for row in rows);parts.append('</tbody></table>')
        if slide['chart']:
            parts.append('<div class="chart">')
            for c in slide['chart']:
                width=100*c['value']/c['denominator'] if c['denominator'] else 0
                parts.append(f'<div class="row"><span>{escape(c["label"])}</span><div class="track"><div class="fill" style="width:{width:.4f}%"></div></div><span>{c["value"]}/{c["denominator"]} · 3/3 {c["all3"]} · 미확보 {c["unknown"]}</span></div>')
            parts.append('</div>')
        parts.extend(['<div class="note">'+escape(slide['note'])+'</div>',f'<div class="footer">AI 작성 준비 자료 · 본인 검토·현장 검증 미확인 / {i:02d} / {len(slides):02d}</div></section>'])
    parts.append('</body></html>');return '\n'.join(parts)+'\n'


def register_noto(text):
    """Embed Noto KR; convert only the used CFF glyphs to quadratic outlines in memory."""
    from fontTools import subset
    from fontTools.fontBuilder import FontBuilder
    from fontTools.pens.cu2quPen import Cu2QuPen
    from fontTools.pens.ttGlyphPen import TTGlyphPen
    from fontTools.ttLib import TTFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont as PDFont
    path=Path('/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc')
    require(path.is_file(), 'Noto Sans CJK font is required for a readable Korean PDF')
    font=TTFont(path,fontNumber=1)
    require('KR' in font['name'].getDebugName(1), 'Expected Noto Korean font face')
    chars=set(ord(c) for c in text if not c.isspace());cmap=font.getBestCmap()
    require(not (chars-set(cmap)), 'Noto is missing required PDF glyphs: '+repr(chars-set(cmap)))
    subsetter=subset.Subsetter();subsetter.populate(text=text);subsetter.subset(font)
    order=font.getGlyphOrder();glyphset=font.getGlyphSet();glyf={}
    for name in order:
        pen=TTGlyphPen(glyphset);glyphset[name].draw(Cu2QuPen(pen,1.0,reverse_direction=True));glyf[name]=pen.glyph()
    builder=FontBuilder(font['head'].unitsPerEm,isTTF=True);builder.setupGlyphOrder(order);builder.setupCharacterMap(font.getBestCmap());builder.setupGlyf(glyf)
    builder.setupHorizontalMetrics(font['hmtx'].metrics);builder.setupHorizontalHeader(ascent=font['hhea'].ascent,descent=font['hhea'].descent)
    builder.setupNameTable({'familyName':'Noto Sans CJK KR PortPlan subset','styleName':'Regular','uniqueFontIdentifier':'PortPlanNotoKR','fullName':'PortPlan Noto KR subset','psName':'PortPlanNotoKR',
                          'copyright':'Noto Sans CJK. Copyright Adobe and Google. SIL Open Font License 1.1. Subset outlines converted for embedding.'})
    builder.setupOS2(sTypoAscender=font['hhea'].ascent,sTypoDescender=font['hhea'].descent,usWinAscent=font['hhea'].ascent,usWinDescent=-font['hhea'].descent)
    builder.setupPost();builder.setupMaxp();buffer=BytesIO();builder.save(buffer);buffer.seek(0);pdfmetrics.registerFont(PDFont('PortPlanNoto',buffer))


def pdf_document(slides):
    from reportlab.lib.colors import HexColor
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfgen import canvas
    from reportlab.platypus import Paragraph, Table, TableStyle
    all_text=json.dumps(slides,ensure_ascii=False)+' PORTPLAN V10 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789'
    register_noto(all_text)
    buffer=BytesIO();c=canvas.Canvas(buffer,pagesize=(1280,720),invariant=1,pageCompression=1)
    c.setTitle('PortPlan V10 · Korean competition preparation');c.setAuthor('AI-assisted PortPlan research; entrant verification pending')
    bg=HexColor('#f7faf5');ink=HexColor('#173c40');green=HexColor('#35826d');light=HexColor('#e6f0e4')
    def paragraph(text,x,y,w,size=20,leading=None):
        style=ParagraphStyle('p',fontName='PortPlanNoto',fontSize=size,leading=leading or size*1.55,textColor=ink,wordWrap='CJK')
        obj=Paragraph(escape(str(text)).replace('\n','<br/>'),style);_,h=obj.wrap(w,700)
        obj.drawOn(c,x,y-h);return y-h
    for index,slide in enumerate(slides,1):
        c.setFillColor(bg);c.rect(0,0,1280,720,stroke=0,fill=1)
        c.setFillColor(green);c.setFont('PortPlanNoto',13);c.drawString(62,672,f'{index:02d} / PORTPLAN V10')
        y=paragraph(slide['title'],62,642,1156,33,43)-15
        y=paragraph(slide['text'],62,y,1156,18,27)-19
        if slide['flow']:
            w=(1156-44)/3;h=118
            for j,text in enumerate(slide['flow']):
                x=62+j*(w+22);c.setFillColor(light);c.roundRect(x,y-h,w,h,7,stroke=0,fill=1);paragraph(text,x+18,y-22,w-36,23,34)
                if j<2:c.setStrokeColor(green);c.setLineWidth(2);c.line(x+w+4,y-h/2,x+w+18,y-h/2);c.line(x+w+13,y-h/2+5,x+w+18,y-h/2);c.line(x+w+13,y-h/2-5,x+w+18,y-h/2)
            y-=h+24
        if slide['chart']:
            for row in slide['chart']:
                paragraph(row['label'],62,y-5,252,17,23)
                c.setFillColor(light);c.rect(322,y-24,580,20,stroke=0,fill=1)
                c.setFillColor(green);c.rect(322,y-24,580*row['value']/row['denominator'] if row['denominator'] else 0,20,stroke=0,fill=1)
                paragraph(f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y-4,294,15,22);y-=53
            y-=8
        if slide['table']:
            headers,rows=slide['table'];dense=len(rows)>7;font_size=11 if dense else 15
            style=ParagraphStyle('cell',fontName='PortPlanNoto',fontSize=font_size,leading=font_size*1.4,textColor=ink,wordWrap='CJK')
            data=[[Paragraph(escape(str(v)),style) for v in row] for row in [headers,*rows]]
            widths=[1156/len(headers)]*len(headers)
            if len(headers)==2:widths=[550,606]
            table=Table(data,colWidths=widths,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),light),('VALIGN',(0,0),(-1,-1),'TOP'),
                ('BOTTOMPADDING',(0,0),(-1,-1),5 if dense else 9),('TOPPADDING',(0,0),(-1,-1),5 if dense else 9),('LEFTPADDING',(0,0),(-1,-1),10),
                ('LINEBELOW',(0,0),(-1,-1),0.5,HexColor('#d3e0d3'))]))
            _,h=table.wrap(1156,600);require(y-h>100,f'PDF table overflows page {index}: {y-h:.1f}');table.drawOn(c,62,y-h);y-=h+18
        note_style=ParagraphStyle('note',fontName='PortPlanNoto',fontSize=12,leading=19,textColor=ink,wordWrap='CJK')
        note=Paragraph(escape(slide['note']),note_style);_,h=note.wrap(1116,300)
        require(y-h-28>64,f'PDF note overflows page {index}: {y-h-28:.1f}')
        c.setFillColor(light);c.roundRect(62,y-h-28,1156,h+25,5,stroke=0,fill=1);note.drawOn(c,82,y-h-14)
        c.setFillColor(HexColor('#638176'));c.setFont('PortPlanNoto',10)
        c.drawString(62,32,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인');c.drawRightString(1218,32,f'PORTPLAN V10 / {index:02d} / {len(slides):02d}');c.showPage()
    c.save();return buffer.getvalue()


def pptx_document(slides):
    """Editable native text, tables and vector shapes; no raster page pictures."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt
    from reportlab.pdfbase import pdfmetrics

    require(len(slides)==18, 'Editable deck must contain exactly 18 slides')
    # pdf_document normally registers the same outline metrics first; direct
    # renderer use remains possible without rendering a PDF or writing a font.
    if 'PortPlanNoto' not in pdfmetrics.getRegisteredFontNames():
        register_noto(json.dumps(slides,ensure_ascii=False)+' PORTPLAN V10 AI 작성 준비 자료 본인 검토 현장 검증 미확인 3/3 미확보 / 0123456789')
    presentation=Presentation();presentation.slide_width=Inches(1280/96);presentation.slide_height=Inches(720/96)
    presentation.core_properties.title='PortPlan V10 · Korean competition preparation'
    presentation.core_properties.author='AI-assisted PortPlan research; entrant verification pending'
    presentation.core_properties.subject='Editable result-driven research deck; synthetic daily model'
    colors={'background':'F7FAF5','ink':'173C40','green':'35826D','light':'E6F0E4','muted':'638176'}
    unit=lambda px: Inches(px/96)

    def wrapped(value,width,size):
        lines=[]
        for raw in str(value).split('\n'):
            line='';used=0
            for char in raw:
                advance=pdfmetrics.stringWidth(char,'PortPlanNoto',size)
                if line and used+advance>width:
                    lines.append(line.rstrip());line='';used=0
                line+=char;used+=advance
            lines.append(line.rstrip())
        return '\n'.join(lines),len(lines)

    def style_text(frame,text,size,leading,color='ink'):
        frame.clear();frame.text=text;frame.word_wrap=False;frame.auto_size=MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        frame.vertical_anchor=MSO_ANCHOR.TOP
        for paragraph in frame.paragraphs:
            paragraph.font.name='Noto Sans CJK KR';paragraph.font.size=Pt(size*.75)
            paragraph.font.color.rgb=RGBColor.from_string(colors[color])
            paragraph.line_spacing=Pt(leading*.75);paragraph.space_before=Pt(0);paragraph.space_after=Pt(0)
            properties=paragraph._p.get_or_add_pPr().get_or_add_defRPr()
            east_asian=properties.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ea')
            if east_asian is None:
                east_asian=OxmlElement('a:ea');properties.append(east_asian)
            east_asian.set('typeface','Noto Sans CJK KR')

    def text(slide,value,x,y,width,size=18,leading=27,color='ink',height=None,name='Editable text'):
        content,line_count=wrapped(value,width-2,size)
        measured=line_count*leading+3
        box=slide.shapes.add_textbox(unit(x),unit(y),unit(width),unit(height or measured));box.name=name
        style_text(box.text_frame,content,size,leading,color)
        return y+(height or measured)

    def rect(slide,x,y,width,height,color='light',kind=MSO_SHAPE.RECTANGLE,name='Editable vector'):
        shape=slide.shapes.add_shape(kind,unit(x),unit(y),unit(width),unit(height));shape.name=name
        shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(colors[color]);shape.line.fill.background()
        return shape

    for index,source in enumerate(slides,1):
        slide=presentation.slides.add_slide(presentation.slide_layouts[6]);slide.background.fill.solid()
        slide.background.fill.fore_color.rgb=RGBColor.from_string(colors['background'])
        text(slide,f'{index:02d} / PORTPLAN V10',62,31,1156,13,20,'green',name='Slide index')
        y=text(slide,source['title'],62,72,1156,33,43,name='Slide title')+15
        y=text(slide,source['text'],62,y,1156,18,27,name='Main explanation')+19
        if source['flow']:
            width=(1156-44)/3
            for j,label in enumerate(source['flow']):
                x=62+j*(width+22);rect(slide,x,y,width,118,name='Flow node '+str(j+1))
                rect(slide,x,y,5,118,'green',name='Flow node accent')
                text(slide,label,x+18,y+22,width-36,23,34,height=80,name='Editable flow label')
                if j<2:rect(slide,x+width+4,y+52,14,14,'green',MSO_SHAPE.CHEVRON,'Editable flow arrow')
            y+=142
        if source['chart']:
            for row in source['chart']:
                text(slide,row['label'],62,y+5,252,17,23,name='Chart method label')
                rect(slide,322,y+4,580,20,name='Editable completion bar track')
                width=580*row['value']/row['denominator'] if row['denominator'] else 0
                if width:rect(slide,322,y+4,width,20,'green',name='Editable completion bar')
                text(slide,f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}",924,y+4,294,15,22,name='Chart exact value')
                y+=53
            y+=8
        if source['table']:
            headers,rows=source['table'];dense=len(rows)>7;size=11 if dense else 15;leading=size*1.4;padding=5 if dense else 9
            widths=[1156/len(headers)]*len(headers)
            if len(headers)==2:widths=[550,606]
            formatted=[];heights=[]
            for row in [headers,*rows]:
                cells=[wrapped(value,widths[j]-20,size) for j,value in enumerate(row)]
                formatted.append([value for value,_ in cells]);heights.append(max(count for _,count in cells)*leading+2*padding+3)
            height=sum(heights);require(y+height<620,f'Editable table overflows slide {index}')
            native=slide.shapes.add_table(len(formatted),len(headers),unit(62),unit(y),unit(1156),unit(height));native.name='Editable evidence table'
            table=native.table;table.first_row=True
            for j,width in enumerate(widths):table.columns[j].width=unit(width)
            for i,row in enumerate(formatted):
                table.rows[i].height=unit(heights[i])
                for j,value in enumerate(row):
                    cell=table.cell(i,j);cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(colors['light' if i==0 else 'background'])
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
                    style_text(cell.text_frame,value,size,leading)
                    # Keep table padding after generic text-frame styling.
                    cell.margin_left=cell.margin_right=unit(10);cell.margin_top=cell.margin_bottom=unit(padding)
            y+=height+18
        note,note_lines=wrapped(source['note'],1116,12);note_height=note_lines*19+3
        require(y+note_height+28<656,f'Editable note overflows slide {index}')
        rect(slide,62,y,1156,note_height+25,name='Editable boundary note background')
        text(slide,note,82,y+11,1116,12,19,height=note_height+3,name='Evidence boundary note')
        text(slide,'AI 작성 준비 자료 · 본인 검토·현장 검증 미확인',62,678,900,10,15,'muted',name='Contribution footer')
        text(slide,f'PORTPLAN V10 / {index:02d} / {len(slides):02d}',1040,678,178,10,15,'muted',name='Page footer')
        for shape in slide.shapes:
            require(shape.shape_type!=MSO_SHAPE_TYPE.PICTURE,'Raster slide pictures are prohibited in the editable deck')
            require(shape.left>=0 and shape.top>=0 and shape.left+shape.width<=presentation.slide_width+1 and
                    shape.top+shape.height<=presentation.slide_height+1,f'Editable shape escapes slide {index}: {shape.name}')
    require(len(presentation.slides)==18,'Editable deck page count changed')
    output=BytesIO();presentation.save(output);return output.getvalue()


def review_generated(evidence, slides):
    """Inspect saved artifacts; do not solve, replay, or modify formal evidence."""
    from datetime import datetime
    from zoneinfo import ZoneInfo
    import unicodedata
    import fitz
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    normalize = lambda value: ''.join(unicodedata.normalize('NFKC', str(value)).split())
    paths = [ROOT / 'docs' / name for name in OUTPUTS]
    require(all(path.is_file() for path in paths), 'Missing generated material')
    document = fitz.open(paths[1])
    presentation = Presentation(paths[5])
    require(len(document) == len(presentation.slides) == len(slides) == 18, 'PDF/PPTX/source page counts differ')
    review = {'schemaVersion': 'portplan-v10-material-review', 'valid': True,
              'reviewedAtKST': datetime.now(ZoneInfo('Asia/Seoul')).isoformat(timespec='seconds'),
              'readOnlyFormalEvidence': True, 'solverRuns': 0, 'newScenarios': 0,
              'formalIndexSha256': evidence['wrapper']['formalIndex']['sha256'],
              'evidenceHashes': evidence['hashes'],
              'generatorSha256': sha(Path(__file__)),
              'slideObjectSha256': hashlib.sha256(json.dumps(slides, ensure_ascii=False, sort_keys=True,
                                                           separators=(',', ':')).encode()).hexdigest(),
              'pdfPageCount': len(document), 'pptxSlideCount': len(presentation.slides),
              'sameSourceSlideObjects': True, 'sourceTextCoverage': True,
              'pdfImages': 0, 'pptxPictures': 0, 'pptxEditableTextShapes': 0, 'pptxEditableTables': 0,
              'pages': [],
              'scope': 'Saved-artifact text, exact native table cells, vector/editable objects, page counts, PDF text bounds and PPTX shape bounds; renderer rejects table/note overflow before writing. PDF/PPTX share the same source slide objects and bound evidence. No PowerPoint application rendering, browser interaction, external clock certification, field validation or entrant understanding is certified.'}
    frozen_generator=(ROOT/evidence['wrapper']['formalIndex']['path']).parent.parent/'source/scripts/build-defense-v10.py'
    require(frozen_generator.is_file(), 'Frozen material-generator source is missing from restored evidence')
    review['frozenGeneratorSha256']=sha(frozen_generator)
    review['generatorRelationToFrozenSource']=('identical' if review['frozenGeneratorSha256']==review['generatorSha256'] else
        'Post-freeze presentation and material-QA changes only: direct native/V7 summaries, precise intervals and saved-output checks. Formal methods, protocol, inputs, receipts and independent reviews remain unchanged.')
    report_text = paths[2].read_text()
    chinese_text = paths[6].read_text()
    expected_ids = Counter()
    adverse_id_count = 0
    completion_loss_id_count = 0
    for panel in evidence['summary']['panels'].values():
        for pair in panel['comparisons'].values():
            for key in ('total', 'incremental'):
                for field in ('negativeScenarioIds', 'negativeEpisodeIds'):
                    expected_ids.update(pair['fees'][key][field])
                    adverse_id_count += len(pair['fees'][key][field])
                require(ci(pair['fees'][key]['uncertainty']['weightedSavingsPct95CI']) in report_text,
                        'Conditional fee interval is absent from Korean report')
                require(ci(pair['fees'][key]['uncertainty']['weightedSavingsPct95CI']) in chinese_text,
                        'Conditional fee interval is absent from Chinese report')
            for field in ('gainedEpisodeIds', 'lostEpisodeIds', 'gainedFromKnownFailureEpisodeIds', 'gainedFromUnknownHistoryEpisodeIds'):
                expected_ids.update(pair['completion'][field])
            completion_loss_id_count += len(pair['completion']['lostEpisodeIds'])
            require(ci(pair['completion']['uncertainty']['meanScenarioRateGain95CI']) in report_text,
                    'Completion interval is absent from Korean report')
    for identifier, count in expected_ids.items():
        require(report_text.count('- `'+identifier+'`') == count,
                'Adverse/completion identifier multiplicity differs in Korean report: ' + identifier)
        require(chinese_text.count('- `'+identifier+'`') == count,
                'Adverse/completion identifier multiplicity differs in Chinese report: ' + identifier)
    require('스트레스 · 새 주문 시드' in report_text and '새 충격 스트레스' not in report_text,
            'Stress mechanism wording incorrectly implies a new shock catalog')
    review['koreanReportAdverseIdentifierOccurrences'] = adverse_id_count
    review['koreanReportLostCompletionIdentifierOccurrences'] = completion_loss_id_count
    review['chineseReportAdverseIdentifierOccurrences'] = adverse_id_count
    review['chineseReportLostCompletionIdentifierOccurrences'] = completion_loss_id_count
    require('尚未填写，未验收' in paths[7].read_text() and '尚未填写，未验收' in paths[8].read_text(),
            'Field and entrant forms must remain explicitly blank and unaccepted')
    review['allDeclaredConditionalIntervalsPresent'] = True
    for i, (page, slide, source) in enumerate(zip(document, presentation.slides, slides), 1):
        pdf_text = normalize(page.get_text())
        pptx_strings = []
        images = len(page.get_images())
        require(images == 0, 'Unexpected raster image in vector PDF page ' + str(i))
        review['pdfImages'] += images
        out_of_bounds = []
        for block in page.get_text('dict')['blocks']:
            for line in block.get('lines', []):
                for span in line.get('spans', []):
                    x0, y0, x1, y1 = span['bbox']
                    if min(x0, y0) < -0.5 or x1 > page.rect.width + 0.5 or y1 > page.rect.height + 0.5:
                        out_of_bounds.append(span['text'])
        require(not out_of_bounds, f'PDF text escapes page {i}: {out_of_bounds[:2]}')
        tables = []
        texts = {}
        for shape in slide.shapes:
            require(shape.shape_type != MSO_SHAPE_TYPE.PICTURE, 'PPTX contains a raster slide picture')
            require(shape.left >= 0 and shape.top >= 0 and shape.left + shape.width <= presentation.slide_width + 1 and
                    shape.top + shape.height <= presentation.slide_height + 1, f'PPTX shape escapes slide {i}')
            if shape.has_text_frame:
                pptx_strings.append(shape.text)
                texts[shape.name] = shape.text
                if shape.text.strip(): review['pptxEditableTextShapes'] += 1
            if shape.has_table:
                tables.append(shape.table)
                review['pptxEditableTables'] += 1
                pptx_strings.extend(cell.text for row in shape.table.rows for cell in row.cells)
        pptx_text = normalize(' '.join(pptx_strings))
        required = [source['title'], source['text'], source['note']]
        required.extend(source['flow'] or [])
        for row in source['chart'] or []:
            required.extend([row['label'], f"{row['value']}/{row['denominator']} · 3/3 {row['all3']} · 미확보 {row['unknown']}"])
        if source['table']:
            headers, rows = source['table']
            required.extend(str(value) for row in [headers, *rows] for value in row)
            require(len(tables) == 1, 'Missing native editable table on slide ' + str(i))
            require(len(tables[0].rows) == len(rows) + 1 and len(tables[0].columns) == len(headers), 'PPTX native table dimensions differ')
            for y, row in enumerate([headers, *rows]):
                for x, value in enumerate(row):
                    require(normalize(tables[0].cell(y, x).text) == normalize(value), f'PPTX exact table cell differs at {i}/{y}/{x}')
        else:
            require(not tables, 'Unexpected native table on slide ' + str(i))
        require(normalize(texts.get('Slide title')) == normalize(source['title']) and
                normalize(texts.get('Main explanation')) == normalize(source['text']) and
                normalize(texts.get('Evidence boundary note')) == normalize(source['note']), 'PPTX main text differs on slide ' + str(i))
        for value in required:
            expected = normalize(value)
            require(expected in pdf_text, f'PDF source text absent on page {i}: {str(value)[:90]}')
            require(expected in pptx_text, f'PPTX source text absent on page {i}: {str(value)[:90]}')
        review['pages'].append({'page': i, 'title': source['title'].replace('\n', ' '),
                                'requiredTextItems': len(required), 'sourceTextCovered': True,
                                'pdfTextWithinPage': True, 'pptxShapesWithinSlide': True,
                                'editableNativeTableCount': len(tables)})
    require(review['pptxEditableTextShapes'] > 100 and review['pptxEditableTables'] >= 8,
            'Deck lacks the expected editable native text and tables')
    contact_path = ROOT / 'artifacts/v10-materials-contact-sheet.png'
    contact_path.parent.mkdir(parents=True, exist_ok=True)
    sheet = fitz.open()
    contact = sheet.new_page(width=1920, height=2292)
    for i in range(len(document)):
        column, row = i % 3, i // 3
        x, y = column * 640, row * 382
        contact.show_pdf_page(fitz.Rect(x+2, y+2, x+638, y+360), document, i)
        contact.insert_text((x+8, y+375), f'PORTPLAN V10 / {i+1:02d} / 18', fontsize=11)
    contact.get_pixmap(alpha=False).save(contact_path)
    sheet.close(); document.close()
    review['outputHashes'] = {str(path.relative_to(ROOT)): sha(path) for path in paths}
    review['contactSheet'] = {'path': str(contact_path.relative_to(ROOT)), 'sha256': sha(contact_path),
                              'scope': 'Full 18-page PDF contact sheet for human visual inspection; no PowerPoint application screenshot.'}
    destination = ROOT / 'data/material-review-v10.json'
    destination.write_text(json.dumps(review, ensure_ascii=False, indent=2)+'\n')
    return review


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--require-results',action='store_true',help='Completeness is always required; retained for release commands.')
    parser.add_argument('--check-only',action='store_true',help='Validate evidence and the narrative without rendering or writing materials.')
    parser.add_argument('--bootstrap-review',default='data/bootstrap-review-v10.json',help='Required independent interval review bound to the same formal index SHA.')
    parser.add_argument('--replace-generated',action='store_true',help='Explicitly regenerate these V10 materials only.')
    args=parser.parse_args(argv)
    evidence=verified_inputs(bootstrap_review=args.bootstrap_review)
    slides=build_slides(evidence)
    if args.check_only:
        print(json.dumps({'valid':True,'pages':len(slides),'writes':0,'evidenceHashes':evidence['hashes']}));return
    targets=[ROOT/'docs'/name for name in OUTPUTS]
    if not args.replace_generated:require(not any(path.exists() for path in targets),'V10 materials already exist; use --replace-generated explicitly')
    report,english,contribution,chinese,field_form,entrant_form=material_texts(evidence,slides)
    pdf=pdf_document(slides)
    pptx=pptx_document(slides)
    require(pdf.startswith(b'%PDF-'), 'PDF generation failed')
    require(pptx.startswith(b'PK'), 'Editable presentation generation failed')
    outputs={targets[0]:html_document(slides).encode(),targets[1]:pdf,targets[2]:report.encode(),targets[3]:english.encode(),
             targets[4]:contribution.encode(),targets[5]:pptx,targets[6]:chinese.encode(),targets[7]:field_form.encode(),targets[8]:entrant_form.encode()}
    for path,data in outputs.items():path.write_bytes(data)
    material_review=review_generated(evidence,slides)
    print(json.dumps({'valid':True,'pages':len(slides),'written':[str(p.relative_to(ROOT)) for p in outputs],
        'outputHashes':{str(p.relative_to(ROOT)):sha(p) for p in outputs},
        'materialReview':'data/material-review-v10.json','editableTextShapes':material_review['pptxEditableTextShapes'],
        'editableTables':material_review['pptxEditableTables']},ensure_ascii=False))


if __name__=='__main__':
    main()
